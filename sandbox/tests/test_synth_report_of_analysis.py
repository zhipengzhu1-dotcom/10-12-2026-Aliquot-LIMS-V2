import pytest
from openpyxl import load_workbook

from synth.lab import LAB_NAME


def labelled(ws):
    """Column A label -> column B value, for the Client and metadata blocks."""
    return {
        row[0].value: row[1].value
        for row in ws.iter_rows(max_col=2)
        if isinstance(row[0].value, str)
    }


def receiving_log_row(generated, number):
    ws = load_workbook(generated / "receiving_log.xlsx").worksheets[0]
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == number:
            return dict(zip(header, row, strict=True))
    raise AssertionError(f"no Receiving Log row #{number}")


def test_metadata_equals_the_receiving_log_row(generated, manifest):
    roa = labelled(load_workbook(generated / "report_of_analysis.xlsx").active)
    log = receiving_log_row(generated, manifest["containers"][0]["number"])
    assert roa["Product"] == log["PRODUCT NAME"]
    assert roa["Lot"] == log["LOT#"]
    assert roa["Product Code"] == log["Product Code"]
    assert roa["Product Type"] == log["TYPE"]
    assert roa["Date Received"] == log["Rec. DATE"]
    assert roa["Time Received"] == log["Rec. Time"]
    assert roa["Client"] == log["CLIENT"]


def test_metadata_the_receiving_log_has_no_column_for(generated, manifest):
    roa = labelled(load_workbook(generated / "report_of_analysis.xlsx").active)
    assert roa["Lab Reference"] == manifest["lab_reference"]
    assert roa["Received By"] == manifest["received_by"]
    assert roa["Document Revision"] == manifest["document_revision"]


def test_layout_has_letterhead_results_and_signature_area(generated, manifest):
    ws = load_workbook(generated / "report_of_analysis.xlsx").active
    assert ws["A1"].value == LAB_NAME
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    for heading in ("REPORT OF ANALYSIS", "Test", "Result", "Analyzed By"):
        assert heading in values
    for person in (manifest["analyst"], manifest["reviewer"], manifest["approver"]):
        assert person in values


def test_the_result_is_the_mean_average_capsule_weight(generated, manifest):
    ws = load_workbook(generated / "report_of_analysis.xlsx").active
    header = next(r for r in ws.iter_rows() if r[0].value == "Test")
    result_row = [c.value for c in ws[header[0].row + 1]]
    result = result_row[[c.value for c in header].index("Result")]
    weights = [w["weight_of_20_capsules_g"] for w in manifest["weighings"]]
    assert result == pytest.approx(sum(weights) / 20 / len(weights), abs=5e-5)
