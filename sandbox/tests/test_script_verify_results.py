"""verify_results on the Result Worksheet that compute_results wrote for seed 1's
expected record, and on a copy of it with one typed average, run the way the
Sandbox runs it but without Docker."""

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_script_compute_results import (
    WEIGHINGS,
    WEIGHINGS_ID,
    nitrosamine_results,
    results,
)
from test_script_create_work_order import XLSX, put, run

WORKSHEET_ID = "d" * 64


def worksheet(tmp_path: Path) -> Path:
    made, outputs = results(tmp_path / "stage4")
    assert made.returncode == 0, made.stderr
    return outputs / "result_worksheet.xlsx"


def verification(tmp_path: Path, sheet: Path, weighings: Path = WEIGHINGS):
    inputs, outputs = tmp_path / "inputs", tmp_path / "outputs"
    put(inputs, WORKSHEET_ID, sheet, "result_worksheet.xlsx")
    put(inputs, WEIGHINGS_ID, weighings, "weighings.json")
    result = run(
        "verify_results",
        inputs,
        outputs,
        [
            {
                "artifact_id": WORKSHEET_ID,
                "filename": "result_worksheet.xlsx",
                "artifact_type": "ResultWorksheet",
                "media_type": XLSX,
            },
            {
                "artifact_id": WEIGHINGS_ID,
                "filename": "weighings.json",
                "artifact_type": "Other",
                "media_type": "application/json",
            },
        ],
    )
    return result, outputs


def record_of(outputs: Path) -> dict:
    return json.loads((outputs / "verification_record.json").read_text())


def test_it_passes_the_worksheet_compute_results_wrote(tmp_path):
    result, outputs = verification(tmp_path, worksheet(tmp_path))

    assert result.returncode == 0, result.stderr
    assert json.loads((outputs / "manifest.json").read_text()) == {
        "verification_record.json": {
            "artifact_type": "VerificationRecord",
            "media_type": "application/json",
        }
    }
    mean = (7.0886 / 20 + 7.2103 / 20) / 2
    assert record_of(outputs) == {
        "schema": "verification_record/1",
        "worksheet": WORKSHEET_ID,
        "weighings": WEIGHINGS_ID,
        "tolerance_g": 0.0001,
        "work_order": "WO-2026-0001",
        "rows": [
            {
                "container": container,
                "row": row,
                "formula": f"=F{row}/20",
                "recomputed_g": pytest.approx(of20 / 20, abs=1e-12),
                "weighing_g": pytest.approx(of20 / 20, abs=1e-12),
                "difference_g": pytest.approx(0, abs=1e-12),
                "within_tolerance": True,
            }
            for container, row, of20 in ((1, 6, 7.0886), (2, 7, 7.2103))
        ],
        "mean": {
            "row": 8,
            "formula": "=AVERAGE(H6:H7)",
            "recomputed_g": pytest.approx(mean, abs=1e-12),
            "weighing_g": pytest.approx(mean, abs=1e-12),
            "difference_g": pytest.approx(0, abs=1e-12),
            "within_tolerance": True,
        },
        "missing_containers": [],
        "tolerance_relative": 1e-06,
        "nitrosamines": [],
        "missing_readings": [],
        "verdict": "pass",
    }
    assert "WO-2026-0001: pass" in result.stdout


def test_it_fails_a_copy_with_a_typed_average_that_disagrees_with_its_weights(
    tmp_path,
):
    copy = tmp_path / "typed.xlsx"
    wb = load_workbook(worksheet(tmp_path))
    wb.active["H6"] = 0.36443
    wb.save(copy)

    result, outputs = verification(tmp_path, copy)

    assert result.returncode == 0, result.stderr
    record = record_of(outputs)
    first, second = record["rows"]
    assert first == {
        "container": 1,
        "row": 6,
        "formula": None,
        "recomputed_g": 0.36443,
        "weighing_g": pytest.approx(0.35443, abs=1e-12),
        "difference_g": pytest.approx(0.01, abs=1e-12),
        "within_tolerance": False,
    }
    assert second["within_tolerance"] is True
    assert record["mean"]["within_tolerance"] is False, "the mean averages H6"
    assert record["verdict"] == "fail"


def test_it_fails_a_worksheet_missing_a_weighed_container(tmp_path):
    extra = tmp_path / "extra.json"
    extra.write_text(
        json.dumps(
            [
                *json.loads(WEIGHINGS.read_text()),
                {"container": 3, "weight_of_20_g": 7.1, "weight_taken_g": 0.35},
            ]
        )
    )

    result, outputs = verification(tmp_path, worksheet(tmp_path), extra)

    assert result.returncode == 0, result.stderr
    record = record_of(outputs)
    assert record["missing_containers"] == [3]
    assert record["verdict"] == "fail"


def test_it_refuses_a_worksheet_without_the_columns_row(tmp_path):
    blank = tmp_path / "blank.xlsx"
    wb = load_workbook(worksheet(tmp_path))
    wb.active.delete_rows(5)
    wb.save(blank)

    result, outputs = verification(tmp_path, blank)

    assert result.returncode == 1
    assert "no row of the Result Worksheet's columns" in result.stderr
    assert list(outputs.iterdir()) == []


def nitrosamine_worksheet(tmp_path: Path) -> tuple[Path, Path]:
    made, outputs = nitrosamine_results(tmp_path / "stage4")
    assert made.returncode == 0, made.stderr
    return outputs / "result_worksheet.xlsx", tmp_path / "stage4" / "readings.json"


def test_it_recomputes_each_nitrosamine_result_and_judgement_from_the_readings(
    tmp_path,
):
    sheet, readings = nitrosamine_worksheet(tmp_path)

    result, outputs = verification(tmp_path, sheet, readings)

    assert result.returncode == 0, result.stderr
    record = record_of(outputs)
    assert record["verdict"] == "pass"
    assert (record["rows"], record["mean"], record["missing_readings"]) == (
        [],
        None,
        [],
    )
    first, *_, last = record["nitrosamines"]
    assert first == {
        "row": 6,
        "container": 1,
        "analyte": "NDEA",
        "limit": "46.4 ng/day",
        "below_reporting_limit": False,
        "content_ppm": {
            "formula": "=G6/F6",
            "recomputed": pytest.approx(0.03),
            "expected": pytest.approx(0.03),
            "within_tolerance": True,
        },
        "intake_ng_day": {
            "formula": "=I6*J6",
            "recomputed": pytest.approx(45.0),
            "expected": pytest.approx(45.0),
            "within_tolerance": True,
        },
        "judgement": "Pass",
        "expected_judgement": "Pass",
    }
    assert (last["container"], last["analyte"], last["intake_ng_day"]) == (
        2,
        "NDMA",
        None,
    )
    assert [n["judgement"] for n in record["nitrosamines"]] == [
        "Pass",
        "Fail",
        "Pass",
        "Pass",
    ]


def test_it_fails_a_typed_measurement_or_a_judgement_that_disagrees(tmp_path):
    sheet, readings = nitrosamine_worksheet(tmp_path)
    copy = tmp_path / "typed.xlsx"
    wb = load_workbook(sheet)
    wb.active["G7"] = 0.9
    wb.active["M9"] = "Fail"
    wb.save(copy)

    result, outputs = verification(tmp_path, copy, readings)

    assert result.returncode == 0, result.stderr
    record = record_of(outputs)
    ndea, ndma, _, last = record["nitrosamines"]
    assert ndea["content_ppm"]["within_tolerance"] is True
    assert ndma["content_ppm"]["recomputed"] == pytest.approx(0.018)
    assert ndma["content_ppm"]["within_tolerance"] is False
    assert (last["judgement"], last["expected_judgement"]) == ("Fail", "Pass")
    assert record["verdict"] == "fail"
