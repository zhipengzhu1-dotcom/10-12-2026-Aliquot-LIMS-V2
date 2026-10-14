"""log_receipt with the lab's Receiving Log as input: the Shadow in the real layout.

The scrubbed Receiving Log in `fixtures/real-layout/` stands in for the lab's
workbook, and the real-layout Fixture's record is the receipt.
"""

import hashlib
import json
import shutil
import subprocess
import sys
from datetime import datetime, time
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / "sandbox" / "scripts" / "log_receipt" / "main.py"
REAL_LAYOUT = ROOT / "fixtures" / "real-layout"
RECORD = REAL_LAYOUT / "submission-form-seed4.expected.json"
LAB_LOG = REAL_LAYOUT / "receiving-log.xlsx"

RECORD_ID = "a" * 64
LOG_ID = "e" * 64
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def run(tmp_path: Path, record: dict, lab_log: Path | None = LAB_LOG):
    inputs, outputs = tmp_path / "inputs", tmp_path / "outputs"
    outputs.mkdir(parents=True)
    (inputs / RECORD_ID).mkdir(parents=True)
    (inputs / RECORD_ID / "submission.json").write_text(json.dumps(record))
    manifest = [
        {
            "artifact_id": RECORD_ID,
            "filename": "submission.json",
            "artifact_type": "SubmissionRecord",
            "media_type": "application/json",
        }
    ]
    if lab_log is not None:
        (inputs / LOG_ID).mkdir()
        staged = inputs / LOG_ID / "receiving-log.xlsx"
        shutil.copyfile(lab_log, staged)
        staged.chmod(0o444)
        manifest.append(
            {
                "artifact_id": LOG_ID,
                "filename": "receiving-log.xlsx",
                "artifact_type": "ReceivingLog",
                "media_type": XLSX,
            }
        )
    (inputs / "manifest.json").write_text(json.dumps(manifest))
    result = subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--inputs",
            str(inputs),
            "--outputs",
            str(outputs),
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    return result, outputs, inputs


def seed4() -> dict:
    return json.loads(RECORD.read_text())


def row(ws, number: int, last: str) -> list:
    return [cell.value for cell in ws[f"A{number}:{last}{number}"][0]]


def every_value(path: Path) -> dict[tuple[str, str], object]:
    wb = load_workbook(path)
    return {
        (ws.title, cell.coordinate): getattr(cell.value, "text", cell.value)
        for ws in wb.worksheets
        for line in ws.iter_rows()
        for cell in line
        if cell.value is not None
    }


def test_the_shadow_adds_one_row_per_container_to_the_year_and_a_new_days_sheet(
    tmp_path,
):
    result, outputs, inputs = run(tmp_path, seed4())

    assert result.returncode == 0, result.stderr
    manifest = json.loads((outputs / "manifest.json").read_text())
    assert manifest["receiving_log_shadow.xlsx"]["artifact_type"] == (
        "ShadowReceivingLog"
    )
    assert manifest["shadow.json"]["artifact_type"] == "Other"
    refs = ["LR-2026-0001", "LR-2026-0002", "LR-2026-0003"]
    shadow = load_workbook(outputs / "receiving_log_shadow.xlsx")

    year = shadow["2026"]
    assert year.max_row == 41, "the three rows follow the lab's last row, 38"
    for i, ref in enumerate(refs):
        assert row(year, 39 + i, "T") == [
            str(36 + i),
            datetime(2026, 9, 28),
            time(13, 17),
            "TLB-5077",
            "Tallowby Zinc Picolinate 30 mg Capsules",
            "FP",
            "LT-267759",
            "Tallowby Botanicals LLC",
            None,
            1,
            "Calla Halloway",
            "IS",
            ref,
            *[None] * 7,
        ], f"year row {39 + i}"

    assert shadow.sheetnames == ["2026", "July 23", "September 28"]
    day = shadow["September 28"]
    assert day["F1"].value == shadow["July 23"]["F1"].value, "the copy keeps the title"
    for i, ref in enumerate(refs):
        assert row(day, 3 + i, "R") == [
            1 + i,
            datetime(2026, 9, 28),
            "Tallowby Zinc Picolinate 30 mg Capsules",
            "FP",
            "LT-267759",
            "Tallowby Botanicals LLC",
            None,
            1,
            "IS",
            ref,
            *[None] * 8,
        ], f"day row {3 + i}"
    assert row(day, 6, "R")[1:] == [None] * 17, "the copied sheet's rows are cleared"

    held = json.loads((outputs / "shadow.json").read_text())
    assert held["year"] == {"sheet": "2026", "rows": [39, 40, 41]}
    assert held["day"] == {"sheet": "September 28", "rows": [3, 4, 5], "made": True}
    sheets = {s["name"]: s for s in held["sheets"]}
    assert sheets["September 28"]["before"] is None
    assert sheets["2026"]["after"]["G40"] == "LT-267759"
    assert sheets["2026"]["after"]["B40"] == "2026-09-28"
    assert sheets["2026"]["before"] == {
        k: v for k, v in sheets["2026"]["after"].items() if int(k[1:]) < 39
    }


def test_each_shadow_row_holds_its_own_samples_values(tmp_path):
    record = json.loads(
        (
            REAL_LAYOUT / "improvised" / "received-2026-06-26-form-2.expected.json"
        ).read_text()
    )

    result, outputs, _ = run(tmp_path, record)

    assert result.returncode == 0, result.stderr
    shadow = load_workbook(outputs / "receiving_log_shadow.xlsx")
    year, day = shadow["2026"], shadow["June 26"]
    held = json.loads((outputs / "shadow.json").read_text())
    expected = [
        (s["product_code"], s["sample_description"], s["lot"])
        for s in record["samples"]
    ]
    assert [
        (year.cell(r, 4).value, year.cell(r, 5).value, year.cell(r, 7).value)
        for r in held["year"]["rows"]
    ] == expected
    assert [
        (day.cell(r, 3).value, day.cell(r, 5).value) for r in held["day"]["rows"]
    ] == [(name, lot) for _, name, lot in expected]
    assert [year.cell(r, 13).value for r in held["year"]["rows"]] == [
        f"LR-2026-000{n}" for n in range(1, 9)
    ]


def test_no_cell_of_the_lab_workbook_changes_in_the_shadow_or_on_disk(tmp_path):
    before = hashlib.sha256(LAB_LOG.read_bytes()).hexdigest()

    result, outputs, inputs = run(tmp_path, seed4())

    assert result.returncode == 0, result.stderr
    staged = inputs / LOG_ID / "receiving-log.xlsx"
    assert hashlib.sha256(staged.read_bytes()).hexdigest() == before
    lab = every_value(LAB_LOG)
    shadow = every_value(outputs / "receiving_log_shadow.xlsx")
    changed = {key for key, value in lab.items() if shadow.get(key) != value}
    assert changed == set()


def test_a_receipt_on_a_day_with_a_sheet_fills_that_sheets_first_blank_rows(
    tmp_path,
):
    record = seed4()
    record["receipt"]["date_received"] = "2026-07-23"

    result, outputs, _ = run(tmp_path, record)

    assert result.returncode == 0, result.stderr
    shadow = load_workbook(outputs / "receiving_log_shadow.xlsx")
    assert shadow.sheetnames == ["2026", "July 23"]
    day = shadow["July 23"]
    assert day["B3"].value.text == "=#REF!", "row 3 holds the lab's formula, kept"
    assert [day.cell(r, 1).value for r in (4, 5, 6)] == [2, 3, 4]
    assert [day.cell(r, 10).value for r in (4, 5, 6)] == [
        "LR-2026-0001",
        "LR-2026-0002",
        "LR-2026-0003",
    ]
    held = json.loads((outputs / "shadow.json").read_text())
    assert held["day"] == {"sheet": "July 23", "rows": [4, 5, 6], "made": False}


def test_without_the_lab_workbook_there_is_no_shadow(tmp_path):
    result, outputs, _ = run(tmp_path, seed4(), lab_log=None)

    assert result.returncode == 0, result.stderr
    assert sorted(p.name for p in outputs.iterdir()) == [
        "lims.sqlite",
        "manifest.json",
        "receipt.json",
        "receiving_log_row.xlsx",
    ]


def test_it_refuses_a_workbook_without_the_years_sheet_and_writes_no_shadow(
    tmp_path,
):
    record = seed4()
    record["receipt"]["date_received"] = "2027-01-04"

    result, outputs, _ = run(tmp_path, record)

    assert result.returncode == 1
    assert "has no sheet 2027" in result.stderr
    assert not (outputs / "receiving_log_shadow.xlsx").exists()
