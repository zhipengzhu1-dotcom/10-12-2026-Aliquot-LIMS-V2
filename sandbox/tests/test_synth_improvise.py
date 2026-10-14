"""The improvised real-layout forms: reproducible, covering the kinds the lab sees."""

import json
from pathlib import Path

from openpyxl import load_workbook

from synth.improvise import main, read_log

REAL_LAYOUT = Path(__file__).resolve().parents[2] / "fixtures" / "real-layout"
IMPROVISED = REAL_LAYOUT / "improvised"
LOG = REAL_LAYOUT / "receiving-log.xlsx"
BLANK = REAL_LAYOUT / "submission-form.docx"


def generate(out: Path) -> None:
    main(["--log", str(LOG), "--blank", str(BLANK), "--out", str(out)])


def records() -> dict[str, dict]:
    return {
        path.name.removesuffix(".expected.json"): json.loads(path.read_text())
        for path in sorted(IMPROVISED.glob("*.expected.json"))
    }


def test_the_committed_forms_are_what_the_generator_writes(tmp_path):
    generate(tmp_path)
    written = sorted(p.name for p in tmp_path.iterdir())
    committed = sorted(
        p.name for p in IMPROVISED.iterdir() if not p.name.endswith(".known-gap.json")
    )
    assert written == committed
    for name in written:
        assert (tmp_path / name).read_bytes() == (IMPROVISED / name).read_bytes(), name


def test_each_form_has_its_expected_record_and_one_weighing_per_container():
    forms = sorted(p.stem for p in IMPROVISED.glob("*.docx"))
    assert len(forms) >= 6
    for stem, record in records().items():
        weighings = json.loads((IMPROVISED / f"{stem}.weighings.json").read_text())
        containers = sum(int(s["no_of_containers"]) for s in record["samples"])
        assert [w["container"] for w in weighings] == list(range(1, containers + 1))
    assert sorted(records()) == forms


def test_each_line_of_the_log_is_one_sample_line_in_one_container():
    lines = sum(len(r["samples"]) for r in records().values())
    assert lines == len(read_log(LOG)) + 1, "the log's rows and the unlogged Sample"
    for stem, record in records().items():
        assert [s["no"] for s in record["samples"]] == [
            str(n) for n in range(1, len(record["samples"]) + 1)
        ], stem
        assert {s["no_of_containers"] for s in record["samples"]} == {"1"}, stem


def test_a_submission_on_two_forms_keeps_each_lots_lines_on_one_form():
    first, second = (
        records()[f"received-2026-06-26-form-{n}"]["samples"] for n in (1, 2)
    )
    lines = {(s["lot"], s["product_code"]) for s in first}
    assert len(first) + len(second) > 10
    assert lines.isdisjoint((s["lot"], s["product_code"]) for s in second)


def test_the_set_covers_the_kinds_the_lab_sees():
    samples = [s for r in records().values() for s in r["samples"]]
    assert {"FP", "RM", "IP"} <= {s["sample_type"] for s in samples}
    lines = {len(r["samples"]) for r in records().values()}
    assert 1 in lines and max(lines) > 1
    classes = {s["tat_class"] for s in samples}
    assert "Regular" in classes and classes - {"Regular"}
    report_to_differs = {
        r["report_to"]["contact"] != r["submitted_by"]["contact"]
        for r in records().values()
    }
    assert report_to_differs == {True, False}
    bill_to_blank = {r["bill_to"] == r["submitted_by"] for r in records().values()}
    assert bill_to_blank == {True, False}


def test_the_faithful_forms_carry_the_logs_rows_and_the_off_ones_do_not():
    logged = {(r.received_on.isoformat(), r.lot, r.product_code) for r in read_log(LOG)}
    for stem, record in records().items():
        faithful = stem.startswith("received-") and "lot-mistyped" not in stem
        for sample in record["samples"]:
            key = (
                record["receipt"]["date_received"],
                sample["lot"],
                sample["product_code"],
            )
            assert (key in logged) == faithful, f"{stem} line {sample['no']}"


def test_the_log_is_read_by_position_from_its_year_sheet():
    rows = read_log(LOG)
    year = load_workbook(LOG)["2026"]
    assert [r.number for r in rows] == list(range(4, year.max_row + 1))
    assert {r.client for r in rows} == {year["H4"].value}
