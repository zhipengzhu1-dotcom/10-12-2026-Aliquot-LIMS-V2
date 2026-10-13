"""render_roa on the Verification Record, Receiving Log row and database that
stages 2 to 5 wrote for seed 1's expected record, run the way the Sandbox runs
it but without Docker."""

import hashlib
import json
import shutil
import sqlite3
from contextlib import closing
from datetime import UTC, date, datetime, time
from pathlib import Path

import pdfplumber
import pytest
from openpyxl import load_workbook
from test_script_compute_results import WEIGHINGS, WEIGHINGS_ID, noted
from test_script_create_work_order import DATABASE_ID, ROW_ID, XLSX, put, query, run
from test_script_verify_results import verification, worksheet

RECORD_ID = "9" * 64


def chain(tmp_path: Path) -> dict[str, Path]:
    """Stages 2 to 5 on seed 1's expected record: the four inputs stage 6 takes."""
    stage5 = tmp_path / "stage5"
    stage5.mkdir()
    made, outputs = verification(stage5, worksheet(stage5))
    assert made.returncode == 0, made.stderr
    return {
        "record": outputs / "verification_record.json",
        "row": stage5 / "stage4" / "receipt-outputs" / "receiving_log_row.xlsx",
        "weighings": WEIGHINGS,
        "database": stage5 / "stage4" / "outputs" / "lims.sqlite",
    }


def render(tmp_path: Path, given: dict[str, Path]):
    inputs, outputs = tmp_path / "inputs", tmp_path / "outputs"
    put(inputs, RECORD_ID, given["record"], "verification_record.json")
    put(inputs, ROW_ID, given["row"], "receiving_log_row.xlsx")
    put(inputs, WEIGHINGS_ID, given["weighings"], "weighings.json")
    put(inputs, DATABASE_ID, given["database"], "lims.sqlite")
    result = run(
        "render_roa",
        inputs,
        outputs,
        [
            {
                "artifact_id": RECORD_ID,
                "filename": "verification_record.json",
                "artifact_type": "VerificationRecord",
                "media_type": "application/json",
            },
            {
                "artifact_id": ROW_ID,
                "filename": "receiving_log_row.xlsx",
                "artifact_type": "ReceivingLogRow",
                "media_type": XLSX,
            },
            {
                "artifact_id": WEIGHINGS_ID,
                "filename": "weighings.json",
                "artifact_type": "Other",
                "media_type": "application/json",
            },
            {
                "artifact_id": DATABASE_ID,
                "filename": "lims.sqlite",
                "artifact_type": "LimsDatabase",
                "media_type": "application/vnd.sqlite3",
            },
        ],
    )
    return result, outputs


@pytest.fixture(scope="module")
def rendered(tmp_path_factory) -> tuple[dict[str, Path], Path]:
    tmp_path = tmp_path_factory.mktemp("render_roa")
    given = chain(tmp_path)
    result, outputs = render(tmp_path, given)
    assert result.returncode == 0, result.stderr
    return given, outputs


def labelled(path: Path) -> dict:
    """Column A label -> column B value, for the Client and metadata blocks.
    The Client's city line has no label, so it reads as ""."""
    ws = load_workbook(path).active
    return {
        row[0].value or "": row[1].value
        for row in ws.iter_rows(min_row=7, max_row=20, max_col=2)
        if row[1].value is not None
    }


def receiving_log_row(path: Path, lab_reference: str) -> dict:
    ws = load_workbook(path).active
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == lab_reference:
            return dict(zip(header, row, strict=True))
    raise AssertionError(f"no Receiving Log row {lab_reference}")


def mean_of(record: Path) -> float:
    return json.loads(record.read_text())["mean"]["recomputed_g"]


def test_it_types_two_reports_the_next_database_and_the_read_back(rendered):
    _, outputs = rendered

    assert json.loads((outputs / "manifest.json").read_text()) == {
        "report_of_analysis.xlsx": {
            "artifact_type": "ReportOfAnalysis",
            "media_type": XLSX,
        },
        "report_of_analysis.pdf": {
            "artifact_type": "ReportOfAnalysis",
            "media_type": "application/pdf",
        },
        "lims.sqlite": {
            "artifact_type": "LimsDatabase",
            "media_type": "application/vnd.sqlite3",
        },
        "report_of_analysis.json": {
            "artifact_type": "Other",
            "media_type": "application/json",
        },
    }


def test_every_metadata_cell_is_the_receiving_log_rows_or_the_samples_rows(rendered):
    given, outputs = rendered
    cells = labelled(outputs / "report_of_analysis.xlsx")
    log = receiving_log_row(given["row"], "LR-2026-0001")

    assert cells == {
        "Client": "Tallowby Botanicals LLC",
        "Attention": "Maren Vantongeren",
        "Address": "39 Juniper Mill Drive",
        "": "Orrin Bay, NY 14999",
        "Lab Reference": "LR-2026-0001",
        "Product": "Tallowby Milk Thistle 250 mg Capsules",
        "Lot": "LT-268123",
        "Product Code": "TLB-1475",
        "Product Type": "60 ct amber glass bottle",
        "Document Revision": "Rev. 03",
        "Date Received": datetime(2026, 1, 6),
        "Time Received": time(15, 17),
        "Received By": "LG",
    }
    for label, column in (
        ("Client", "CLIENT"),
        ("Lab Reference", "#"),
        ("Product", "PRODUCT NAME"),
        ("Lot", "LOT#"),
        ("Product Code", "Product Code"),
        ("Product Type", "TYPE"),
        ("Date Received", "Rec. DATE"),
        ("Time Received", "Rec. Time"),
    ):
        assert cells[label] == log[column], label
    assert query(
        given["database"],
        "SELECT report_to_contact, report_to_address, report_to_city, received_by"
        " FROM samples WHERE lab_reference = 'LR-2026-0001'",
    ) == [(cells["Attention"], cells["Address"], cells[""], cells["Received By"])]


def test_the_result_is_the_records_mean_with_the_test_and_unit_of_the_database(
    rendered,
):
    given, outputs = rendered
    ws = load_workbook(outputs / "report_of_analysis.xlsx").active

    assert [c.value for c in ws[22]][:5] == [
        "Test",
        "Method",
        "Specification",
        "Result",
        "Unit",
    ]
    test, method, spec, result, unit = [c.value for c in ws[23]][:5]
    assert (test, method, spec, unit) == (
        "Average capsule weight",
        "QAL-SOP-114",
        "Report value",
        "g",
    )
    assert result == round(mean_of(given["record"]), 4)
    assert result == pytest.approx((7.0886 / 20 + 7.2103 / 20) / 2, abs=5e-5)
    assert [ws.cell(26, c).value for c in (1, 3)] == ["Rendered By", "Verified By"]
    assert ws.cell(27, 1).value == "render_roa"
    assert ws.cell(27, 3).value == f"Verification Record {RECORD_ID[:12]}, verdict pass"


def test_the_pdf_read_with_pdfplumber_holds_lab_reference_lot_product_code_result(
    rendered,
):
    given, outputs = rendered
    with pdfplumber.open(outputs / "report_of_analysis.pdf") as pdf:
        assert len(pdf.pages) == 1
        text = pdf.pages[0].extract_text()

    result = f"{round(mean_of(given['record']), 4):.4f}"
    for value in ("LR-2026-0001", "LT-268123", "TLB-1475", result):
        assert value in text
    assert "Tallowby Milk Thistle 250 mg Capsules" in text
    assert "2026-01-06" in text


def test_the_reports_row_names_the_sample_and_the_pdf_by_its_artifact_id(rendered):
    given, outputs = rendered
    pdf_id = hashlib.sha256(
        (outputs / "report_of_analysis.pdf").read_bytes()
    ).hexdigest()

    assert query(given["database"], "SELECT * FROM reports") == []
    [(ref, issued, revision, source)] = query(
        outputs / "lims.sqlite", "SELECT * FROM reports"
    )
    assert (ref, revision, source) == ("LR-2026-0001", "Rev. 03", pdf_id)
    assert date.fromisoformat(issued) in {
        datetime.now(UTC).date(),
        date.fromordinal(datetime.now(UTC).date().toordinal() - 1),
    }
    for table in ("lots", "samples", "work_orders", "results"):
        assert query(outputs / "lims.sqlite", f"SELECT * FROM {table}") == query(
            given["database"], f"SELECT * FROM {table}"
        ), table


def test_the_read_back_is_the_receiving_log_row_and_the_saved_sheet(rendered):
    _, outputs = rendered
    read = json.loads((outputs / "report_of_analysis.json").read_text())

    assert read["schema"] == "report_of_analysis/1"
    assert read["receiving_log_row"] == {
        "#": "LR-2026-0001",
        "Rec. DATE": "2026-01-06",
        "Rec. Time": "15:17",
        "Product Code": "TLB-1475",
        "PRODUCT NAME": "Tallowby Milk Thistle 250 mg Capsules",
        "TYPE": "60 ct amber glass bottle",
        "LOT#": "LT-268123",
        "CLIENT": "Tallowby Botanicals LLC",
        "EXP DATE": "",
        "QTY RECD": "1",
    }
    sheet = read["sheet"]
    assert sheet[0][0] == "Quillmere Analytical Laboratories"
    assert sheet[17][:2] == ["Date Received", "2026-01-06"]
    assert sheet[18][:2] == ["Time Received", "15:17"]
    assert (
        sheet[22][3]
        == load_workbook(outputs / "report_of_analysis.xlsx").active["D23"].value
    )


def blank_receiving_log_type(given: dict[str, Path], tmp_path: Path) -> None:
    copy = tmp_path / "row.xlsx"
    wb = load_workbook(given["row"])
    ws = wb.active
    column = [c.value for c in ws[1]].index("TYPE") + 1
    ws.cell(2, column).value = None
    wb.save(copy)
    given["row"] = copy


def blank_samples_received_by(given: dict[str, Path], tmp_path: Path) -> None:
    copy = tmp_path / "carried.sqlite"
    shutil.copyfile(given["database"], copy)
    with closing(sqlite3.connect(copy)) as db:
        db.execute("UPDATE samples SET received_by = ' '")
        db.commit()
    given["database"] = copy


@pytest.mark.parametrize(
    ("blank", "reason"),
    [
        (
            blank_receiving_log_type,
            "Product Type: the Receiving Log row of LR-2026-0001 has no TYPE",
        ),
        (
            blank_samples_received_by,
            "Received By: the samples row of LR-2026-0001 has no received_by",
        ),
    ],
)
def test_it_refuses_a_missing_metadata_value_and_writes_nothing(
    tmp_path, blank, reason
):
    given = chain(tmp_path)
    blank(given, tmp_path)

    result, outputs = render(tmp_path, given)

    assert result.returncode == 1
    assert reason in result.stderr
    assert list(outputs.iterdir()) == []


def test_the_signature_area_carries_the_weighings_notebook_reference(tmp_path):
    given = chain(tmp_path)
    given["weighings"] = noted(tmp_path, *["NB-07 p.31"] * 2)

    result, outputs = render(tmp_path, given)

    assert result.returncode == 0, result.stderr
    ws = load_workbook(outputs / "report_of_analysis.xlsx").active
    assert [ws.cell(31, c).value for c in (1, 2)] == [
        "Notebook Reference",
        "NB-07 p.31",
    ]
    sheet = json.loads((outputs / "report_of_analysis.json").read_text())["sheet"]
    assert sheet[30][:2] == ["Notebook Reference", "NB-07 p.31"]
    with pdfplumber.open(outputs / "report_of_analysis.pdf") as pdf:
        text = pdf.pages[0].extract_text()
    assert "Notebook Reference NB-07 p.31" in text
    assert text.index("Date: ______________") < text.index("NB-07 p.31")


def test_a_fixtures_weighings_put_no_notebook_reference_in_the_report(rendered):
    _, outputs = rendered
    ws = load_workbook(outputs / "report_of_analysis.xlsx").active

    assert ws.max_row == 29
    with pdfplumber.open(outputs / "report_of_analysis.pdf") as pdf:
        assert "Notebook" not in pdf.pages[0].extract_text()


def test_it_refuses_a_record_whose_work_order_has_no_sample_on_the_receiving_log_row(
    tmp_path,
):
    given = chain(tmp_path)
    record = json.loads(given["record"].read_text())
    record["work_order"] = "WO-2026-0999"
    given["record"] = tmp_path / "record.json"
    given["record"].write_text(json.dumps(record))

    result, outputs = render(tmp_path, given)

    assert result.returncode == 1
    assert "no Sample of Work Order WO-2026-0999" in result.stderr
    assert list(outputs.iterdir()) == []
