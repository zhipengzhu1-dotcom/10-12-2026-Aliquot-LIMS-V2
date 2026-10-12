import re

import pytest
from openpyxl import load_workbook

COLUMNS = [
    "#",
    "Code",
    "Name",
    "Status",
    "Lot#",
    "Weight of 20 Capsules (g)",
    "Weight Taken (g)",
    "Average Weight (g)",
]


def table(path):
    ws = load_workbook(path).active
    header = next(r for r in ws.iter_rows() if r[0].value == "#")
    rows = []
    for row in ws.iter_rows(min_row=header[0].row + 1):
        if not isinstance(row[0].value, int):
            break
        rows.append(row)
    return ws, [c.value for c in header], rows


def test_the_table_has_the_documented_columns(generated):
    _, header, _ = table(generated / "result_worksheet.xlsx")
    assert header == COLUMNS


def test_average_weight_is_a_formula_over_column_f(generated):
    _, _, rows = table(generated / "result_worksheet.xlsx")
    assert rows
    for row in rows:
        r = row[0].row
        assert row[7].value == f"=F{r}/20"


def test_recomputing_the_formulas_from_the_weighings(generated, manifest):
    ws, _, rows = table(generated / "result_worksheet.xlsx")
    weighings = manifest["weighings"]
    assert len(rows) == len(weighings) == len(manifest["containers"])
    averages = []
    for row, weighing in zip(rows, weighings, strict=True):
        assert row[0].value == weighing["container"]
        assert row[1].value == manifest["product_code"]
        assert row[2].value == manifest["product_name"]
        assert row[4].value == manifest["lot"]
        assert row[5].value == weighing["weight_of_20_capsules_g"]
        assert row[6].value == weighing["weight_taken_g"]
        cell, divisor = re.fullmatch(r"=([A-Z]+\d+)/(\d+)", row[7].value).groups()
        average = ws[cell].value / int(divisor)
        assert average == pytest.approx(weighing["weight_of_20_capsules_g"] / 20)
        assert 0.3 < average < 1.0  # a plausible capsule fill in grams
        averages.append(average)

    mean_row = rows[-1][0].row + 1
    first, last = rows[0][0].row, rows[-1][0].row
    assert ws.cell(mean_row, 7).value == "Mean"
    assert ws.cell(mean_row, 8).value == f"=AVERAGE(H{first}:H{last})"
