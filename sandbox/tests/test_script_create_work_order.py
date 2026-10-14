"""create_work_order on the Receiving Log row and database that log_receipt
wrote for a Fixture's expected record, run the way the Sandbox runs it but
without Docker."""

import json
import shutil
import sqlite3
import subprocess
import sys
from contextlib import closing
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / "sandbox" / "scripts"
RECORD = ROOT / "fixtures" / "submission-form-seed1.expected.json"

RECORD_ID = "a" * 64
ROW_ID = "c" * 64
DATABASE_ID = "e" * 64
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TEST = "Average capsule weight"
NITROSAMINES = "Nitrosamine impurities"
COLUMNS = (
    "Lab Reference",
    "Test",
    "Lot",
    "Product Code",
    "TAT Class",
    "Due Date",
    "Analyte",
    "Limit",
)


def run(script: str, inputs: Path, outputs: Path, manifest: list[dict]):
    (inputs / "manifest.json").write_text(json.dumps(manifest))
    outputs.mkdir(parents=True)
    return subprocess.run(
        [
            sys.executable,
            str(SCRIPTS / script / "main.py"),
            "--inputs",
            str(inputs),
            "--outputs",
            str(outputs),
        ],
        capture_output=True,
        text=True,
        check=False,
    )


def put(inputs: Path, artifact_id: str, source: Path, filename: str) -> None:
    (inputs / artifact_id).mkdir(parents=True)
    shutil.copyfile(source, inputs / artifact_id / filename)
    (inputs / artifact_id / filename).chmod(0o444)


def receipt(
    tmp_path: Path,
    tat_class: str,
    carried: Path | None = None,
    more: tuple[dict, ...] = (),
) -> Path:
    """Stage 2 on seed 1's expected record with `tat_class`, and `more` Samples
    after its own, each a change to seed 1's Sample: its outputs directory."""
    record = json.loads(RECORD.read_text())
    [sample] = record["samples"]
    sample["tat_class"] = tat_class
    record["samples"] += [
        {**sample, "no": str(n), **change} for n, change in enumerate(more, start=2)
    ]
    tmp_path.mkdir(exist_ok=True)
    (tmp_path / "submission.json").write_text(json.dumps(record))
    inputs = tmp_path / "receipt-inputs"
    put(inputs, RECORD_ID, tmp_path / "submission.json", "submission.json")
    manifest = [
        {
            "artifact_id": RECORD_ID,
            "filename": "submission.json",
            "artifact_type": "SubmissionRecord",
            "media_type": "application/json",
        }
    ]
    if carried is not None:
        put(inputs, DATABASE_ID, carried, "lims.sqlite")
        manifest.append(
            {
                "artifact_id": DATABASE_ID,
                "filename": "lims.sqlite",
                "artifact_type": "LimsDatabase",
                "media_type": "application/vnd.sqlite3",
            }
        )
    outputs = tmp_path / "receipt-outputs"
    result = run("log_receipt", inputs, outputs, manifest)
    assert result.returncode == 0, result.stderr
    return outputs


def work_order(tmp_path: Path, receipt_outputs: Path):
    """Stage 3 on the outputs of `receipt` and the record it logged."""
    inputs, outputs = tmp_path / "inputs", tmp_path / "outputs"
    record = receipt_outputs.parent / "submission.json"
    put(inputs, ROW_ID, receipt_outputs / "receiving_log_row.xlsx", "row.xlsx")
    put(inputs, RECORD_ID, record, "submission.json")
    put(inputs, DATABASE_ID, receipt_outputs / "lims.sqlite", "lims.sqlite")
    result = run(
        "create_work_order",
        inputs,
        outputs,
        [
            {
                "artifact_id": ROW_ID,
                "filename": "row.xlsx",
                "artifact_type": "ReceivingLogRow",
                "media_type": XLSX,
            },
            {
                "artifact_id": RECORD_ID,
                "filename": "submission.json",
                "artifact_type": "SubmissionRecord",
                "media_type": "application/json",
            },
            {
                "artifact_id": DATABASE_ID,
                "filename": "lims.sqlite",
                "artifact_type": "LimsDatabase",
                "media_type": "application/vnd.sqlite3",
            },
        ],
    )
    return result, outputs


def query(database: Path, sql: str) -> list[tuple]:
    with closing(sqlite3.connect(database)) as db:
        return db.execute(sql).fetchall()


# Seed 1 is received on Tuesday 2026-01-06.
@pytest.mark.parametrize(
    ("tat_class", "due"),
    [
        ("Rush1to2Days", "2026-01-08"),
        ("Rush3to4Days", "2026-01-12"),
        ("Regular", "2026-01-20"),
    ],
)
def test_the_due_date_counts_the_tat_class_business_days_after_receipt(
    tmp_path, tat_class, due
):
    result, outputs = work_order(tmp_path, receipt(tmp_path, tat_class))

    assert result.returncode == 0, result.stderr
    summary = json.loads((outputs / "work_order.json").read_text())
    assert {(row[4], row[5]) for row in summary["rows"]} == {(tat_class, due)}
    assert summary["header"]["Due Date"] == due
    assert query(
        outputs / "lims.sqlite", "SELECT DISTINCT due_date FROM work_orders"
    ) == [(due,)]


def test_it_writes_the_header_block_one_test_per_sample_and_the_work_orders_rows(
    tmp_path,
):
    result, outputs = work_order(tmp_path, receipt(tmp_path, "Rush3to4Days"))

    assert result.returncode == 0, result.stderr
    assert sorted(p.name for p in outputs.iterdir()) == [
        "lims.sqlite",
        "manifest.json",
        "work_order.json",
        "work_order.xlsx",
    ]
    assert json.loads((outputs / "manifest.json").read_text()) == {
        "work_order.xlsx": {"artifact_type": "WorkOrder", "media_type": XLSX},
        "lims.sqlite": {
            "artifact_type": "LimsDatabase",
            "media_type": "application/vnd.sqlite3",
        },
        "work_order.json": {"artifact_type": "Other", "media_type": "application/json"},
    }

    ws = load_workbook(outputs / "work_order.xlsx").active
    assert ws.title == "WO-2026-0001"
    due = datetime(2026, 1, 12)
    blank = (None,) * 6
    assert list(ws.iter_rows(values_only=True)) == [
        ("Work Order", "WO-2026-0001", *blank),
        ("Client", "Tallowby Botanicals LLC", *blank),
        ("Received", datetime(2026, 1, 6), *blank),
        ("Due Date", due, *blank),
        (None,) * 8,
        COLUMNS,
        ("LR-2026-0001", TEST, "LT-268123", "TLB-1475", "Rush3to4Days", due)
        + (None, None),
        ("LR-2026-0002", TEST, "LT-268123", "TLB-1475", "Rush3to4Days", due)
        + (None, None),
    ]
    row = [TEST, "LT-268123", "TLB-1475", "Rush3to4Days", "2026-01-12", "", ""]
    assert json.loads((outputs / "work_order.json").read_text()) == {
        "schema": "work_order/2",
        "receiving_log_row": [
            ["LR-2026-0001", "2026-01-06"],
            ["LR-2026-0002", "2026-01-06"],
        ],
        "header": {
            "Work Order": "WO-2026-0001",
            "Client": "Tallowby Botanicals LLC",
            "Received": "2026-01-06",
            "Due Date": "2026-01-12",
        },
        "columns": list(COLUMNS),
        "rows": [["LR-2026-0001", *row], ["LR-2026-0002", *row]],
    }
    assert query(
        outputs / "lims.sqlite",
        "SELECT * FROM work_orders ORDER BY lab_reference",
    ) == [
        ("WO-2026-0001", ref, TEST, "2026-01-12", ROW_ID)
        for ref in ("LR-2026-0001", "LR-2026-0002")
    ]


def test_one_work_order_plans_each_sample_with_its_own_lot_and_due_date(tmp_path):
    other = {
        "lot": "LT-269999",
        "product_code": "TLB-2000",
        "tat_class": "Rush1to2Days",
        "no_of_containers": "1",
    }

    result, outputs = work_order(tmp_path, receipt(tmp_path, "Regular", more=(other,)))

    assert result.returncode == 0, result.stderr
    summary = json.loads((outputs / "work_order.json").read_text())
    assert [row[:6] for row in summary["rows"]] == [
        ["LR-2026-0001", TEST, "LT-268123", "TLB-1475", "Regular", "2026-01-20"],
        ["LR-2026-0002", TEST, "LT-268123", "TLB-1475", "Regular", "2026-01-20"],
        ["LR-2026-0003", TEST, "LT-269999", "TLB-2000", "Rush1to2Days", "2026-01-08"],
    ]
    assert summary["header"]["Due Date"] == "2026-01-20", "due with its last Sample"
    assert query(
        outputs / "lims.sqlite",
        "SELECT work_order, lab_reference, due_date FROM work_orders"
        " ORDER BY lab_reference",
    ) == [
        ("WO-2026-0001", "LR-2026-0001", "2026-01-20"),
        ("WO-2026-0001", "LR-2026-0002", "2026-01-20"),
        ("WO-2026-0001", "LR-2026-0003", "2026-01-08"),
    ]


def requested(tmp_path: Path, *lines: dict):
    """Stage 3 on seed 1's record with `lines` after its own, one container each."""
    more = tuple(
        {"lot": f"LT-27000{n}", "no_of_containers": "1", **line}
        for n, line in enumerate(lines)
    )
    return work_order(tmp_path, receipt(tmp_path, "Regular", more=more))


def planned_for(outputs: Path, ref: str) -> list[tuple[str, str, str]]:
    summary = json.loads((outputs / "work_order.json").read_text())
    return [(test, a, limit) for r, test, *_, a, limit in summary["rows"] if r == ref]


def test_an_rm_sample_requesting_ndma_and_ndea_gets_both_analytes_and_limits(
    tmp_path,
):
    raw = {
        "sample_type": "RM",
        "tests_requested": "Nitrosamines: NDMA, NDEA",
        "limits": "NDMA: 96 ng/day; NDEA NMT 0.03 ppm",
    }

    result, outputs = requested(tmp_path, raw)

    assert result.returncode == 0, result.stderr
    assert planned_for(outputs, "LR-2026-0003") == [
        (NITROSAMINES, "NDMA", "96 ng/day"),
        (NITROSAMINES, "NDEA", "0.03 ppm"),
    ]
    assert query(
        outputs / "lims.sqlite",
        "SELECT test FROM work_orders WHERE lab_reference = 'LR-2026-0003'",
    ) == [(NITROSAMINES,)], "one work_orders row per Sample and Test"


def test_an_fp_sample_requesting_nitrosamines_and_capsule_weight_gets_both(
    tmp_path,
):
    fp = {"tests_requested": "Average capsule weight\nN-nitroso-varenicline"}

    result, outputs = requested(tmp_path, fp)

    assert result.returncode == 0, result.stderr
    assert planned_for(outputs, "LR-2026-0003") == [
        (TEST, "", ""),
        (NITROSAMINES, "N-nitroso-varenicline", ""),
    ]


def test_a_stability_condition_stays_on_the_sample_and_an_odd_limit_is_flagged(
    tmp_path,
):
    stability = {
        "sample_type": "Initial-1",
        "tests_requested": "NDMA",
        "limits": "below LOQ",
    }

    result, outputs = requested(tmp_path, stability)

    assert result.returncode == 0, result.stderr
    assert planned_for(outputs, "LR-2026-0003") == [
        (NITROSAMINES, "NDMA", "unread: below LOQ")
    ]
    assert query(
        outputs / "lims.sqlite",
        "SELECT sample_type FROM samples WHERE lab_reference = 'LR-2026-0003'",
    ) == [("Initial-1",)]


@pytest.mark.parametrize(
    ("line", "refusal"),
    [
        (
            {"tests_requested": "Water content"},
            "LR-2026-0003: no Test definition for the request 'Water content'",
        ),
        (
            {"tests_requested": "Nitrosamine impurities"},
            "LR-2026-0003: the request 'Nitrosamine impurities' names no"
            " nitrosamine; which should be tested?",
        ),
        (
            {"sample_type": "RM"},
            "LR-2026-0003: the line requests no Tests, and only an FP Sample has"
            " a default; Sample type 'RM'",
        ),
    ],
)
def test_it_refuses_a_request_it_cannot_map_with_its_words(tmp_path, line, refusal):
    result, outputs = requested(tmp_path, line)

    assert result.returncode == 1
    assert refusal in result.stderr
    assert list(outputs.iterdir()) == []


def test_numbering_continues_from_the_work_orders_of_the_carried_database(
    tmp_path,
):
    first, earlier = work_order(
        tmp_path / "first", receipt(tmp_path / "first", "Regular")
    )
    assert first.returncode == 0, first.stderr

    second = tmp_path / "second"
    result, outputs = work_order(
        second, receipt(second, "Regular", carried=earlier / "lims.sqlite")
    )

    assert result.returncode == 0, result.stderr
    assert query(
        outputs / "lims.sqlite",
        "SELECT work_order, lab_reference FROM work_orders ORDER BY lab_reference",
    ) == [
        ("WO-2026-0001", "LR-2026-0001"),
        ("WO-2026-0001", "LR-2026-0002"),
        ("WO-2026-0002", "LR-2026-0003"),
        ("WO-2026-0002", "LR-2026-0004"),
    ]


def test_it_refuses_a_receiving_log_row_without_a_database(tmp_path):
    receipt_outputs = receipt(tmp_path, "Regular")
    inputs, outputs = tmp_path / "inputs", tmp_path / "outputs"
    put(inputs, ROW_ID, receipt_outputs / "receiving_log_row.xlsx", "row.xlsx")

    result = run(
        "create_work_order",
        inputs,
        outputs,
        [
            {
                "artifact_id": ROW_ID,
                "filename": "row.xlsx",
                "artifact_type": "ReceivingLogRow",
                "media_type": XLSX,
            }
        ],
    )

    assert result.returncode == 1
    assert "exactly one LimsDatabase input, got 0" in result.stderr
    assert list(outputs.iterdir()) == []
