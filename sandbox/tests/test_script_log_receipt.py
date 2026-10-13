"""log_receipt on the stage 1 record of a Fixture, run the way the Sandbox runs
it but without Docker, with no database and with a carried one."""

import importlib.util
import json
import shutil
import sqlite3
import subprocess
import sys
from contextlib import closing
from datetime import datetime, time
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SCRIPT_DIR = ROOT / "sandbox" / "scripts" / "log_receipt"
RECORD = ROOT / "fixtures" / "submission-form-seed1.expected.json"

RECORD_ID = "a" * 64
DATABASE_ID = "b" * 64
EARLIER_RECORD_ID = "d" * 64
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

COLUMNS = [
    "#",
    "Rec. DATE",
    "Rec. Time",
    "Product Code",
    "PRODUCT NAME",
    "TYPE",
    "LOT#",
    "CLIENT",
    "EXP DATE",
    "QTY RECD",
]

_spec = importlib.util.spec_from_file_location(
    "lims_schema", SCRIPT_DIR / "lims_schema.py"
)
assert _spec and _spec.loader
lims_schema = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(lims_schema)


def stage(inputs: Path, database: Path | None = None) -> None:
    (inputs / RECORD_ID).mkdir(parents=True)
    shutil.copyfile(RECORD, inputs / RECORD_ID / "submission.json")
    manifest = [
        {
            "artifact_id": RECORD_ID,
            "filename": "submission.json",
            "artifact_type": "SubmissionRecord",
            "media_type": "application/json",
        }
    ]
    if database is not None:
        (inputs / DATABASE_ID).mkdir()
        shutil.copyfile(database, inputs / DATABASE_ID / "lims.sqlite")
        (inputs / DATABASE_ID / "lims.sqlite").chmod(0o444)
        manifest.append(
            {
                "artifact_id": DATABASE_ID,
                "filename": "lims.sqlite",
                "artifact_type": "LimsDatabase",
                "media_type": "application/vnd.sqlite3",
            }
        )
    (inputs / "manifest.json").write_text(json.dumps(manifest))


def run_script(tmp_path: Path, database: Path | None = None):
    inputs, outputs = tmp_path / "inputs", tmp_path / "outputs"
    outputs.mkdir(parents=True)
    stage(inputs, database)
    result = subprocess.run(
        [
            sys.executable,
            str(SCRIPT_DIR / "main.py"),
            "--inputs",
            str(inputs),
            "--outputs",
            str(outputs),
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    return result, outputs


def query(database: Path, sql: str) -> list[tuple]:
    with closing(sqlite3.connect(database)) as db:
        return db.execute(sql).fetchall()


def sheet_rows(outputs: Path) -> tuple[str, list[tuple]]:
    ws = load_workbook(outputs / "receiving_log_row.xlsx").active
    return ws.title, list(ws.iter_rows(values_only=True))


def database_with_a_2026_sample(path: Path) -> Path:
    """A database holding LR-2026-0007 and LR-2025-0050 from an earlier record."""
    party = {
        f"{p}_{f}": "x"
        for p in ("submitted_by", "bill_to", "report_to")
        for f in ("company", "contact", "address", "city", "phone", "email")
    }
    with closing(sqlite3.connect(path)) as db:
        lims_schema.create_schema(db)
        db.execute(
            "INSERT INTO lots VALUES ('LT-250001', 'CVN-1000', 'Earlier Capsules',"
            " '30 ct HDPE bottle', 'Corvane Nutrition Co.', NULL, ?)",
            (EARLIER_RECORD_ID,),
        )
        for ref, on in (("LR-2026-0007", "2026-01-02"), ("LR-2025-0050", "2025-12-30")):
            row = {
                "lab_reference": ref,
                "lot": "LT-250001",
                "container": 1,
                "qty_received": 1,
                "sample_type": "FP",
                "tat_code": "1",
                "tat_class": "Regular",
                "date_submitted": on,
                "date_received": on,
                "time_received": "09:00",
                "received_by": "LG",
                **party,
                "source_artifact": EARLIER_RECORD_ID,
            }
            db.execute(
                f"INSERT INTO samples ({', '.join(row)})"
                f" VALUES ({', '.join('?' for _ in row)})",
                tuple(row.values()),
            )
        db.commit()
    return path


def test_with_no_database_it_creates_the_five_tables_and_numbers_from_1(tmp_path):
    result, outputs = run_script(tmp_path)

    assert result.returncode == 0, result.stderr
    assert sorted(p.name for p in outputs.iterdir()) == [
        "lims.sqlite",
        "manifest.json",
        "receipt.json",
        "receiving_log_row.xlsx",
    ]
    assert json.loads((outputs / "manifest.json").read_text()) == {
        "receiving_log_row.xlsx": {
            "artifact_type": "ReceivingLogRow",
            "media_type": XLSX,
        },
        "lims.sqlite": {
            "artifact_type": "LimsDatabase",
            "media_type": "application/vnd.sqlite3",
        },
        "receipt.json": {"artifact_type": "Other", "media_type": "application/json"},
    }

    title, rows = sheet_rows(outputs)
    assert title == "01-06"
    assert list(rows[0]) == COLUMNS
    assert rows[1:] == [
        (
            ref,
            datetime(2026, 1, 6),
            time(15, 17),
            "TLB-1475",
            "Tallowby Milk Thistle 250 mg Capsules",
            "60 ct amber glass bottle",
            "LT-268123",
            "Tallowby Botanicals LLC",
            None,
            1,
        )
        for ref in ("LR-2026-0001", "LR-2026-0002")
    ]
    assert json.loads((outputs / "receipt.json").read_text()) == {
        "schema": "receipt/1",
        "sheet": "01-06",
        "columns": COLUMNS,
        "rows": [
            [
                ref,
                "2026-01-06",
                "15:17",
                "TLB-1475",
                "Tallowby Milk Thistle 250 mg Capsules",
                "60 ct amber glass bottle",
                "LT-268123",
                "Tallowby Botanicals LLC",
                "",
                "1",
            ]
            for ref in ("LR-2026-0001", "LR-2026-0002")
        ],
    }

    database = outputs / "lims.sqlite"
    assert query(
        database, "SELECT name FROM sqlite_schema WHERE type = 'table' ORDER BY name"
    ) == [("lots",), ("reports",), ("results",), ("samples",), ("work_orders",)]
    assert query(database, "PRAGMA user_version") == [(lims_schema.SCHEMA_VERSION,)]
    assert query(database, "SELECT * FROM lots") == [
        (
            "LT-268123",
            "TLB-1475",
            "Tallowby Milk Thistle 250 mg Capsules",
            "60 ct amber glass bottle",
            "Tallowby Botanicals LLC",
            None,
            RECORD_ID,
        )
    ]
    assert query(
        database,
        "SELECT lab_reference, lot, container, date_received, time_received,"
        " received_by, tat_class, submitted_by_company, report_to_email,"
        " source_artifact FROM samples ORDER BY lab_reference",
    ) == [
        (
            ref,
            "LT-268123",
            container,
            "2026-01-06",
            "15:17",
            "LG",
            "Rush3to4Days",
            "Tallowby Botanicals LLC",
            "maren.vantongeren@tallowby.example",
            RECORD_ID,
        )
        for container, ref in enumerate(("LR-2026-0001", "LR-2026-0002"), start=1)
    ]


def test_numbering_continues_from_the_carried_database_for_the_year_received(
    tmp_path,
):
    carried = database_with_a_2026_sample(tmp_path / "carried.sqlite")
    before = carried.read_bytes()

    result, outputs = run_script(tmp_path, carried)

    assert result.returncode == 0, result.stderr
    _, rows = sheet_rows(outputs)
    assert [row[0] for row in rows[1:]] == ["LR-2026-0008", "LR-2026-0009"]
    assert query(
        outputs / "lims.sqlite",
        "SELECT lab_reference, lot, source_artifact FROM samples"
        " ORDER BY lab_reference",
    ) == [
        ("LR-2025-0050", "LT-250001", EARLIER_RECORD_ID),
        ("LR-2026-0007", "LT-250001", EARLIER_RECORD_ID),
        ("LR-2026-0008", "LT-268123", RECORD_ID),
        ("LR-2026-0009", "LT-268123", RECORD_ID),
    ]
    assert query(outputs / "lims.sqlite", "SELECT lot FROM lots ORDER BY lot") == [
        ("LT-250001",),
        ("LT-268123",),
    ]
    assert carried.read_bytes() == before, "the carried version is never changed"


def test_a_second_receipt_of_the_same_lot_keeps_its_one_lots_row(tmp_path):
    first, outputs = run_script(tmp_path / "first")
    assert first.returncode == 0, first.stderr

    second, again = run_script(tmp_path / "second", outputs / "lims.sqlite")

    assert second.returncode == 0, second.stderr
    assert query(again / "lims.sqlite", "SELECT lab_reference FROM samples") == [
        ("LR-2026-0001",),
        ("LR-2026-0002",),
        ("LR-2026-0003",),
        ("LR-2026-0004",),
    ]
    assert query(again / "lims.sqlite", "SELECT count(*) FROM lots") == [(1,)]


def test_it_refuses_inputs_without_a_submission_record(tmp_path):
    inputs, outputs = tmp_path / "inputs", tmp_path / "outputs"
    inputs.mkdir()
    outputs.mkdir()
    (inputs / "manifest.json").write_text("[]")

    result = subprocess.run(
        [
            sys.executable,
            str(SCRIPT_DIR / "main.py"),
            "--inputs",
            str(inputs),
            "--outputs",
            str(outputs),
        ],
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 1
    assert "exactly one SubmissionRecord" in result.stderr
    assert list(outputs.iterdir()) == []
