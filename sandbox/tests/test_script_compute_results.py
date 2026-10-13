"""compute_results on the Work Order and database that create_work_order wrote
for seed 1's expected record, with seed 1's committed weighings, run the way
the Sandbox runs it but without Docker."""

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_script_create_work_order import (
    DATABASE_ID,
    ROOT,
    XLSX,
    put,
    query,
    receipt,
    run,
    work_order,
)

WEIGHINGS = ROOT / "fixtures" / "submission-form-seed1.weighings.json"
WORK_ORDER_ID = "f" * 64
WEIGHINGS_ID = "b" * 64
TEST = "Average capsule weight"
MEAN_TEST = "Average capsule weight, Work Order mean"


def results(tmp_path: Path, weighings: Path = WEIGHINGS, with_weighings=True):
    """Stages 2, 3 and 4 on seed 1's expected record: stage 4's run and outputs."""
    made, stage3 = work_order(tmp_path / "stage3", receipt(tmp_path, "Rush3to4Days"))
    assert made.returncode == 0, made.stderr
    inputs, outputs = tmp_path / "inputs", tmp_path / "outputs"
    put(inputs, WORK_ORDER_ID, stage3 / "work_order.xlsx", "work_order.xlsx")
    put(inputs, DATABASE_ID, stage3 / "lims.sqlite", "lims.sqlite")
    manifest = [
        {
            "artifact_id": WORK_ORDER_ID,
            "filename": "work_order.xlsx",
            "artifact_type": "WorkOrder",
            "media_type": XLSX,
        },
        {
            "artifact_id": DATABASE_ID,
            "filename": "lims.sqlite",
            "artifact_type": "LimsDatabase",
            "media_type": "application/vnd.sqlite3",
        },
    ]
    if with_weighings:
        put(inputs, WEIGHINGS_ID, weighings, "weighings.json")
        manifest.append(
            {
                "artifact_id": WEIGHINGS_ID,
                "filename": "weighings.json",
                "artifact_type": "Other",
                "media_type": "application/json",
            }
        )
    return run("compute_results", inputs, outputs, manifest), outputs


def test_it_writes_the_worksheet_with_formulas_and_the_weighings_as_values(tmp_path):
    result, outputs = results(tmp_path)

    assert result.returncode == 0, result.stderr
    assert sorted(p.name for p in outputs.iterdir()) == [
        "lims.sqlite",
        "manifest.json",
        "result_worksheet.json",
        "result_worksheet.xlsx",
    ]
    assert json.loads((outputs / "manifest.json").read_text()) == {
        "result_worksheet.xlsx": {
            "artifact_type": "ResultWorksheet",
            "media_type": XLSX,
        },
        "lims.sqlite": {
            "artifact_type": "LimsDatabase",
            "media_type": "application/vnd.sqlite3",
        },
        "result_worksheet.json": {
            "artifact_type": "Other",
            "media_type": "application/json",
        },
    }
    name = "Tallowby Milk Thistle 250 mg Capsules"
    sheet = [
        ["Result Worksheet: Average Capsule Weight", None, *[None] * 6],
        ["Work Order", "WO-2026-0001", *[None] * 6],
        ["Test", TEST, *[None] * 6],
        [None] * 8,
        [
            "#",
            "Code",
            "Name",
            "Status",
            "Lot#",
            "Weight of 20 Capsules (g)",
            "Weight Taken (g)",
            "Average Weight (g)",
        ],
        [1, "TLB-1475", name, "Complete", "LT-268123", 7.0886, 0.3541, "=F6/20"],
        [2, "TLB-1475", name, "Complete", "LT-268123", 7.2103, 0.3422, "=F7/20"],
        [None] * 6 + ["Mean", "=AVERAGE(H6:H7)"],
    ]
    ws = load_workbook(outputs / "result_worksheet.xlsx").active
    assert [list(row) for row in ws.iter_rows(values_only=True)] == sheet
    assert json.loads((outputs / "result_worksheet.json").read_text()) == {
        "schema": "result_worksheet/1",
        "sheet": sheet,
    }


def test_the_results_rows_hold_each_containers_average_and_the_mean(tmp_path):
    result, outputs = results(tmp_path)

    assert result.returncode == 0, result.stderr
    rows = query(
        outputs / "lims.sqlite",
        "SELECT lab_reference, test, value, unit, source_artifact FROM results"
        " ORDER BY lab_reference, test",
    )
    mean = (7.0886 / 20 + 7.2103 / 20) / 2
    assert rows == [
        ("LR-2026-0001", TEST, pytest.approx(0.35443, abs=1e-9), "g", WEIGHINGS_ID),
        ("LR-2026-0001", MEAN_TEST, pytest.approx(mean, abs=1e-9), "g", WEIGHINGS_ID),
        ("LR-2026-0002", TEST, pytest.approx(0.360515, abs=1e-9), "g", WEIGHINGS_ID),
        ("LR-2026-0002", MEAN_TEST, pytest.approx(mean, abs=1e-9), "g", WEIGHINGS_ID),
    ]
    assert query(outputs / "lims.sqlite", "SELECT count(*) FROM work_orders") == [
        (2,)
    ], "the carried rows are kept"


def test_it_refuses_weighings_that_miss_a_container_of_the_work_order(tmp_path):
    short = tmp_path / "short.json"
    short.write_text(json.dumps(json.loads(WEIGHINGS.read_text())[:1]))

    result, outputs = results(tmp_path, weighings=short)

    assert result.returncode == 1
    assert "LR-2026-0002: no weighing for container 2" in result.stderr
    assert list(outputs.iterdir()) == []


def test_it_refuses_a_work_order_without_weighings(tmp_path):
    result, outputs = results(tmp_path, with_weighings=False)

    assert result.returncode == 1
    assert "exactly one Other input, got 0" in result.stderr
    assert list(outputs.iterdir()) == []


def noted(tmp_path: Path, *references: str) -> Path:
    """Seed 1's weighings, each naming the next of `references`, as
    parse_results_message writes them."""
    entries = json.loads(WEIGHINGS.read_text())
    path = tmp_path / "noted.json"
    path.write_text(
        json.dumps(
            [
                {**entry, "notebook_reference": ref}
                for entry, ref in zip(entries, references, strict=True)
            ]
        )
    )
    return path


def test_it_writes_the_weighings_notebook_reference_in_the_header(tmp_path):
    result, outputs = results(tmp_path, weighings=noted(tmp_path, *["NB-07 p.31"] * 2))

    assert result.returncode == 0, result.stderr
    ws = load_workbook(outputs / "result_worksheet.xlsx").active
    header = [list(row[:2]) for row in ws.iter_rows(max_row=5, values_only=True)]
    assert header == [
        ["Result Worksheet: Average Capsule Weight", None],
        ["Work Order", "WO-2026-0001"],
        ["Test", TEST],
        ["Notebook Reference", "NB-07 p.31"],
        ["#", "Code"],
    ]
    read_back = json.loads((outputs / "result_worksheet.json").read_text())["sheet"]
    assert read_back[3][:2] == ["Notebook Reference", "NB-07 p.31"]


def test_it_refuses_weighings_that_name_two_notebook_references(tmp_path):
    result, outputs = results(
        tmp_path, weighings=noted(tmp_path, "NB-07 p.31", "NB-07 p.32")
    )

    assert result.returncode == 1
    assert "do not all name one notebook reference" in result.stderr
    assert list(outputs.iterdir()) == []
