from datetime import date, datetime, time

from openpyxl import load_workbook

COLUMNS = [
    "#",
    "Rec. DATE",
    "Rec. Time",
    "Product Code",
    "PRODUCT NAME",
    "TYPE",
    "LOT#",
    "CLIENT",
    "EXP DATE",
    "QTY RECD",
]


def test_every_sheet_has_exactly_the_documented_columns(generated):
    wb = load_workbook(generated / "receiving_log.xlsx")
    for ws in wb.worksheets:
        assert [c.value for c in ws[1]] == COLUMNS, ws.title


def test_one_row_per_container_on_the_yearly_and_the_daily_sheet(generated, manifest):
    wb = load_workbook(generated / "receiving_log.xlsx")
    received = date.fromisoformat(manifest["received_on"])
    assert wb.sheetnames == [str(received.year), f"{received:%m-%d}"]
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(rows) == len(manifest["containers"]) >= 1
        assert [r[0] for r in rows] == [c["number"] for c in manifest["containers"]]


def test_each_row_carries_the_receipt(generated, manifest):
    ws = load_workbook(generated / "receiving_log.xlsx").worksheets[0]
    hh, mm = map(int, manifest["received_at"].split(":"))
    for row in ws.iter_rows(min_row=2, values_only=True):
        assert row[1:9] == (
            datetime.fromisoformat(manifest["received_on"]),
            time(hh, mm),
            manifest["product_code"],
            manifest["product_name"],
            manifest["product_type"],
            manifest["lot"],
            manifest["client"],
            datetime.fromisoformat(manifest["expires_on"]),
        )
        assert row[9] == 1
