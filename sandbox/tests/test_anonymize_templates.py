"""The scrubbed copies in fixtures/real-layout against the real templates.

Runs only where the real templates are: ALIQUOT_TEMPLATES names their folder,
outside the repository (ADR-0009). Elsewhere, CI included, every test here is
skipped with the reason. A failure names a part and a length, never a value.
"""

import html
import os
import re
import struct
import zipfile
import zlib
from pathlib import Path

import docx
import openpyxl
import pytest

from synth.anonymize.cli import KINDS, kind_of, main
from synth.anonymize.labels import LABELS, is_field_code, is_generic_sheet_name, is_kept

REPO = Path(__file__).resolve().parents[2]
FIXTURES = REPO / "fixtures"
COMMITTED = FIXTURES / "real-layout"
TEMPLATES = Path(os.environ.get("ALIQUOT_TEMPLATES", "/nonexistent"))

pytestmark = pytest.mark.skipif(
    not TEMPLATES.is_dir(),
    reason="ALIQUOT_TEMPLATES does not name a folder: the real templates live only "
    "on the lab's machine (ADR-0009), so the leak and layout checks run there",
)

# Elements whose text is layout or the application's own bookkeeping, kept as written.
STRUCTURE = re.compile(
    r"wp:posOffset|wp:align|wp14:pct\w+|xdr:(?:col|colOff|row|rowOff)|Application"
    r"|AppVersion|DocSecurity|ScaleCrop|LinksUpToDate|SharedDoc|HyperlinksChanged"
    r"|Template|TotalTime|Pages|Words|Characters|CharactersWithSpaces|Lines|Paragraphs"
    r"|vt:i4|cp:revision|Version|MinVersion|Culture"
)
# Parts that hold only styles, numbering, fonts, theme and calculation order.
STRUCTURE_PARTS = re.compile(
    r"\[Content_Types\]\.xml|.*theme\d+\.xml|.*/(?:styles|numbering|fontTable|webSettings"
    r"|calcChain|metadata)\.xml"
)
# What docProps/app.xml calls the parts of a workbook or a document.
OFFICE_WORDS = {"worksheets", "named ranges", "title"}
FORMULAS = ("f", "definedName", "formula", "formula1", "formula2")
EXCEL_ERROR = re.compile(r"#(?:REF!|N/A|VALUE!|DIV/0!|NAME\?|NUM!|NULL!)")
FORMULA_LITERAL = re.compile(
    r'"(?:[^"]|"")*"|(?<![\w$.\]\'])\d+\.\d+|(?<![\w$.:\]\'])\d{3,}(?![\w:\'])'
)
VALUE_ATTRIBUTES = re.compile(
    r'\s(?:url|descr|title|displayName|userId|userName|o:gfxdata)="([^"]*)"'
    r'|<w:(?:alias|tag) w:val="([^"]*)"'
    r'|<Relationship [^>]*Target="([^"]*)" TargetMode="External"'
)
NUMERIC = re.compile(r"[\d.,:/\- ]+")


def _values(path: Path) -> list[tuple[str, str]]:
    """Every value-bearing text of a document, by part: cells, runs, paragraphs,
    properties, formula literals, external paths and the attributes that hold names."""
    out: list[tuple[str, str]] = []
    with zipfile.ZipFile(path) as z:
        for part in z.namelist():
            if not part.endswith((".xml", ".rels")) or STRUCTURE_PARTS.fullmatch(part):
                continue
            xml = z.read(part).decode("utf-8", errors="replace")
            for m in re.finditer(r"<([\w:]+)(?:\s[^>]*)?>([^<]+)</\1>", xml):
                tag, text = m.group(1), html.unescape(m.group(2))
                if tag in FORMULAS:
                    out += [
                        (part, lit.strip('"')) for lit in FORMULA_LITERAL.findall(text)
                    ]
                elif not STRUCTURE.fullmatch(tag):
                    out.append((part, text))
            for m in re.finditer(
                r"<(w:p|si|a:p|is)[\s>](?:(?!<\1[\s>]).)*?</\1>", xml, re.S
            ):
                runs = re.findall(
                    r"<(?:w:|a:)?t(?:\s[^>]*)?>([^<]*)</(?:w:|a:)?t>", m.group(0)
                )
                out.append((part, html.unescape("".join(runs))))
            out += [
                (part, html.unescape(next(g for g in groups if g)))
                for groups in VALUE_ATTRIBUTES.findall(xml)
                if any(groups)
            ]
    pieces = []
    for part, value in out:
        pieces.append((part, value.strip()))
        pieces += [
            (part, t) for t in re.split(r"[\s,;]+", value) if re.search(r"\d", t)
        ]
    return pieces


def _is_structure(v: str) -> bool:
    print_area = re.fullmatch(r"'?(.+?)'?!Print_\w+", v)
    return (
        is_kept(v)
        or is_field_code(v)
        or EXCEL_ERROR.fullmatch(v) is not None
        or v.lower() in OFFICE_WORDS
        or is_generic_sheet_name(v)
        or (print_area is not None and is_generic_sheet_name(print_area.group(1)))
        or any(v.lower() in label.lower() for label in LABELS)
    )


def _candidates(template: Path) -> list[tuple[str, str]]:
    """Every value of 4+ characters in the template that is not a label, by part."""
    return [
        (p, v) for p, v in _values(template) if len(v) >= 4 and not _is_structure(v)
    ]


@pytest.fixture(scope="module")
def pairs(tmp_path_factory) -> dict[str, tuple[Path, Path]]:
    fresh = tmp_path_factory.mktemp("anonymized")
    main([str(TEMPLATES), "--out", str(fresh)])
    found = {
        kind_of(p): p
        for p in TEMPLATES.iterdir()
        if p.suffix in (".docx", ".xlsx") and not p.name.startswith("~$")
    }
    return {k: (found[k], fresh / f"{k}{found[k].suffix}") for k in KINDS}


def _texts(path: Path):
    """Every part of a file, or the file itself, as UTF-8 and as UTF-16 text."""
    blobs = []
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as z:
            blobs += [(f"{path.name}:{n}", z.read(n)) for n in z.namelist()]
    else:
        blobs.append((path.name, path.read_bytes()))
    for where, data in blobs:
        yield where, html.unescape(data.decode("utf-8", errors="replace")).lower()
        yield where, data.decode("utf-16-le", errors="replace").lower()


def leaks(template: Path, scrubbed: Path, everywhere: list[Path]) -> list[str]:
    """Where a template value survives, as '<file>:<part> holds a value of N
    characters', never the value."""
    values = {v.lower() for _, v in _candidates(template)}
    name_parts = re.split(r"[\s_.]+", template.stem) + re.split(
        r"[\s_.\-]+", template.stem
    )
    values |= {t.lower() for t in name_parts if len(t) >= 4 and not _is_structure(t)}
    numbers = {v for v in values if NUMERIC.fullmatch(v)}
    strings = values - numbers
    found = []
    for path in [scrubbed, *everywhere]:
        if zipfile.is_zipfile(path):
            held = {v.lower() for _, v in _values(path)}
            found += [
                f"{path.name} holds a number of {len(v)} characters"
                for v in numbers
                if v in held
            ]
        for where, text in _texts(path):
            found += [
                f"{where} holds a value of {len(v)} characters"
                for v in strings
                if v in text
            ]
            if not zipfile.is_zipfile(path):
                found += [
                    f"{where} holds a number of {len(v)} characters"
                    for v in numbers
                    if re.search(rf"(?<![\w.]){re.escape(v)}(?![\w])", text)
                ]
        found += [
            f"{path.name}'s name holds a value of {len(v)} characters"
            for v in values
            if v in path.name.lower()
        ]
    return sorted(set(found))


def _fixture_files() -> list[Path]:
    return sorted(p for p in FIXTURES.rglob("*") if p.is_file())


@pytest.mark.parametrize("kind", KINDS)
def test_no_template_value_survives_in_the_copy_or_any_fixture(pairs, kind):
    template, scrubbed = pairs[kind]
    assert _candidates(template), "the template yields values to look for"
    assert leaks(template, scrubbed, _fixture_files()) == []


@pytest.mark.parametrize("kind", KINDS)
def test_the_committed_copy_is_what_the_tool_writes(pairs, kind):
    _, fresh = pairs[kind]
    assert (COMMITTED / fresh.name).read_bytes() == fresh.read_bytes()


@pytest.mark.parametrize("kind", KINDS)
def test_images_thumbnails_and_printer_settings_are_blank(pairs, kind):
    _, scrubbed = pairs[kind]
    with zipfile.ZipFile(scrubbed) as z:
        for part in z.namelist():
            data = z.read(part)
            if part.endswith(".png"):
                idat = data[data.index(b"IDAT") + 4 : data.index(b"IEND") - 8]
                assert set(zlib.decompress(idat)) <= {0, 255}, part
            elif part.endswith(".emf"):
                assert struct.unpack_from("<I", data, 52)[0] == 2, part
            elif part.endswith(".bin"):
                assert data[:64] == bytes(64), part
                assert struct.unpack_from("<H", data, 70)[0] == 0, part


def _formula_shape(f: str) -> str:
    return FORMULA_LITERAL.sub("#", f)


def _kind_of_value(v) -> str:
    if v is None:
        return "empty"
    if isinstance(v, str) and v.startswith("="):
        return "formula"
    return type(v).__name__


@pytest.mark.parametrize("kind", [k for k in KINDS if k != "submission-form"])
def test_the_workbook_layout_is_the_templates(pairs, kind):
    template, scrubbed = pairs[kind]
    real = openpyxl.load_workbook(template)
    copy = openpyxl.load_workbook(scrubbed)
    assert copy.sheetnames == real.sheetnames
    for r, c in zip(real.worksheets, copy.worksheets, strict=True):
        assert c.dimensions == r.dimensions, r.title
        assert sorted(map(str, c.merged_cells.ranges)) == sorted(
            map(str, r.merged_cells.ranges)
        )
        assert {k: d.width for k, d in c.column_dimensions.items()} == {
            k: d.width for k, d in r.column_dimensions.items()
        }
        for row_r, row_c in zip(r.iter_rows(), c.iter_rows(), strict=True):
            for a, b in zip(row_r, row_c, strict=True):
                where = f"{r.title}!{a.coordinate}"
                assert _kind_of_value(b.value) == _kind_of_value(a.value), where
                assert b.style_id == a.style_id, where
                if _kind_of_value(a.value) == "formula":
                    assert _formula_shape(b.value) == _formula_shape(a.value), where
                elif isinstance(a.value, str) and is_kept(a.value):
                    assert b.value == a.value, where


def test_the_submission_form_layout_is_the_templates(pairs):
    template, scrubbed = pairs["submission-form"]
    real, copy = docx.Document(template), docx.Document(scrubbed)
    assert len(copy.paragraphs) == len(real.paragraphs)
    for a, b in zip(real.paragraphs, copy.paragraphs, strict=True):
        if is_kept(a.text):
            assert b.text == a.text
    assert len(copy.tables) == len(real.tables)
    for t_r, t_c in zip(real.tables, copy.tables, strict=True):
        assert (len(t_c.rows), len(t_c.columns)) == (len(t_r.rows), len(t_r.columns))
        for cell_r, cell_c in zip(t_r._cells, t_c._cells, strict=True):
            if is_kept(cell_r.text):
                assert cell_c.text == cell_r.text
    for s_r, s_c in zip(real.sections, copy.sections, strict=True):
        for part in ("header", "footer", "first_page_header", "first_page_footer"):
            a, b = getattr(s_r, part), getattr(s_c, part)
            assert len(b.paragraphs) == len(a.paragraphs), part
