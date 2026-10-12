from docx import Document

from synth.submission_form import (
    REQUIRED_PARTY_FIELDS,
    REQUIRED_RECEIPT_FIELDS,
    REQUIRED_SAMPLE_FIELDS,
)
from synth.tat import TAT_CODES


def tables(generated):
    doc = Document(str(generated / "submission_form.docx"))
    return [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]


def test_three_tables_header_samples_and_lab_use_only_receipt(generated):
    header, samples, receipt = tables(generated)
    assert header[0][1:] == ["Submitted By", "Bill To", "Report To"]
    assert len(samples) == 11
    assert all(len(row) == 8 for row in samples)
    assert receipt[0] == ["Date Received", "Time Received", "Received By (initials)"]
    text = "\n".join(
        p.text for p in Document(str(generated / "submission_form.docx")).paragraphs
    )
    assert "FOR LAB USE ONLY" in text


def test_the_parties_come_from_the_record(generated, manifest):
    header, _, _ = tables(generated)
    by_label = {row[0]: row[1:] for row in header[1:]}
    assert by_label["Company"] == [manifest["client"]] * 3
    assert by_label["Contact"] == [
        manifest[p]["contact"] for p in ("submitter", "bill_to", "report_to")
    ]
    assert by_label["Email"][2] == manifest["report_to"]["email"]


def test_the_first_sample_row_carries_the_record_and_a_valid_tat_code(
    generated, manifest
):
    _, samples, _ = tables(generated)
    row = dict(zip(samples[0], samples[1], strict=True))
    assert row == {
        "No.": "1",
        "Sample Description": manifest["product_name"],
        "Sample Type": "FP",
        "TAT Code": manifest["tat_code"],
        "Lot": manifest["lot"],
        "Product Code": manifest["product_code"],
        "Product Type": manifest["product_type"],
        "No. of Containers": str(len(manifest["containers"])),
    }
    assert row["TAT Code"] in TAT_CODES
    assert [r[0] for r in samples[1:]] == [str(n) for n in range(1, 11)]
    assert all(cell == "" for r in samples[2:] for cell in r[1:])


def test_the_receipt_block_is_filled_with_date_time_and_initials(generated, manifest):
    _, _, receipt = tables(generated)
    y, m, d = manifest["received_on"].split("-")
    assert receipt[1] == [
        f"{m}/{d}/{y}",
        manifest["received_at"],
        manifest["received_by"],
    ]


def test_every_required_field_is_filled(generated):
    header, samples, receipt = tables(generated)
    columns = header[0]
    for party, fields in REQUIRED_PARTY_FIELDS.items():
        col = columns.index(party)
        for field in fields:
            assert next(r for r in header if r[0] == field)[col].strip(), (party, field)
    first = dict(zip(samples[0], samples[1], strict=True))
    for field in REQUIRED_SAMPLE_FIELDS:
        assert first[field].strip(), field
    received = dict(zip(receipt[0], receipt[1], strict=True))
    for field in REQUIRED_RECEIPT_FIELDS:
        assert received[field].strip(), field
