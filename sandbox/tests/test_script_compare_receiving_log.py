"""compare_receiving_log on a Shadow and the lab's workbook with rows entered by hand.

The scrubbed Receiving Log stands in for the lab's workbook. The Shadow is
what log_receipt writes for the real-layout Fixture's three containers. The
lab's copy then gets hand-entered rows for the same Samples: one that matches,
one with a changed Lot, and none for the third container.
"""

import hashlib
import json
import shutil
import subprocess
import sys
from datetime import datetime, time
from pathlib import Path

from openpyxl import load_workbook
from test_script_log_receipt_shadow import LAB_LOG, run, seed4

ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / "sandbox" / "scripts" / "compare_receiving_log" / "main.py"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def entered_by_hand(row: int, lot: str, lab_reference: str) -> list:
    return [
        str(row - 3),
        datetime(2026, 9, 28),
        time(13, 17),
        "TLB-5077",
        "Tallowby Zinc Picolinate 30 mg Capsules",
        "FP",
        lot,
        "Tallowby Botanicals LLC",
        None,
        1,
        "Calla Halloway",
        "IS",
        lab_reference,
    ]


def labs_workbook(tmp_path: Path, rows: list[list]) -> Path:
    path = tmp_path / "lab" / "receiving-log.xlsx"
    path.parent.mkdir()
    shutil.copyfile(LAB_LOG, path)
    wb = load_workbook(path)
    ws = wb["2026"]
    for i, values in enumerate(rows):
        for col, value in enumerate(values, start=1):
            ws.cell(39 + i, col, value)
    wb.save(path)
    path.chmod(0o444)
    return path


def compare(tmp_path: Path, lab: Path) -> tuple[subprocess.CompletedProcess, Path]:
    made, shadow_outputs, _ = run(tmp_path / "stage2", seed4())
    assert made.returncode == 0, made.stderr
    inputs, outputs = tmp_path / "inputs", tmp_path / "outputs"
    outputs.mkdir()
    manifest = []
    for artifact_id, source, artifact_type in (
        ("1" * 64, shadow_outputs / "receiving_log_shadow.xlsx", "ShadowReceivingLog"),
        ("2" * 64, shadow_outputs / "shadow.json", "Other"),
        ("3" * 64, lab, "ReceivingLog"),
    ):
        (inputs / artifact_id).mkdir(parents=True)
        shutil.copyfile(source, inputs / artifact_id / source.name)
        manifest.append(
            {
                "artifact_id": artifact_id,
                "filename": source.name,
                "artifact_type": artifact_type,
                "media_type": XLSX,
            }
        )
    (inputs / "manifest.json").write_text(json.dumps(manifest))
    result = subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--inputs",
            str(inputs),
            "--outputs",
            str(outputs),
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    return result, outputs


def test_it_reports_a_changed_lot_and_a_missing_row_by_field_and_none_on_a_match(
    tmp_path,
):
    lab = labs_workbook(
        tmp_path,
        [
            entered_by_hand(39, "LT-267759", "QAL-26-0412"),
            entered_by_hand(40, "LT-267795", "QAL-26-0413"),
        ],
    )
    before = hashlib.sha256(lab.read_bytes()).hexdigest()

    result, outputs = compare(tmp_path, lab)

    assert result.returncode == 0, result.stderr
    report = json.loads((outputs / "comparison.json").read_text())
    assert report["sheet"] == "2026"
    assert report["rows"] == [
        {
            "status": "matched",
            "shadow_row": 39,
            "lab_row": 39,
            "lab_reference": "LR-2026-0001",
            "differences": [],
        },
        {
            "status": "differs",
            "shadow_row": 40,
            "lab_row": 40,
            "lab_reference": "LR-2026-0002",
            "differences": [
                {
                    "column": "G",
                    "field": "LOT#",
                    "shadow": "LT-267759",
                    "lab": "LT-267795",
                }
            ],
        },
        {
            "status": "missing_in_lab",
            "shadow_row": 41,
            "lab_row": None,
            "lab_reference": "LR-2026-0003",
            "differences": [],
        },
    ]
    assert report["counts"] == {
        "matched": 1,
        "differs": 1,
        "missing_in_lab": 1,
        "missing_in_shadow": 0,
    }
    text = (outputs / "comparison.txt").read_text()
    lines = [line.split() for line in text.splitlines()]
    assert [
        "40",
        "40",
        "LR-2026-0002",
        "differs",
        "LOT#",
        "LT-267759",
        "LT-267795",
    ] in (lines), text
    assert ["41", "-", "LR-2026-0003", "missing", "in", "the", "lab's"] in lines, text
    manifest = json.loads((outputs / "manifest.json").read_text())
    assert {k: v["artifact_type"] for k, v in manifest.items()} == {
        "comparison.json": "Comparison",
        "comparison.txt": "Comparison",
    }
    assert hashlib.sha256(lab.read_bytes()).hexdigest() == before


def test_a_lab_row_for_the_same_samples_beyond_the_shadows_is_missing_in_the_shadow(
    tmp_path,
):
    lab = labs_workbook(
        tmp_path,
        [entered_by_hand(39 + i, "LT-267759", f"QAL-26-04{i}") for i in range(4)],
    )

    result, outputs = compare(tmp_path, lab)

    assert result.returncode == 0, result.stderr
    report = json.loads((outputs / "comparison.json").read_text())
    assert [r["status"] for r in report["rows"]] == [
        "matched",
        "matched",
        "matched",
        "missing_in_shadow",
    ]
    assert report["rows"][-1]["lab_row"] == 42


def test_with_no_lab_rows_for_the_samples_every_shadow_row_is_missing_in_the_lab(
    tmp_path,
):
    result, outputs = compare(tmp_path, LAB_LOG)

    assert result.returncode == 0, result.stderr
    report = json.loads((outputs / "comparison.json").read_text())
    assert [r["status"] for r in report["rows"]] == ["missing_in_lab"] * 3
