"""parse_submission on every committed Fixture, run the way the Sandbox runs it
but without Docker."""

import json
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / "sandbox" / "scripts" / "parse_submission" / "main.py"
FIXTURES = sorted(
    form
    for form in (ROOT / "fixtures").rglob("*.docx")
    if form.with_suffix(".expected.json").exists()
)
REAL_LAYOUT = ROOT / "fixtures" / "real-layout"

FORM_ID = "c" * 64
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def stage(inputs: Path, form: Path) -> None:
    (inputs / FORM_ID).mkdir(parents=True)
    shutil.copyfile(form, inputs / FORM_ID / form.name)
    manifest = [
        {
            "artifact_id": FORM_ID,
            "filename": form.name,
            "artifact_type": "SubmissionForm",
            "media_type": DOCX,
        }
    ]
    (inputs / "manifest.json").write_text(json.dumps(manifest))


def run_script(inputs: Path, outputs: Path) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--inputs",
            str(inputs),
            "--outputs",
            str(outputs),
        ],
        capture_output=True,
        text=True,
        check=False,
    )


def test_there_is_at_least_one_fixture():
    assert FIXTURES, "fixtures/ holds no Submission Form docx"


@pytest.mark.parametrize("form", FIXTURES, ids=lambda p: p.name)
def test_parse_submission_writes_the_expected_record(form, tmp_path):
    inputs, outputs = tmp_path / "inputs", tmp_path / "outputs"
    outputs.mkdir()
    stage(inputs, form)

    result = run_script(inputs, outputs)

    assert result.returncode == 0, result.stderr
    expected = json.loads(form.with_suffix(".expected.json").read_text())
    assert json.loads((outputs / "submission.json").read_text()) == expected
    assert json.loads((outputs / "manifest.json").read_text()) == {
        "submission.json": {
            "artifact_type": "SubmissionRecord",
            "media_type": "application/json",
        }
    }
    assert sorted(p.name for p in outputs.iterdir()) == [
        "manifest.json",
        "submission.json",
    ]


def test_parse_submission_refuses_inputs_without_a_docx(tmp_path):
    inputs, outputs = tmp_path / "inputs", tmp_path / "outputs"
    inputs.mkdir()
    outputs.mkdir()
    (inputs / "manifest.json").write_text("[]")

    result = run_script(inputs, outputs)

    assert result.returncode == 1
    assert "exactly one Submission Form" in result.stderr
    assert list(outputs.iterdir()) == []


def parsed(form: Path, tmp_path: Path) -> dict:
    inputs, outputs = tmp_path / "inputs", tmp_path / "outputs"
    outputs.mkdir()
    stage(inputs, form)
    result = run_script(inputs, outputs)
    assert result.returncode == 0, result.stderr
    return json.loads((outputs / "submission.json").read_text())


def test_the_blank_real_form_is_read_as_the_real_layout_with_every_field_blank(
    tmp_path,
):
    record = parsed(REAL_LAYOUT / "submission-form.docx", tmp_path)

    values = [
        value
        for section in record.values()
        if isinstance(section, dict)
        for value in section.values()
    ]
    assert values and all(value == "" for value in values), record
    assert record["date_submitted"] == ""
    assert record["samples"] == []


def test_every_sample_line_of_a_form_is_a_sample_with_its_own_containers(tmp_path):
    form = REAL_LAYOUT / "improvised" / "received-2026-06-26-form-2.docx"

    samples = parsed(form, tmp_path)["samples"]

    assert [(s["no"], s["lot"], s["product_code"]) for s in samples] == [
        *[(str(n), "H159531", "1988U") for n in range(1, 7)],
        ("7", "A383081", "6442V"),
        ("8", "F782437", "2251V"),
    ]
    assert [(s["no_of_containers"], s["product_type"]) for s in samples] == [
        *[("1", "Sjtgdw")] * 4,
        *[("1", "Xsmmejc")] * 2,
        *[("1", "Sjtgdw")] * 2,
    ]


def test_a_blank_bill_to_block_on_the_real_form_is_the_submitters_party(tmp_path):
    from docx import Document

    doc = Document(str(REAL_LAYOUT / "submission-form-seed4.docx"))
    [bill_to] = [
        cell
        for cell in doc.tables[0].rows[2].cells
        if cell.text.startswith("Tallowby") and "Attn:" in cell.text
    ]
    for paragraph in bill_to.paragraphs:
        for run in list(paragraph.runs):
            run._r.getparent().remove(run._r)
    form = tmp_path / "blank-bill-to.docx"
    doc.save(str(form))

    record = parsed(form, tmp_path)

    expected = json.loads(
        (REAL_LAYOUT / "submission-form-seed4.expected.json").read_text()
    )
    assert record["bill_to"] == expected["submitted_by"]
