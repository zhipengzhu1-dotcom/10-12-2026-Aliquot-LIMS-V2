"""Stage 3 Script: write the Work Order for the Samples on a Receiving Log row.

Sandbox Script convention: run as
    python /script/main.py --inputs /inputs --outputs /outputs
with /inputs/manifest.json listing the inputs and each input at
/inputs/<artifact_id>/<filename>. The inputs must hold exactly one
`ReceivingLogRow` (the sheet `log_receipt` writes), exactly one
`SubmissionRecord` (the `submission/2` record that row was logged from) and
exactly one `LimsDatabase`, the toy LIMS database version the stage carries,
whose `samples` rows hold each Lab Reference's Lot, Sample type and TAT class.

One Work Order plans every Sample on the row, which may hold several Lots
and TAT classes: all the Samples one receipt logged. The Work Order number
is `WO-<year received>-<n>` with n written as four digits or more: n is one
more than the highest n for that year in `work_orders`.

Each Sample gets the Tests its record line requests. The row's Lab
References belong to the record's lines in order, `no_of_containers` each,
as `log_receipt` numbered them. A line's `tests_requested` holds requests
separated by `;` or line breaks, and each maps to a Test definition by its
words: "Average capsule weight" to that Test, and a request naming
nitrosamines (the word, an abbreviation in `NITROSAMINE`, or a name starting
"N-nitroso") to "Nitrosamine impurities", whose analytes are the
nitrosamines it names. A request naming none is refused, asking which, and
any other request is refused with its words. A line requesting nothing
keeps the slice 1 Test: "Average capsule weight" for an FP Sample, a refusal
for any other. The line's `limits`, split the same way, give each analyte
the part that names it, or the only part when the line names one analyte:
a number with `ng/day` (an acceptable intake), `ppm` or `ng/g` is written
`<number> <unit>`, and any other text is kept as `unread: <text>`. A part
that names no requested analyte is refused. The Sample type is read only
for that default, so a stability condition in it stays on the Sample.

A Sample's due date is its TAT class's number of business days
(`TAT_BUSINESS_DAYS`) after the date received, counting Monday to Friday
with no holiday calendar, and the Work Order is due when its last Sample is.

Outputs, typed in `/outputs/manifest.json`:

- `work_order.xlsx` (`WorkOrder`): one sheet named for the Work Order
  number. Rows 1 to 4 are the header block, a label in column A and its
  value in column B (`HEADER`); row 5 is blank; row 6 holds `COLUMNS`; then
  one row per Sample, Test and analyte, with the Sample's Lot, Product Code,
  TAT class and due date, and the analyte and its limit (blank for a Test
  without analytes).
- `lims.sqlite` (`LimsDatabase`): the carried database with one
  `work_orders` row per Sample and Test, each with the Sample's due date and,
  as `source_artifact`, the Receiving Log row's Artifact id.
- `work_order.json` (`Other`): for the stage 3 Evaluation, which has no
  xlsx reader, the Receiving Log row's Lab References and dates as this
  Script read them, and the Work Order as it reads it back from the saved
  xlsx: `{"schema": "work_order/2", "receiving_log_row": [[lab reference,
  YYYY-MM-DD], ...], "header": {label: text, ...}, "columns": [...],
  "rows": [[lab reference, test, lot, product code, TAT class, due date,
  analyte, limit], ...]}`.

The Script refuses inputs it cannot read or plan one Work Order from (exit
1, the reason on stderr) and judges nothing else: the stage 3 Evaluation
decides whether the due dates and the Tests are right.
"""

import argparse
import json
import re
import shutil
import sqlite3
import sys
import tempfile
from contextlib import closing
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook

SCHEMA_VERSION = 1
TAT_BUSINESS_DAYS = {"Regular": 10, "Rush1to2Days": 2, "Rush3to4Days": 4}
AVERAGE_CAPSULE_WEIGHT = "Average capsule weight"
NITROSAMINE_IMPURITIES = "Nitrosamine impurities"
NITROSAMINE = re.compile(
    r"\b(?:NDMA|NDEA|NMBA|NDIPA|NEIPA|NDBA|NMPA)\b|\bN-nitroso[\w-]*", re.IGNORECASE
)
NAMES_NITROSAMINES = re.compile(r"nitrosamine", re.IGNORECASE)
LIMIT = re.compile(r"(\d+(?:\.\d+)?)\s*(ng/day|ppm|ng/g)\b", re.IGNORECASE)
SEPARATOR = re.compile(r"[;\n]")
HEADER = ("Work Order", "Client", "Received", "Due Date")
COLUMNS = (
    "Lab Reference",
    "Test",
    "Lot",
    "Product Code",
    "TAT Class",
    "Due Date",
    "Analyte",
    "Limit",
)
DUE = COLUMNS.index("Due Date")
WORK_ORDER = re.compile(r"^WO-(\d{4})-(\d+)$")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class WorkOrderError(Exception):
    """The inputs are not a Receiving Log row and database this Script can read."""


@dataclass(frozen=True)
class Sample:
    lab_reference: str
    received: date
    lot: str
    product_code: str
    client: str
    sample_type: str
    tat_class: str


@dataclass(frozen=True)
class Planned:
    """One Test of a Sample, for one analyte, or for none ("")."""

    test: str
    analyte: str = ""
    limit: str = ""


def inputs_of(inputs: Path) -> tuple[Path, str, Path, Path]:
    """The Receiving Log row's path and Artifact id, the record's path, and
    the carried database's path."""
    manifest = json.loads((inputs / "manifest.json").read_text())

    def only(artifact_type: str) -> dict:
        found = [e for e in manifest if e["artifact_type"] == artifact_type]
        if len(found) != 1:
            raise WorkOrderError(
                f"expected exactly one {artifact_type} input, got {len(found)}"
            )
        return found[0]

    def path(entry: dict) -> Path:
        return inputs / entry["artifact_id"] / entry["filename"]

    row, database = only("ReceivingLogRow"), only("LimsDatabase")
    record = only("SubmissionRecord")
    return path(row), row["artifact_id"], path(record), path(database)


def read_receiving_log_row(path: Path) -> list[tuple[str, date]]:
    """Each row's Lab Reference (`#`) and date received (`Rec. DATE`)."""
    rows = list(load_workbook(path).active.iter_rows(values_only=True))
    if not rows:
        raise WorkOrderError("the Receiving Log row's sheet is empty")
    columns = list(rows[0])
    for name in ("#", "Rec. DATE"):
        if name not in columns:
            raise WorkOrderError(f"the Receiving Log row has no {name!r} column")
    ref, on = columns.index("#"), columns.index("Rec. DATE")
    read = []
    for row in rows[1:]:
        if not isinstance(row[on], datetime):
            raise WorkOrderError(f"{row[ref]}: Rec. DATE {row[on]!r} is not a date")
        read.append((str(row[ref]), row[on].date()))
    if not read:
        raise WorkOrderError("the Receiving Log row has no Samples")
    return read


def lines_of(record: dict, refs: list[str]) -> list[dict]:
    """The record line each Lab Reference was logged from, in row order."""
    lines = []
    for line in record["samples"]:
        count = line["no_of_containers"]
        if not count.isdigit():
            raise WorkOrderError(
                f"record line {line['no']}: no_of_containers {count!r} is not a count"
            )
        lines += [line] * int(count)
    if len(lines) != len(refs):
        raise WorkOrderError(
            f"the record's lines hold {len(lines)} containers, and the Receiving"
            f" Log row {len(refs)}"
        )
    return lines


def samples_of(db: sqlite3.Connection, rows: list[tuple[str, date]]) -> list[Sample]:
    samples = []
    for ref, received in rows:
        found = db.execute(
            "SELECT samples.lot, lots.product_code, lots.client,"
            " samples.sample_type, samples.tat_class"
            " FROM samples JOIN lots USING (lot) WHERE lab_reference = ?",
            (ref,),
        ).fetchone()
        if found is None:
            raise WorkOrderError(f"the carried database has no samples row for {ref}")
        samples.append(Sample(ref, received, *found))
    return samples


def parts(text: str) -> list[str]:
    return [part.strip() for part in SEPARATOR.split(text) if part.strip()]


def limit_of(text: str) -> str:
    found = LIMIT.search(text)
    return f"{found[1]} {found[2].lower()}" if found else f"unread: {text}"


def plan(line: dict, sample_type: str) -> list[Planned]:
    """The Tests a record line requests, or why it cannot be planned."""
    requests = parts(line.get("tests_requested", ""))
    if not requests:
        if sample_type == "FP":
            return [Planned(AVERAGE_CAPSULE_WEIGHT)]
        raise WorkOrderError(
            f"the line requests no Tests, and only an FP Sample has a default;"
            f" Sample type {sample_type!r}"
        )
    planned: list[Planned] = []
    analytes: list[str] = []
    for request in requests:
        named = [
            name if name.lower().startswith("n-nitroso") else name.upper()
            for name in NITROSAMINE.findall(request)
        ]
        if request.casefold() == AVERAGE_CAPSULE_WEIGHT.casefold():
            planned.append(Planned(AVERAGE_CAPSULE_WEIGHT))
        elif named:
            analytes += [a for a in dict.fromkeys(named) if a not in analytes]
        elif NAMES_NITROSAMINES.search(request):
            raise WorkOrderError(
                f"the request {request!r} names no nitrosamine; which should be tested?"
            )
        else:
            raise WorkOrderError(f"no Test definition for the request {request!r}")
    limits: dict[str, str] = {}
    for part in parts(line.get("limits", "")) if analytes else []:
        owners = [a for a in analytes if a.casefold() in part.casefold()]
        if not owners and len(analytes) == 1:
            owners = analytes
        if not owners:
            raise WorkOrderError(
                f"the limit {part!r} names none of the nitrosamines requested,"
                f" {', '.join(analytes)}"
            )
        limits.update(dict.fromkeys(owners, limit_of(part)))
    return planned + [
        Planned(NITROSAMINE_IMPURITIES, a, limits.get(a, "")) for a in analytes
    ]


def due_date(received: date, tat_class: str) -> date:
    day, left = received, TAT_BUSINESS_DAYS[tat_class]
    while left:
        day += timedelta(days=1)
        if day.weekday() < 5:
            left -= 1
    return day


def tests_of(samples: list[Sample], lines: list[dict]) -> list[tuple[object, ...]]:
    """One row per Sample, Test and analyte, in `COLUMNS` order."""
    rows = []
    for s, line in zip(samples, lines, strict=True):
        if line["lot"] != s.lot:
            raise WorkOrderError(
                f"{s.lab_reference}: the record's line {line['no']} is Lot"
                f" {line['lot']!r}, not the logged Lot {s.lot!r}"
            )
        try:
            planned = plan(line, s.sample_type)
        except WorkOrderError as error:
            raise WorkOrderError(f"{s.lab_reference}: {error}") from None
        due = due_date(s.received, s.tat_class)
        rows += [
            (s.lab_reference, p.test, s.lot, s.product_code, s.tat_class, due)
            + (p.analyte, p.limit)
            for p in planned
        ]
    return rows


def next_work_order(db: sqlite3.Connection, year: int) -> str:
    highest = 0
    for (number,) in db.execute("SELECT DISTINCT work_order FROM work_orders"):
        match = WORK_ORDER.match(number)
        if match and int(match[1]) == year:
            highest = max(highest, int(match[2]))
    return f"WO-{year}-{highest + 1:04d}"


def header_of(
    number: str, samples: list[Sample], rows: list[tuple[object, ...]]
) -> dict[str, object]:
    """The header block's values, or a refusal when the Samples were not
    received together from one Client."""
    first = samples[0]
    for field in ("received", "client"):
        values = sorted({str(getattr(s, field)) for s in samples})
        if len(values) > 1:
            raise WorkOrderError(
                f"the Receiving Log row's Samples differ in {field}: {values};"
                " one Work Order needs one"
            )
    values = (number, first.client, first.received, max(row[DUE] for row in rows))
    return dict(zip(HEADER, values, strict=True))


def write_sheet(
    path: Path, header: dict[str, object], rows: list[tuple[object, ...]]
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = str(header["Work Order"])
    for label, value in header.items():
        ws.append((label, value))
        if isinstance(value, date):
            ws.cell(ws.max_row, 2).number_format = "mm/dd/yyyy"
    ws.append(())
    ws.append(COLUMNS)
    for row in rows:
        ws.append(row)
        ws.cell(ws.max_row, DUE + 1).number_format = "mm/dd/yyyy"
    wb.save(path)


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    return str(value)


def read_back(path: Path) -> dict:
    rows = [
        [cell_text(v) for v in row]
        for row in load_workbook(path).active.iter_rows(values_only=True)
    ]
    blank = len(HEADER)
    return {
        "header": {label: value for label, value, *_ in rows[:blank]},
        "columns": rows[blank + 1],
        "rows": rows[blank + 2 :],
    }


def create_work_order(inputs: Path, outputs: Path) -> str:
    row_path, source, record_path, carried = inputs_of(inputs)
    receiving_log_row = read_receiving_log_row(row_path)
    lines = lines_of(
        json.loads(record_path.read_text()), [ref for ref, _ in receiving_log_row]
    )

    with tempfile.TemporaryDirectory() as work:
        database = Path(work) / "lims.sqlite"
        shutil.copyfile(carried, database)
        with closing(sqlite3.connect(database)) as db:
            (version,) = db.execute("PRAGMA user_version").fetchone()
            if version != SCHEMA_VERSION:
                raise WorkOrderError(
                    f"the carried database has schema version {version};"
                    f" this Script reads version {SCHEMA_VERSION}"
                )
            db.execute("PRAGMA foreign_keys = ON")
            samples = samples_of(db, receiving_log_row)
            number = next_work_order(db, samples[0].received.year)
            rows = tests_of(samples, lines)
            header = header_of(number, samples, rows)
            planned = dict.fromkeys((row[0], row[1], row[DUE]) for row in rows)
            db.executemany(
                "INSERT INTO work_orders (work_order, lab_reference, test, due_date,"
                " source_artifact) VALUES (?, ?, ?, ?, ?)",
                [
                    (number, ref, test, due.isoformat(), source)
                    for ref, test, due in planned
                ],
            )
            db.commit()
        shutil.copyfile(database, outputs / "lims.sqlite")

    write_sheet(outputs / "work_order.xlsx", header, rows)
    summary = {
        "schema": "work_order/2",
        "receiving_log_row": [[ref, on.isoformat()] for ref, on in receiving_log_row],
        **read_back(outputs / "work_order.xlsx"),
    }
    (outputs / "work_order.json").write_text(json.dumps(summary, indent=2) + "\n")
    manifest = {
        "work_order.xlsx": {"artifact_type": "WorkOrder", "media_type": XLSX},
        "lims.sqlite": {
            "artifact_type": "LimsDatabase",
            "media_type": "application/vnd.sqlite3",
        },
        "work_order.json": {"artifact_type": "Other", "media_type": "application/json"},
    }
    (outputs / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    return number


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--inputs", type=Path, required=True)
    parser.add_argument("--outputs", type=Path, required=True)
    args = parser.parse_args(argv)

    try:
        number = create_work_order(args.inputs, args.outputs)
    except (WorkOrderError, KeyError, sqlite3.Error) as error:
        print(f"create_work_order refused its inputs: {error!r}", file=sys.stderr)
        return 1
    print(f"wrote Work Order {number}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
