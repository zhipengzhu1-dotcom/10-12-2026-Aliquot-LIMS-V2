"""The shadow Receiving Log: a copy of the lab's workbook with the receipt's rows added.

The lab's Receiving Log workbook holds a sheet for each year, named by the
year, and a sheet for each day, named by the month and day as `July 23`. Its
headers that named the lab are not kept in the Fixtures, so every column is
found by its position, the one the lab's workbook gives it:

- The year's sheet: the column headers in row 3 and one row per container
  from row 4. A `#`, B `Rec. DATE`, C `Rec. Time`, D `Product Code`,
  E `PRODUCT NAME`, F `TYPE`, G `LOT#`, H `CLIENT`, I `EXP DATE`, J `QTY RECD`,
  K `SENT BY`, L `RECD BY`, M the Lab Reference; N to T are filled later by the
  lab (report date, book reference, comments, invoice, PO#, revision).
- The day's sheet: the column headers in row 2 and the rows from row 3, with A
  numbered ahead. A `#`, B `DATE`, C `PRODUCT NAME`, D `TYPE`, E `LOT#`,
  F `CLIENT`, G `EXP DATE`, H `QTY RECD`, I `RECD BY`, J the Lab Reference; K to
  R are filled later by the lab.

The rows go after the last filled row of the year's sheet, numbered on from
its `#`, and into the first rows of the day's sheet whose B to R are blank.
When the workbook has no sheet for the day, the day's sheet is a copy of the
last day's sheet with its rows from row 3 cleared, except the numbers in A.
New cells take the style of the cell above them. No cell that held a value
before is written.
"""

import re
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.worksheet import Worksheet

YEAR_FIRST_ROW = 4
YEAR_COLUMNS = 20  # A to T
DAY_FIRST_ROW = 3
DAY_COLUMNS = 18  # A to R
DAY_SHEET = re.compile(
    r"^(January|February|March|April|May|June|July|August|September|October"
    r"|November|December) \d{1,2}$"
)


class ShadowError(Exception):
    """The lab's workbook is not a Receiving Log this Script can add rows to."""


def cell_text(value: object) -> str:
    """A cell's value as text: dates YYYY-MM-DD, times HH:MM, blank as ""."""
    if value is None:
        return ""
    if isinstance(value, ArrayFormula):
        return value.text or ""
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def cells_of(ws: Worksheet) -> dict[str, str]:
    """Every cell that holds a value, by coordinate."""
    return {
        cell.coordinate: text
        for row in ws.iter_rows()
        for cell in row
        if (text := cell_text(cell.value)) != ""
    }


def _blank(ws: Worksheet, row: int, first: int, last: int) -> bool:
    return all(ws.cell(row, col).value in (None, "") for col in range(first, last + 1))


def _number(text: str) -> int | None:
    return int(text) if text.isdigit() else None


def _put(ws: Worksheet, row: int, values: dict[int, object]) -> None:
    for col, value in values.items():
        cell = ws.cell(row, col)
        above = ws.cell(row - 1, col)
        if above.has_style:
            cell._style = above._style
        cell.value = value


def _year_rows(ws: Worksheet, count: int) -> list[int]:
    last = YEAR_FIRST_ROW - 1
    for row in range(YEAR_FIRST_ROW, ws.max_row + 1):
        if not _blank(ws, row, 1, YEAR_COLUMNS):
            last = row
    return list(range(last + 1, last + 1 + count))


def _day_rows(ws: Worksheet, count: int) -> list[int]:
    rows: list[int] = []
    row = DAY_FIRST_ROW
    while len(rows) < count:
        if _blank(ws, row, 2, DAY_COLUMNS):
            rows.append(row)
        row += 1
    return rows


def _day_sheet(wb, name: str) -> tuple[Worksheet, bool]:
    """The day's sheet, and whether this Script made it."""
    if name in wb.sheetnames:
        return wb[name], False
    days = [ws for ws in wb.worksheets if DAY_SHEET.match(ws.title)]
    if not days:
        raise ShadowError(
            f"the Receiving Log has no sheet for {name} and no day's sheet to copy"
        )
    ws = wb.copy_worksheet(days[-1])
    ws.title = name
    for row in ws.iter_rows(min_row=DAY_FIRST_ROW, min_col=2, max_col=DAY_COLUMNS):
        for cell in row:
            cell.value = None
    return ws, True


def write_shadow(
    lab: Path,
    shadow: Path,
    record: dict,
    on: date,
    at: time,
    refs: list[str],
) -> dict:
    """Write the shadow and return what it holds, for the stage 2 Evaluation."""
    try:
        wb = load_workbook(lab)
    except Exception as error:
        raise ShadowError(f"the Receiving Log is not a workbook: {error}") from error
    before = {ws.title: cells_of(ws) for ws in wb.worksheets}
    year_name, day_name = str(on.year), f"{on:%B} {on.day}"
    if year_name not in wb.sheetnames:
        raise ShadowError(f"the Receiving Log has no sheet {year_name}")
    year = wb[year_name]
    day, made = _day_sheet(wb, day_name)

    sample, receipt = record["sample"], record["receipt"]
    received = datetime.combine(on, time(0, 0))
    client = record["submitted_by"]["company"]
    year_rows = _year_rows(year, len(refs))
    previous = (
        0
        if year_rows[0] == YEAR_FIRST_ROW
        else _number(cell_text(year.cell(year_rows[0] - 1, 1).value))
    )
    for i, (row, ref) in enumerate(zip(year_rows, refs, strict=True)):
        number = "" if previous is None else str(previous + 1 + i)
        _put(
            year,
            row,
            {
                1: number,
                2: received,
                3: at,
                4: sample["product_code"],
                5: sample["sample_description"],
                6: sample["sample_type"],
                7: sample["lot"],
                8: client,
                10: 1,
                11: record["submitted_by"]["contact"],
                12: receipt["received_by"],
                13: ref,
            },
        )
    day_rows = _day_rows(day, len(refs))
    for row, ref in zip(day_rows, refs, strict=True):
        values: dict[int, object] = {
            2: received,
            3: sample["sample_description"],
            4: sample["sample_type"],
            5: sample["lot"],
            6: client,
            8: 1,
            9: receipt["received_by"],
            10: ref,
        }
        if day.cell(row, 1).value in (None, ""):
            above = _number(cell_text(day.cell(row - 1, 1).value))
            values[1] = 1 if above is None else above + 1
        _put(day, row, values)
    wb.save(shadow)

    after = {ws.title: cells_of(ws) for ws in load_workbook(shadow).worksheets}
    return {
        "schema": "shadow/1",
        "year": {"sheet": year_name, "rows": year_rows},
        "day": {"sheet": day_name, "rows": day_rows, "made": made},
        "sheets": [
            {"name": name, "before": before.get(name), "after": after.get(name)}
            for name in dict.fromkeys([*before, *after])
        ],
    }
