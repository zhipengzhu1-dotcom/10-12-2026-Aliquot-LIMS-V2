"""Stage 2 Script: log a record's containers as Receiving Log rows and Sample rows.

The Sample rows go into the toy LIMS database.

Sandbox Script convention: run as
    python /script/main.py --inputs /inputs --outputs /outputs
with /inputs/manifest.json listing the inputs and each input at
/inputs/<artifact_id>/<filename>. The inputs must hold exactly one
`SubmissionRecord` (the `submission/1` record `parse_submission` writes) and
at most one `LimsDatabase`, the version of the toy LIMS database the stage
carries. With none, this Script creates the database with the schema in
`lims_schema.py`.

Each container gets a Lab Reference, `LR-<year received>-<n>` with n written
as four digits or more: n is one more than the highest n for that year in
`samples`, and the containers take consecutive numbers in order.

Outputs, typed in `/outputs/manifest.json`:

- `receiving_log_row.xlsx` (`ReceivingLogRow`): one sheet named for the day
  received (MM-DD), the ten Receiving Log columns in row 1, and one row per
  container. `#` holds the container's Lab Reference. `EXP DATE` is blank,
  because the Submission Form gives no expiry date, and `QTY RECD` is 1, one
  container per row.
- `lims.sqlite` (`LimsDatabase`): the carried database, or a new one, with
  one `lots` row for the Lot (an existing row for the Lot is kept) and one
  `samples` row per container. Each new row's `source_artifact` is the
  record's Artifact id.
- `receipt.json` (`Other`): the sheet as this Script reads it back from the
  saved xlsx, for the stage 2 Evaluation, which has no xlsx reader:
  `{"schema": "receipt/1", "sheet": "MM-DD", "columns": [...],
  "rows": [[cell text, ...], ...]}`. Dates read back as YYYY-MM-DD, times
  as HH:MM, a blank cell as "".

The inputs may also hold one `ReceivingLog`, the lab's Receiving Log workbook
as the Conductor read it (ADR-0018). The Script never writes it. With one, it
also writes:

- `receiving_log_shadow.xlsx` (`ShadowReceivingLog`): the Shadow, a copy of
  the lab's workbook with one row per container added to the year's sheet and
  the day's sheet, in the real columns (`shadow.py` names them).
- `shadow.json` (`Other`): what the Shadow holds, for the stage 2 Evaluation
  and the Comparison: `{"schema": "shadow/1", "year": {"sheet", "rows"},
  "day": {"sheet", "rows", "made"}, "sheets": [{"name", "before", "after"}]}`.
  `rows` are the new rows' numbers, `made` says whether the day's sheet is
  new, and `sheets` lists every sheet of either workbook, whose `before` and
  `after` map each coordinate that holds a value to its text in the lab's
  workbook and in the Shadow (null where the workbook lacks the sheet).

The Script refuses inputs it cannot read (exit 1, the reason on stderr) and
judges nothing else: the stage 2 Evaluation decides whether the rows match
the record.
"""

import argparse
import json
import re
import shutil
import sqlite3
import sys
import tempfile
from contextlib import closing
from datetime import date, datetime, time
from pathlib import Path

from lims_schema import SCHEMA_VERSION, create_schema
from openpyxl import Workbook, load_workbook
from shadow import ShadowError, write_shadow

COLUMNS = (
    "#",
    "Rec. DATE",
    "Rec. Time",
    "Product Code",
    "PRODUCT NAME",
    "TYPE",
    "LOT#",
    "CLIENT",
    "EXP DATE",
    "QTY RECD",
)
LAB_REFERENCE = re.compile(r"^LR-(\d{4})-(\d+)$")
PARTIES = ("submitted_by", "bill_to", "report_to")
PARTY_FIELDS = ("company", "contact", "address", "city", "phone", "email")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ReceiptError(Exception):
    """The inputs are not a record and database this Script can read."""


def inputs_of(inputs: Path) -> tuple[Path, str, Path | None, Path | None]:
    """The record's path and Artifact id, and the carried database's and
    the lab's Receiving Log's paths."""
    manifest = json.loads((inputs / "manifest.json").read_text())

    def of_type(artifact_type: str) -> list[dict]:
        return [e for e in manifest if e["artifact_type"] == artifact_type]

    def path(entry: dict) -> Path:
        return inputs / entry["artifact_id"] / entry["filename"]

    def optional(artifact_type: str) -> Path | None:
        entries = of_type(artifact_type)
        if len(entries) > 1:
            raise ReceiptError(
                f"expected at most one {artifact_type} input, got {len(entries)}"
            )
        return path(entries[0]) if entries else None

    records = of_type("SubmissionRecord")
    if len(records) != 1:
        raise ReceiptError(
            f"expected exactly one SubmissionRecord input, got {len(records)}"
        )
    [record] = records
    return (
        path(record),
        record["artifact_id"],
        optional("LimsDatabase"),
        optional("ReceivingLog"),
    )


def received(record: dict) -> tuple[date, time, int]:
    """The date and time received and the number of containers."""
    receipt, sample = record["receipt"], record["sample"]
    try:
        on = date.fromisoformat(receipt["date_received"])
        at = time.fromisoformat(receipt["time_received"])
    except ValueError as error:
        raise ReceiptError(f"the record's receipt is unreadable: {error}") from error
    containers = sample["no_of_containers"]
    if not containers.isdigit() or int(containers) < 1:
        raise ReceiptError(
            f"No. of Containers {containers!r} is not a whole number of at least 1"
        )
    return on, at, int(containers)


def next_lab_references(db: sqlite3.Connection, year: int, count: int) -> list[str]:
    highest = 0
    for (ref,) in db.execute("SELECT lab_reference FROM samples"):
        match = LAB_REFERENCE.match(ref)
        if match and int(match[1]) == year:
            highest = max(highest, int(match[2]))
    return [f"LR-{year}-{n:04d}" for n in range(highest + 1, highest + count + 1)]


def write_database(
    db: sqlite3.Connection, record: dict, source: str, refs: list[str]
) -> None:
    sample, receipt = record["sample"], record["receipt"]
    db.execute(
        "INSERT INTO lots (lot, product_code, product_name, product_type, client,"
        " exp_date, source_artifact) VALUES (?, ?, ?, ?, ?, NULL, ?)"
        " ON CONFLICT (lot) DO NOTHING",
        (
            sample["lot"],
            sample["product_code"],
            sample["sample_description"],
            sample["product_type"],
            record["submitted_by"]["company"],
            source,
        ),
    )
    parties = {
        f"{party}_{field}": record[party].get(field, "")
        for party in PARTIES
        for field in PARTY_FIELDS
    }
    for container, ref in enumerate(refs, start=1):
        row = {
            "lab_reference": ref,
            "lot": sample["lot"],
            "container": container,
            "qty_received": 1,
            "sample_type": sample["sample_type"],
            "tat_code": sample["tat_code"],
            "tat_class": sample["tat_class"],
            "date_submitted": record["date_submitted"],
            "date_received": receipt["date_received"],
            "time_received": receipt["time_received"],
            "received_by": receipt["received_by"],
            **parties,
            "source_artifact": source,
        }
        db.execute(
            f"INSERT INTO samples ({', '.join(row)})"
            f" VALUES ({', '.join('?' for _ in row)})",
            tuple(row.values()),
        )


def write_rows(path: Path, record: dict, on: date, at: time, refs: list[str]) -> None:
    sample = record["sample"]
    wb = Workbook()
    ws = wb.active
    ws.title = f"{on:%m-%d}"
    ws.append(COLUMNS)
    for ref in refs:
        ws.append(
            (
                ref,
                on,
                at,
                sample["product_code"],
                sample["sample_description"],
                sample["product_type"],
                sample["lot"],
                record["submitted_by"]["company"],
                None,
                1,
            )
        )
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "mm/dd/yyyy"
        row[2].number_format = "hh:mm"
    wb.save(path)


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return str(value)


def read_back(path: Path) -> dict:
    ws = load_workbook(path).active
    rows = [[cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
    return {
        "schema": "receipt/1",
        "sheet": ws.title,
        "columns": rows[0],
        "rows": rows[1:],
    }


def log_receipt(inputs: Path, outputs: Path) -> list[str]:
    record_path, source, carried, lab_log = inputs_of(inputs)
    record = json.loads(record_path.read_text())
    on, at, containers = received(record)

    with tempfile.TemporaryDirectory() as work:
        database = Path(work) / "lims.sqlite"
        if carried is not None:
            shutil.copyfile(carried, database)
        with closing(sqlite3.connect(database)) as db:
            if carried is None:
                create_schema(db)
            (version,) = db.execute("PRAGMA user_version").fetchone()
            if version != SCHEMA_VERSION:
                raise ReceiptError(
                    f"the carried database has schema version {version};"
                    f" this Script writes version {SCHEMA_VERSION}"
                )
            db.execute("PRAGMA foreign_keys = ON")
            refs = next_lab_references(db, on.year, containers)
            write_database(db, record, source, refs)
            db.commit()
        shutil.copyfile(database, outputs / "lims.sqlite")

    write_rows(outputs / "receiving_log_row.xlsx", record, on, at, refs)
    receipt = read_back(outputs / "receiving_log_row.xlsx")
    (outputs / "receipt.json").write_text(json.dumps(receipt, indent=2) + "\n")
    manifest = {
        "receiving_log_row.xlsx": {
            "artifact_type": "ReceivingLogRow",
            "media_type": XLSX,
        },
        "lims.sqlite": {
            "artifact_type": "LimsDatabase",
            "media_type": "application/vnd.sqlite3",
        },
        "receipt.json": {"artifact_type": "Other", "media_type": "application/json"},
    }
    if lab_log is not None:
        held = write_shadow(
            lab_log, outputs / "receiving_log_shadow.xlsx", record, on, at, refs
        )
        (outputs / "shadow.json").write_text(json.dumps(held, indent=2) + "\n")
        manifest["receiving_log_shadow.xlsx"] = {
            "artifact_type": "ShadowReceivingLog",
            "media_type": XLSX,
        }
        manifest["shadow.json"] = {
            "artifact_type": "Other",
            "media_type": "application/json",
        }
    (outputs / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    return refs


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--inputs", type=Path, required=True)
    parser.add_argument("--outputs", type=Path, required=True)
    args = parser.parse_args(argv)

    try:
        refs = log_receipt(args.inputs, args.outputs)
    except (ReceiptError, ShadowError, KeyError, sqlite3.Error) as error:
        print(f"log_receipt refused its inputs: {error!r}", file=sys.stderr)
        return 1
    print(f"logged {len(refs)} containers: {', '.join(refs)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
