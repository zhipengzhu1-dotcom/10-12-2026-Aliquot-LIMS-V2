"""Stage 5 Script: verify a Result Worksheet against the weighings.

Sandbox Script convention: run as
    python /script/main.py --inputs /inputs --outputs /outputs
with /inputs/manifest.json listing the inputs and each input at
/inputs/<artifact_id>/<filename>. The inputs must hold exactly one
`ResultWorksheet` (the sheet `compute_results` writes) and exactly one
`Other`, the weighings: one JSON array of `{"container": n, ...}`, each
entry weighed for Average capsule weight (`"weight_of_20_g": g,
"weight_taken_g": g`) or carrying nitrosamine readings (`"nitrosamines":
{...}`), the shapes `compute_results` documents.

The worksheet holds a block of `COLUMNS`, one row per container, `#`
holding the container number, then a Mean row with "Mean" in column G, or
a block of `NITROSAMINE_COLUMNS`, one row per container and analyte, or
both. The Script recomputes each cell it checks by evaluating it on the
worksheet's own cells: a formula of cells and numbers joined by `*` and
`/`, or `=AVERAGE(<cell>:<cell>)`, or a typed number as it stands. It
compares each container row's Average Weight (g) with the weighing of the
row's container divided by 20, and the Mean row with the mean of the
weighings' averages. It compares each nitrosamine row's Content (ppm) and
Intake (ng/day) with its own arithmetic on the row's readings (the measured
ng/mL over the weight taken / VF * dilution, and that times the maximum
daily dose), within `TOLERANCE_RELATIVE` of it, and the row's Judgement
with its own judgement of the row's Limit, as `compute_results` documents
the rule.

Output, typed in `/outputs/manifest.json`:

- `verification_record.json` (`VerificationRecord`):
  `{"schema": "verification_record/1", "worksheet": <artifact id>,
  "weighings": <artifact id>, "work_order": <number>, "tolerance_g": 0.0001,
  "rows": [{"row", "container", "formula", "recomputed_g", "weighing_g",
  "difference_g", "within_tolerance"}, ...], "mean": {"row", "formula",
  "recomputed_g", "weighing_g", "difference_g", "within_tolerance"} | null,
  "missing_containers": [n, ...], "tolerance_relative": 1e-06,
  "nitrosamines": [{"row", "container", "analyte", "limit",
  "below_reporting_limit", "content_ppm": <checked>, "intake_ng_day":
  <checked> | null, "judgement", "expected_judgement"}, ...],
  "missing_readings": ["container <n>: <analyte>", ...], "verdict": "pass"
  | "fail"}`, where a checked value is `{"formula", "recomputed",
  "expected", "within_tolerance"}`. `formula` is null when the cell holds
  a typed value; a recomputed value is null when the cell is neither a
  number nor a formula this Script evaluates; `weighing_g` and `expected`
  are null when no weighing or reading has the row's container and
  analyte. `mean` is null when the worksheet has no Average capsule weight
  block, and `intake_ng_day` when the row has neither an Intake nor a daily
  dose. `missing_containers` lists the weighed containers with no row, and
  `missing_readings` the readings with no row. The verdict is `pass` when
  the worksheet has a row, every row and the mean are within tolerance,
  every judgement is the expected one, and every weighing and reading has a
  row.

The Script refuses inputs it cannot read as a worksheet and weighings
(exit 1, the reason on stderr).
"""

import argparse
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.worksheet import Worksheet

TOLERANCE_G = 0.0001
TOLERANCE_RELATIVE = 1e-6
COLUMNS = (
    "#",
    "Code",
    "Name",
    "Status",
    "Lot#",
    "Weight of 20 Capsules (g)",
    "Weight Taken (g)",
    "Average Weight (g)",
)
NITROSAMINE_COLUMNS = (
    "#",
    "Analyte",
    "Wt Taken (mg)",
    "VF (mL)",
    "Dilution",
    "Conc. (mg/mL)",
    "Measured (ng/mL)",
    "Reporting Limit (ng/mL)",
    "Content (ppm)",
    "Max Daily Dose (mg/day)",
    "Intake (ng/day)",
    "Limit",
    "Judgement",
)
CONTAINER, MEAN_LABEL, AVERAGE = 1, 7, 8
ANALYTE, CONTENT, INTAKE, LIMIT, JUDGEMENT = "B", "I", "K", "L", "M"
OPERAND = r"[A-Z]+\d+|\d+(?:\.\d+)?"
PRODUCT = re.compile(rf"=({OPERAND})((?:[*/](?:{OPERAND}))*)")
TERM = re.compile(rf"([*/])({OPERAND})")
CELL = re.compile(r"([A-Z]+)(\d+)")
AVERAGE_OF = re.compile(r"=AVERAGE\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)")
# Phrased as compute_results phrases its judgements.
PASS, FAIL = "Pass", "Fail"
NO_LIMIT, LIMIT_UNREAD, NO_DAILY_DOSE = "No limit", "Limit unread", "No daily dose"
PPM_PER = {"ppm": 1.0, "ng/g": 1000.0}


class VerificationError(Exception):
    """The inputs are not a worksheet and weighings this Script can read."""


@dataclass(frozen=True)
class Inputs:
    worksheet: Path
    worksheet_id: str
    weighings: Path
    weighings_id: str


@dataclass(frozen=True)
class Weighings:
    """Each container's weight of 20 capsules, and each container's
    nitrosamine preparation as given, by container number."""

    capsules: dict[int, float]
    nitrosamines: dict[int, dict]


def inputs_of(inputs: Path) -> Inputs:
    manifest = json.loads((inputs / "manifest.json").read_text())

    def only(artifact_type: str) -> dict:
        found = [e for e in manifest if e["artifact_type"] == artifact_type]
        if len(found) != 1:
            raise VerificationError(
                f"expected exactly one {artifact_type} input, got {len(found)}"
            )
        return found[0]

    def path(entry: dict) -> Path:
        return inputs / entry["artifact_id"] / entry["filename"]

    worksheet, weighings = only("ResultWorksheet"), only("Other")
    return Inputs(
        path(worksheet),
        worksheet["artifact_id"],
        path(weighings),
        weighings["artifact_id"],
    )


def is_number(value: object) -> bool:
    return isinstance(value, int | float) and not isinstance(value, bool)


def read_weighings(path: Path) -> Weighings:
    try:
        entries = json.loads(path.read_text())
    except json.JSONDecodeError as error:
        raise VerificationError(f"the weighings are not JSON: {error}") from error
    if not isinstance(entries, list):
        raise VerificationError("the weighings are not a JSON array")
    capsules: dict[int, float] = {}
    nitrosamines: dict[int, dict] = {}
    seen: set[int] = set()
    for entry in entries:
        container = entry.get("container") if isinstance(entry, dict) else None
        if not isinstance(container, int) or isinstance(container, bool):
            raise VerificationError(f"weighing {entry!r} has no whole container")
        if container in seen:
            raise VerificationError(f"container {container} is weighed twice")
        seen.add(container)
        if "nitrosamines" in entry:
            given = entry["nitrosamines"]
            if not isinstance(given, dict) or not isinstance(
                given.get("readings"), dict
            ):
                raise VerificationError(f"container {container} has no readings")
            nitrosamines[container] = given
        if "nitrosamines" not in entry or "weight_of_20_g" in entry:
            weight = entry.get("weight_of_20_g")
            if not is_number(weight):
                raise VerificationError(f"weighing {entry!r} has no weight_of_20_g")
            capsules[container] = float(weight)
    return Weighings(capsules, nitrosamines)


def value_of(ws: Worksheet, column: str, row: int, seen: frozenset = frozenset()):
    """The number a cell gives, evaluating the formula shapes this Script
    knows on the worksheet's own cells, or None."""
    if (column, row) in seen:
        return None
    seen = seen | {(column, row)}
    cell = ws.cell(row, column_index_from_string(column)).value
    if is_number(cell):
        return float(cell)
    if not isinstance(cell, str):
        return None

    def operand(text: str):
        if match := CELL.fullmatch(text):
            return value_of(ws, match[1], int(match[2]), seen)
        return float(text)

    if match := PRODUCT.fullmatch(cell):
        value = operand(match[1])
        for op, text in TERM.findall(match[2]):
            other = operand(text)
            if value is None or other is None or (op == "/" and other == 0):
                return None
            value = value * other if op == "*" else value / other
        return value
    if match := AVERAGE_OF.fullmatch(cell):
        first_col, first, last_col, last = match.groups()
        if first_col != last_col or int(first) > int(last):
            return None
        values = [
            value_of(ws, first_col, at, seen) for at in range(int(first), int(last) + 1)
        ]
        if any(v is None for v in values):
            return None
        return sum(values) / len(values)
    return None


def checked(row: int, formula, recomputed, weighing) -> dict:
    difference = (
        None if recomputed is None or weighing is None else recomputed - weighing
    )
    return {
        "row": row,
        "formula": formula,
        "recomputed_g": recomputed,
        "weighing_g": weighing,
        "difference_g": difference,
        "within_tolerance": difference is not None and abs(difference) <= TOLERANCE_G,
    }


def close(recomputed, expected) -> bool:
    return (
        recomputed is not None
        and expected is not None
        and abs(recomputed - expected) <= TOLERANCE_RELATIVE * max(1.0, abs(expected))
    )


def expected_judgement(
    limit: str, content: float | None, intake: float | None, below: bool
) -> str | None:
    if not limit:
        return NO_LIMIT
    if limit.startswith("unread: "):
        return LIMIT_UNREAD
    if below:
        return PASS
    number, _, unit = limit.partition(" ")
    if unit == "ng/day":
        if intake is None:
            return NO_DAILY_DOSE
        measured = intake
    elif unit in PPM_PER and content is not None:
        measured = content * PPM_PER[unit]
    else:
        return None
    return PASS if measured <= float(number) else FAIL


def verify_nitrosamine_row(ws: Worksheet, row: int, given: dict | None) -> dict:
    container = ws.cell(row, CONTAINER).value
    analyte = ws[f"{ANALYTE}{row}"].value
    limit = ws[f"{LIMIT}{row}"].value or ""
    reading = (given or {}).get("readings", {}).get(analyte)
    content = intake = None
    below = False
    if given is not None and isinstance(reading, dict):
        mg_ml = given["weight_taken_mg"] / given["final_volume_ml"] * given["dilution"]
        content = reading["concentration_ng_ml"] / mg_ml
        dose = given.get("max_daily_dose_mg")
        intake = None if dose is None else content * dose
        below = reading["concentration_ng_ml"] < reading["reporting_limit_ng_ml"]

    def check(column: str, expected) -> dict:
        cell = ws[f"{column}{row}"].value
        recomputed = value_of(ws, column, row)
        return {
            "formula": cell if isinstance(cell, str) and cell.startswith("=") else None,
            "recomputed": recomputed,
            "expected": expected,
            "within_tolerance": close(recomputed, expected),
        }

    dosed = given is not None and given.get("max_daily_dose_mg") is not None
    has_intake = ws[f"{INTAKE}{row}"].value is not None
    return {
        "row": row,
        "container": container,
        "analyte": analyte,
        "limit": limit,
        "below_reporting_limit": below,
        "content_ppm": check(CONTENT, content),
        "intake_ng_day": check(INTAKE, intake) if dosed or has_intake else None,
        "judgement": ws[f"{JUDGEMENT}{row}"].value,
        "expected_judgement": (
            None
            if content is None
            else expected_judgement(str(limit), content, intake, below)
        ),
    }


def verify(worksheet: Path, weighings: Weighings) -> dict:
    """The record's work order, rows, mean, nitrosamines, missing weighings
    and readings, and verdict."""
    ws = load_workbook(worksheet).active
    rows = range(1, ws.max_row + 1)

    def cells(row: int, count: int) -> tuple:
        return tuple(ws.cell(row, col).value for col in range(1, count + 1))

    header = next((r for r in rows if cells(r, len(COLUMNS)) == COLUMNS), None)
    nitro_header = next(
        (r for r in rows if cells(r, len(NITROSAMINE_COLUMNS)) == NITROSAMINE_COLUMNS),
        None,
    )
    if header is None and nitro_header is None:
        raise VerificationError(
            "the worksheet has no row of the Result Worksheet's columns"
        )
    number = next(
        (ws.cell(r, 2).value for r in rows if ws.cell(r, 1).value == "Work Order"),
        None,
    )

    def formula_at(row: int):
        cell = ws.cell(row, AVERAGE).value
        return cell if isinstance(cell, str) and cell.startswith("=") else None

    checks, mean = [], None
    if header is not None:
        mean_at = next(
            (r for r in rows if r > header and ws.cell(r, MEAN_LABEL).value == "Mean"),
            None,
        )
        if mean_at is None:
            raise VerificationError("the worksheet has no Mean row")
        for row in range(header + 1, mean_at):
            container = ws.cell(row, CONTAINER).value
            weight = (
                weighings.capsules.get(container)
                if isinstance(container, int)
                else None
            )
            checks.append(
                {"container": container}
                | checked(
                    row,
                    formula_at(row),
                    value_of(ws, "H", row),
                    None if weight is None else weight / 20,
                )
            )
        averages = [w / 20 for w in weighings.capsules.values()]
        mean = checked(
            mean_at,
            formula_at(mean_at),
            value_of(ws, "H", mean_at),
            sum(averages) / len(averages) if averages else None,
        )
    on_sheet = {c["container"] for c in checks}
    missing = sorted(set(weighings.capsules) - on_sheet)

    nitrosamines = []
    if nitro_header is not None:
        row = nitro_header + 1
        while ws.cell(row, CONTAINER).value is not None:
            container = ws.cell(row, CONTAINER).value
            nitrosamines.append(
                verify_nitrosamine_row(ws, row, weighings.nitrosamines.get(container))
            )
            row += 1
    listed = {(n["container"], n["analyte"]) for n in nitrosamines}
    missing_readings = [
        f"container {container}: {analyte}"
        for container, given in sorted(weighings.nitrosamines.items())
        for analyte in given["readings"]
        if (container, analyte) not in listed
    ]

    passed = (
        bool(checks or nitrosamines)
        and all(c["within_tolerance"] for c in checks)
        and (mean is None or mean["within_tolerance"])
        and not missing
        and all(
            n["content_ppm"]["within_tolerance"]
            and (n["intake_ng_day"] is None or n["intake_ng_day"]["within_tolerance"])
            and n["judgement"] is not None
            and n["judgement"] == n["expected_judgement"]
            for n in nitrosamines
        )
        and not missing_readings
    )
    return {
        "work_order": number,
        "rows": checks,
        "mean": mean,
        "missing_containers": missing,
        "tolerance_relative": TOLERANCE_RELATIVE,
        "nitrosamines": nitrosamines,
        "missing_readings": missing_readings,
        "verdict": "pass" if passed else "fail",
    }


def verify_results(inputs: Path, outputs: Path) -> dict:
    given = inputs_of(inputs)
    found = verify(given.worksheet, read_weighings(given.weighings))
    record = {
        "schema": "verification_record/1",
        "worksheet": given.worksheet_id,
        "weighings": given.weighings_id,
        "tolerance_g": TOLERANCE_G,
    } | found
    text = json.dumps(record, indent=2) + "\n"
    (outputs / "verification_record.json").write_text(text)
    manifest = {
        "verification_record.json": {
            "artifact_type": "VerificationRecord",
            "media_type": "application/json",
        }
    }
    (outputs / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    return record


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--inputs", type=Path, required=True)
    parser.add_argument("--outputs", type=Path, required=True)
    args = parser.parse_args(argv)

    try:
        record = verify_results(args.inputs, args.outputs)
    except (VerificationError, KeyError, TypeError, ZeroDivisionError) as error:
        print(f"verify_results refused its inputs: {error!r}", file=sys.stderr)
        return 1
    print(
        f"verified the Result Worksheet of Work Order {record['work_order']}:"
        f" {record['verdict']}"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
