"""Stage 6 Script: render the Report of Analysis for a verified Work Order.

Sandbox Script convention: run as
    python /script/main.py --inputs /inputs --outputs /outputs
with /inputs/manifest.json listing the inputs and each input at
/inputs/<artifact_id>/<filename>. The inputs must hold exactly one
`VerificationRecord` (the `verification_record/1` record `verify_results`
writes), exactly one `ReceivingLogRow` (the sheet `log_receipt` writes),
exactly one `Other`, the weighings stage 4 computed from, and exactly one
`LimsDatabase`, the toy LIMS database version the stage carries.

The report is for the first Sample of the record's Work Order, by Lab
Reference, that the Receiving Log row lists. Every metadata value is read
from that Sample's Receiving Log row or from its `samples` row, as `CLIENT`
and `METADATA` name the source of each; the Tests from its `work_orders`
rows. For Average capsule weight the result is the record's recomputed
mean, rounded to 4 places, and the unit is that of the Sample's `results`
row for the Test. For Nitrosamine impurities there is one row per analyte
the record lists for the Sample's container, `nitrosamine_row()`: the
result in the unit of its limit (the intake in ng/day for a limit in ng/day
when a daily dose was given, the content in ng/g for a limit in ng/g, else
the content in ppm), recomputed by stage 5 and rounded to 4 places, or
"< RL" below the reporting limit, and the record's judgement. The
letterhead, the Methods, the Specification of a result without a limit and
the Document Revision belong to the report form, not to the Sample. The
signature area names this Script and the Verification Record, whose Runs'
actors the Ledger holds, and, below them, the notebook reference the
weighings name, when they name one.

Outputs, typed in `/outputs/manifest.json`. The xlsx and the PDF are drawn
from one `Report`, so they cannot disagree:

- `report_of_analysis.xlsx` (`ReportOfAnalysis`): the layout of
  `sandbox/synth/report_of_analysis.py`. Letterhead in rows 1 to 5, the
  Client block from row 7 and the metadata block from row 12 (labels in
  column A, values in column B), the results table's columns in row 22 and
  its rows from row 23, and the signature area from two rows below the
  last result (row 26 for one result), with the notebook reference's label
  and value five rows below that.
- `report_of_analysis.pdf` (`ReportOfAnalysis`): the same report, drawn by
  reportlab with uncompressed page streams and invariant metadata, so the
  stage 6 Evaluation reads its text from the bytes and the same report
  gives the same bytes.
- `lims.sqlite` (`LimsDatabase`): the carried database with one `reports`
  row for the Sample: today's date (UTC) as `issued_date`, the form's
  Document Revision, and the PDF's Artifact id, the sha256 of its bytes, as
  `source_artifact`.
- `report_of_analysis.json` (`Other`): for the stage 6 Evaluation, which has
  no xlsx reader: `{"schema": "report_of_analysis/1", "receiving_log_row":
  {column: text}, "sheet": [[cell, ...], ...]}`, the Sample's Receiving Log
  row as this Script read it and the report's sheet as it reads back the
  saved xlsx, row 1 first. Dates read back as YYYY-MM-DD, times as HH:MM.

The Script refuses inputs it cannot read or render a complete report from,
including any metadata value that is missing or blank (exit 1, the reason on
stderr, no outputs). It judges nothing else: the stage 6 Evaluation decides
whether the report agrees with its sources and whether the verdict was pass.
"""

import argparse
import hashlib
import json
import shutil
import sqlite3
import sys
import tempfile
from contextlib import closing
from dataclasses import dataclass
from datetime import UTC, date, datetime, time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen.canvas import Canvas

SCHEMA_VERSION = 1
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

LAB_NAME = "Quillmere Analytical Laboratories"
LAB_ADDRESS = "27 Wrenfield Court, Velmora, NJ 07999"
LAB_CONTACT = "Tel (201) 555-0100 · reports@quillmere.example"
TITLE = "REPORT OF ANALYSIS"
SHEET = "Report of Analysis"
DOCUMENT_REVISION = "Rev. 03"
METHOD = "QAL-SOP-114"
NITROSAMINE_METHOD = "QAL-SOP-207"
SPECIFICATION = "Report value"
RESULT_COLUMNS = ("Test", "Method", "Specification", "Result", "Unit", "Judgement")
AVERAGE_CAPSULE_WEIGHT = "Average capsule weight"
NITROSAMINE_IMPURITIES = "Nitrosamine impurities"
BELOW_REPORTING_LIMIT = "< RL"
UNREAD = "unread: "
SIGNATURE_ROLES = ("Rendered By", "Verified By")
NOTEBOOK_LABEL = "Notebook Reference"

RECEIVING_LOG_COLUMNS = (
    "#",
    "Rec. DATE",
    "Rec. Time",
    "Product Code",
    "PRODUCT NAME",
    "TYPE",
    "LOT#",
    "CLIENT",
    "EXP DATE",
    "QTY RECD",
)


@dataclass(frozen=True)
class ReceivingLog:
    """A metadata value read from the Sample's Receiving Log row, by column."""

    column: str


@dataclass(frozen=True)
class Samples:
    """A metadata value read from the Sample's `samples` row, by column."""

    column: str


@dataclass(frozen=True)
class Form:
    """A value of the report form itself, not of the Sample."""

    value: str


Source = ReceivingLog | Samples | Form

CLIENT: tuple[tuple[str, Source], ...] = (
    ("Client", ReceivingLog("CLIENT")),
    ("Attention", Samples("report_to_contact")),
    ("Address", Samples("report_to_address")),
    ("", Samples("report_to_city")),
)
METADATA: tuple[tuple[str, Source], ...] = (
    ("Lab Reference", ReceivingLog("#")),
    ("Product", ReceivingLog("PRODUCT NAME")),
    ("Lot", ReceivingLog("LOT#")),
    ("Product Code", ReceivingLog("Product Code")),
    ("Product Type", ReceivingLog("TYPE")),
    ("Document Revision", Form(DOCUMENT_REVISION)),
    ("Date Received", ReceivingLog("Rec. DATE")),
    ("Time Received", ReceivingLog("Rec. Time")),
    ("Received By", Samples("received_by")),
)

Value = str | date | time
Result = tuple[str, str, str, float | str, str, str | None]


class ReportError(Exception):
    """The inputs are not ones this Script can render a complete report from."""


@dataclass(frozen=True)
class Report:
    """The one structure both renderings draw."""

    client: tuple[tuple[str, Value], ...]
    metadata: tuple[tuple[str, Value], ...]
    results: tuple[Result, ...]
    signatures: tuple[tuple[str, str], ...]
    notebook: str | None

    @property
    def lab_reference(self) -> str:
        return str(dict(self.metadata)["Lab Reference"])


@dataclass(frozen=True)
class Inputs:
    record: Path
    record_id: str
    receiving_log: Path
    weighings: Path
    database: Path


def inputs_of(inputs: Path) -> Inputs:
    manifest = json.loads((inputs / "manifest.json").read_text())

    def only(artifact_type: str) -> dict:
        found = [e for e in manifest if e["artifact_type"] == artifact_type]
        if len(found) != 1:
            raise ReportError(
                f"expected exactly one {artifact_type} input, got {len(found)}"
            )
        return found[0]

    def path(entry: dict) -> Path:
        return inputs / entry["artifact_id"] / entry["filename"]

    record = only("VerificationRecord")
    return Inputs(
        path(record),
        record["artifact_id"],
        path(only("ReceivingLogRow")),
        path(only("Other")),
        path(only("LimsDatabase")),
    )


def read_notebook_reference(path: Path) -> str | None:
    """The notebook reference every weighing names, or None when none names one."""
    try:
        entries = json.loads(path.read_text())
    except json.JSONDecodeError as error:
        raise ReportError(f"the weighings are not JSON: {error}") from error
    if not isinstance(entries, list) or not all(isinstance(e, dict) for e in entries):
        raise ReportError("the weighings are not a JSON array of objects")
    named = {entry.get("notebook_reference") for entry in entries}
    if named <= {None}:
        return None
    if len(named) > 1 or not all(isinstance(n, str) and n for n in named):
        raise ReportError(
            "the weighings do not all name one notebook reference:"
            f" {sorted(map(str, named))}"
        )
    return named.pop()


def text_of(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def read_receiving_log(path: Path) -> dict[str, dict[str, object]]:
    """Each row of the Receiving Log sheet, by its `#`, the Lab Reference."""
    ws = load_workbook(path).active
    rows = list(ws.iter_rows(values_only=True))
    if not rows or tuple(rows[0][: len(RECEIVING_LOG_COLUMNS)]) != (
        RECEIVING_LOG_COLUMNS
    ):
        raise ReportError("the Receiving Log row has no row of its columns")
    return {
        text_of(row[0]): dict(zip(RECEIVING_LOG_COLUMNS, row, strict=False))
        for row in rows[1:]
    }


@dataclass(frozen=True)
class Verified:
    work_order: str
    mean: float | None
    verdict: str
    nitrosamines: list[dict]


def read_record(path: Path) -> Verified:
    """The record's Work Order, recomputed mean, verdict and nitrosamine rows."""
    try:
        record = json.loads(path.read_text())
    except json.JSONDecodeError as error:
        raise ReportError(f"the Verification Record is not JSON: {error}") from error
    if record.get("schema") != "verification_record/1":
        raise ReportError("the Verification Record is not verification_record/1")
    work_order, mean = record.get("work_order"), record.get("mean") or {}
    value = mean.get("recomputed_g") if isinstance(mean, dict) else None
    if not isinstance(work_order, str) or not work_order:
        raise ReportError("the Verification Record names no Work Order")
    if value is not None and (
        not isinstance(value, int | float) or isinstance(value, bool)
    ):
        raise ReportError("the Verification Record's mean is not a number")
    nitrosamines = record.get("nitrosamines", [])
    if not isinstance(nitrosamines, list):
        raise ReportError("the Verification Record's nitrosamines is not a list")
    return Verified(
        work_order,
        None if value is None else float(value),
        str(record.get("verdict")),
        nitrosamines,
    )


def nitrosamine_row(entry: dict) -> Result:
    """One analyte's row of the results table, from the record's entry."""
    limit = entry["limit"]
    content = entry["content_ppm"]["recomputed"]
    intake = (entry["intake_ng_day"] or {}).get("recomputed")
    unit = limit.partition(" ")[2]
    if unit == "ng/day" and intake is not None:
        value = intake
    elif unit == "ng/g":
        value = None if content is None else content * 1000
    else:
        value, unit = content, "ppm"
    if not isinstance(value, int | float):
        raise ReportError(f"the record has no recomputed {entry['analyte']} {unit}")
    if not limit:
        specification = SPECIFICATION
    elif limit.startswith(UNREAD):
        specification = limit.removeprefix(UNREAD)
    else:
        specification = f"NMT {limit}"
    return (
        f"{NITROSAMINE_IMPURITIES}: {entry['analyte']}",
        NITROSAMINE_METHOD,
        specification,
        BELOW_REPORTING_LIMIT if entry["below_reporting_limit"] else round(value, 4),
        unit,
        entry["judgement"],
    )


def build_report(
    db: sqlite3.Connection,
    receiving_log: dict[str, dict[str, object]],
    record: Path,
    record_id: str,
    notebook: str | None,
) -> Report:
    verified = read_record(record)
    work_order = verified.work_order
    tests = db.execute(
        "SELECT lab_reference, test FROM work_orders WHERE work_order = ?"
        " ORDER BY lab_reference",
        (work_order,),
    ).fetchall()
    on_log = [ref for ref, _ in tests if ref in receiving_log]
    if not on_log:
        raise ReportError(
            f"no Sample of Work Order {work_order} is on the Receiving Log row"
        )
    ref = on_log[0]
    sample = db.execute(
        "SELECT * FROM samples WHERE lab_reference = ?", (ref,)
    ).fetchone()
    if sample is None:
        raise ReportError(f"the carried database has no samples row for {ref}")
    results: list[Result] = []
    for test in sorted({t for r, t in tests if r == ref}):
        if test == NITROSAMINE_IMPURITIES:
            results += [
                nitrosamine_row(entry)
                for entry in verified.nitrosamines
                if entry.get("container") == sample["container"]
            ]
            continue
        unit = db.execute(
            "SELECT unit FROM results WHERE lab_reference = ? AND test = ?",
            (ref, test),
        ).fetchone()
        if unit is None or not unit["unit"]:
            raise ReportError(f"the carried database has no {test} result for {ref}")
        if verified.mean is None:
            raise ReportError("the Verification Record has no recomputed mean")
        results.append(
            (test, METHOD, SPECIFICATION, round(verified.mean, 4), unit["unit"], None)
        )
    if not results:
        raise ReportError(f"the Verification Record has no result for {ref}")
    row = receiving_log[ref]

    def value(label: str, source: Source) -> Value:
        match source:
            case ReceivingLog(column):
                found, where = row.get(column), f"the Receiving Log row of {ref}"
            case Samples(column):
                found, where = sample[column], f"the samples row of {ref}"
            case Form(form_value):
                return form_value
        if isinstance(found, datetime):
            found = found.date()
        if isinstance(found, date | time):
            return found
        if found is None or str(found).strip() == "":
            name = label or "Address, second line"
            raise ReportError(f"{name}: {where} has no {source.column}")
        return str(found)

    return Report(
        client=tuple((label, value(label, s)) for label, s in CLIENT),
        metadata=tuple((label, value(label, s)) for label, s in METADATA),
        results=tuple(results),
        signatures=(
            ("Rendered By", "render_roa"),
            (
                "Verified By",
                f"Verification Record {record_id[:12]}, verdict {verified.verdict}",
            ),
        ),
        notebook=notebook,
    )


def write_xlsx(report: Report, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws["A1"] = LAB_NAME
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = LAB_ADDRESS
    ws["A3"] = LAB_CONTACT
    ws["A5"] = TITLE
    ws["A5"].font = Font(size=12, bold=True)
    for row, (label, value) in enumerate(report.client, start=7):
        ws.cell(row, 1, label)
        ws.cell(row, 2, value)
    for row, (label, value) in enumerate(report.metadata, start=12):
        ws.cell(row, 1, label).font = Font(bold=True)
        cell = ws.cell(row, 2, value)
        if isinstance(value, date):
            cell.number_format = "mm/dd/yyyy"
        elif isinstance(value, time):
            cell.number_format = "hh:mm"
    results_header = 22
    for col, title in enumerate(RESULT_COLUMNS, start=1):
        ws.cell(results_header, col, title).font = Font(bold=True)
    for row, result in enumerate(report.results, start=results_header + 1):
        for col, value in enumerate(result, start=1):
            ws.cell(row, col, value)
    signed = results_header + len(report.results) + 3
    for i, (role, who) in enumerate(report.signatures):
        col = 1 + 2 * i
        ws.cell(signed, col, role).font = Font(bold=True)
        ws.cell(signed + 1, col, who)
        ws.cell(signed + 2, col, "Signature: ______________")
        ws.cell(signed + 3, col, "Date: ______________")
    if report.notebook is not None:
        ws.cell(signed + 5, 1, NOTEBOOK_LABEL).font = Font(bold=True)
        ws.cell(signed + 5, 2, report.notebook)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 44
    wb.save(path)


def write_pdf(report: Report, path: Path) -> None:
    canvas = Canvas(str(path), pagesize=letter, pageCompression=0, invariant=1)
    y = 740

    def line(text: str, x: float = 72, font: str = "Helvetica", size: int = 10):
        canvas.setFont(font, size)
        canvas.drawString(x, y, text)

    def down(by: int) -> None:
        nonlocal y
        y -= by

    line(LAB_NAME, font="Helvetica-Bold", size=14)
    for text in (LAB_ADDRESS, LAB_CONTACT):
        down(14)
        line(text)
    down(28)
    line(TITLE, font="Helvetica-Bold", size=12)
    down(10)
    for block in (report.client, report.metadata):
        down(14)
        for label, value in block:
            down(14)
            line(label, font="Helvetica-Bold")
            line(text_of(value), x=200)
    down(28)
    xs = (72, 262, 330, 420, 480, 530)
    for x, title in zip(xs, RESULT_COLUMNS, strict=True):
        line(title, x=x, font="Helvetica-Bold")
    for result in report.results:
        down(14)
        for x, value in zip(xs, result, strict=True):
            line(f"{value:.4f}" if isinstance(value, float) else value or "", x=x)
    down(42)
    top = y
    for i, (role, who) in enumerate(report.signatures):
        x, y = 72 + 270 * i, top
        line(role, x=x, font="Helvetica-Bold")
        for text in (who, "Signature: ______________", "Date: ______________"):
            down(14)
            line(text, x=x)
    if report.notebook is not None:
        down(28)
        line(NOTEBOOK_LABEL, font="Helvetica-Bold")
        line(report.notebook, x=200)
    canvas.showPage()
    canvas.save()


def read_back(path: Path) -> list[list]:
    """The saved sheet, row 1 first: numbers as numbers, dates and times as text."""
    ws = load_workbook(path).active
    return [
        [v if v is None or isinstance(v, int | float) else text_of(v) for v in row]
        for row in ws.iter_rows(values_only=True)
    ]


def render_roa(inputs: Path, outputs: Path) -> Report:
    given = inputs_of(inputs)
    receiving_log = read_receiving_log(given.receiving_log)

    with tempfile.TemporaryDirectory() as work:
        out = Path(work)
        database = out / "lims.sqlite"
        shutil.copyfile(given.database, database)
        with closing(sqlite3.connect(database)) as db:
            db.row_factory = sqlite3.Row
            (version,) = db.execute("PRAGMA user_version").fetchone()
            if version != SCHEMA_VERSION:
                raise ReportError(
                    f"the carried database has schema version {version};"
                    f" this Script writes version {SCHEMA_VERSION}"
                )
            report = build_report(
                db,
                receiving_log,
                given.record,
                given.record_id,
                read_notebook_reference(given.weighings),
            )
            write_xlsx(report, out / "report_of_analysis.xlsx")
            write_pdf(report, out / "report_of_analysis.pdf")
            pdf_id = hashlib.sha256(
                (out / "report_of_analysis.pdf").read_bytes()
            ).hexdigest()
            db.execute("PRAGMA foreign_keys = ON")
            db.execute(
                "INSERT INTO reports (lab_reference, issued_date, document_revision,"
                " source_artifact) VALUES (?, ?, ?, ?)",
                (
                    report.lab_reference,
                    datetime.now(UTC).date().isoformat(),
                    DOCUMENT_REVISION,
                    pdf_id,
                ),
            )
            db.commit()

        read = {
            "schema": "report_of_analysis/1",
            "receiving_log_row": {
                column: text_of(value)
                for column, value in receiving_log[report.lab_reference].items()
            },
            "sheet": read_back(out / "report_of_analysis.xlsx"),
        }
        (out / "report_of_analysis.json").write_text(json.dumps(read, indent=2) + "\n")
        manifest = {
            "report_of_analysis.xlsx": {
                "artifact_type": "ReportOfAnalysis",
                "media_type": XLSX,
            },
            "report_of_analysis.pdf": {
                "artifact_type": "ReportOfAnalysis",
                "media_type": "application/pdf",
            },
            "lims.sqlite": {
                "artifact_type": "LimsDatabase",
                "media_type": "application/vnd.sqlite3",
            },
            "report_of_analysis.json": {
                "artifact_type": "Other",
                "media_type": "application/json",
            },
        }
        (out / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
        for name in [*manifest, "manifest.json"]:
            shutil.copyfile(out / name, outputs / name)
    return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--inputs", type=Path, required=True)
    parser.add_argument("--outputs", type=Path, required=True)
    args = parser.parse_args(argv)

    try:
        report = render_roa(args.inputs, args.outputs)
    except (ReportError, KeyError, sqlite3.Error) as error:
        print(f"render_roa refused its inputs: {error!r}", file=sys.stderr)
        return 1
    print(f"rendered the Report of Analysis of {report.lab_reference}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
