"""Read a Message reporting results into the weighings stage 4 reads.

Sandbox Script convention: run as
    python /script/main.py --inputs /inputs --outputs /outputs
with /inputs/manifest.json listing the inputs and each input at
/inputs/<artifact_id>/<filename>. The inputs must hold exactly one Message,
stored as `Other` JSON in the shape `store_message` writes:

    {"seq": 31, "room": "RM-lab", "said_by": "Operator",
     "ts": "2026-10-13T11:02:00.000Z", "body": "Lab Ref 26-0412, ..."}

and exactly one `WorkOrder`, the sheet `create_work_order` writes, whose
rows list one Lab Reference per container, each with its Lot.

The body is free text. Each value is read by a labelled pattern, never
guessed:

- the notebook reference, after "notebook", up to the next comma,
  semicolon, line end or full stop that ends a sentence;
- the Lot, after "Lot" (a value holding a digit);
- per container, "Container <n>: <w> g for 20 capsules, <t> g taken",
  in either order within the container's part of the body.

The Script refuses (exit 1, the reason on stderr, no outputs) a Message
that names no notebook reference, names a Lot the Work Order does not
list, names a container twice, gives a container no weight of 20
capsules or no weight taken, or names a number of containers other than
the number of Samples on the Work Order. A Lab Reference the Message
names is not read: the Work Order input says which Samples it is about.

Output `/outputs/weighings.json`, typed `Other`: one JSON array of
`{"container": n, "weight_of_20_g": g, "weight_taken_g": g,
"notebook_reference": "<ref>"}` in container order, the shape synth's
`weighings.json` has plus the notebook reference, so stages 4 and 5 read
it as they read the Fixture's weighings.
"""

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

NOTEBOOK = re.compile(r"\bnotebook\s*[:#]?\s*(.+?)\s*(?=[,;\n]|\.(?:\s|$)|$)", re.I)
LOT = re.compile(r"\bLot\s*(?:#|No\.?|:)?\s*([A-Za-z0-9-]*\d[A-Za-z0-9-]*)", re.I)
CONTAINER = re.compile(r"\bContainer\s+(\d+)\s*:", re.I)
GRAMS = r"(\d+(?:\.\d+)?)\s*g\b"
OF_20 = re.compile(rf"{GRAMS}\s+for\s+20\s+capsules", re.I)
TAKEN = re.compile(rf"{GRAMS}\s+taken", re.I)
WORK_ORDER_COLUMNS = ("Lab Reference", "Test", "Lot")


class ResultsMessageError(Exception):
    """The inputs are not a results Message and Work Order this Script can read."""


def only(manifest: list[dict], artifact_type: str) -> dict:
    found = [e for e in manifest if e["artifact_type"] == artifact_type]
    if len(found) != 1:
        raise ResultsMessageError(
            f"expected exactly one {artifact_type} input, got {len(found)}"
        )
    return found[0]


def message_of(path: Path) -> dict:
    try:
        message = json.loads(path.read_text())
    except json.JSONDecodeError as error:
        raise ResultsMessageError(f"the Message is not JSON: {error}") from error
    keys = {"seq", "room", "said_by", "ts", "body"}
    if not isinstance(message, dict) or not keys <= message.keys():
        raise ResultsMessageError(
            f"the Other input is not a Message: it needs {', '.join(sorted(keys))}"
        )
    return message


def work_order_of(path: Path) -> tuple[list[str], int]:
    """The Work Order's Lots, in row order, and the number of Samples it lists."""
    rows = list(load_workbook(path).active.iter_rows(values_only=True))
    starts = [i for i, row in enumerate(rows) if tuple(row[:3]) == WORK_ORDER_COLUMNS]
    planned = [row for row in rows[starts[0] + 1 :] if row[0]] if starts else []
    if not planned or not all(row[2] for row in planned):
        raise ResultsMessageError("the Work Order has no Lab Reference rows with a Lot")
    lots = list(dict.fromkeys(str(row[2]) for row in planned))
    return lots, len({row[0] for row in planned})


def one(label: str, values: list[str]) -> str:
    distinct = sorted(set(values))
    if len(distinct) > 1:
        raise ResultsMessageError(
            f"the Message names {len(distinct)} {label}s: {', '.join(distinct)}"
        )
    return distinct[0] if distinct else ""


def weighings_of(body: str) -> dict[int, tuple[float, float]]:
    """Each container's weight of 20 capsules and weight taken, by number."""
    marks = list(CONTAINER.finditer(body))
    weighed: dict[int, tuple[float, float]] = {}
    for i, mark in enumerate(marks):
        container = int(mark.group(1))
        end = marks[i + 1].start() if i + 1 < len(marks) else len(body)
        part = body[mark.end() : end]
        if container in weighed:
            raise ResultsMessageError(f"the Message names container {container} twice")
        of_20 = [float(m.group(1)) for m in OF_20.finditer(part)]
        taken = [float(m.group(1)) for m in TAKEN.finditer(part)]
        lacking = [
            what
            for what, found in (
                ("weight of 20 capsules", of_20),
                ("weight taken", taken),
            )
            if len(found) != 1
        ]
        if lacking:
            raise ResultsMessageError(
                f"container {container} has no single {' and no single '.join(lacking)}"
            )
        weighed[container] = (of_20[0], taken[0])
    return weighed


def parse(message: dict, lots: list[str], samples: int) -> list[dict]:
    lot = ", ".join(lots)
    body = message["body"]
    notebook = one("notebook reference", [m.group(1) for m in NOTEBOOK.finditer(body)])
    if not notebook:
        raise ResultsMessageError("the Message names no notebook reference")
    named = one("Lot", [m.group(1) for m in LOT.finditer(body)])
    if named and named not in lots:
        raise ResultsMessageError(
            f"the Message names Lot {named}, and the Work Order is for Lot {lot}"
        )
    weighed = weighings_of(body)
    if len(weighed) != samples:
        raise ResultsMessageError(
            f"the Message weighs {len(weighed)} containers, and Lot {lot}'s"
            f" Work Order has {samples} Samples"
        )
    return [
        {
            "container": container,
            "weight_of_20_g": of_20,
            "weight_taken_g": taken,
            "notebook_reference": notebook,
        }
        for container, (of_20, taken) in sorted(weighed.items())
    ]


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--inputs", type=Path, required=True)
    parser.add_argument("--outputs", type=Path, required=True)
    args = parser.parse_args(argv)

    try:
        manifest = json.loads((args.inputs / "manifest.json").read_text())

        def path(entry: dict) -> Path:
            return args.inputs / entry["artifact_id"] / entry["filename"]

        message = message_of(path(only(manifest, "Other")))
        lots, samples = work_order_of(path(only(manifest, "WorkOrder")))
        weighings = parse(message, lots, samples)
    except (ResultsMessageError, KeyError) as error:
        print(f"parse_results_message refused its inputs: {error}", file=sys.stderr)
        return 1

    (args.outputs / "weighings.json").write_text(json.dumps(weighings, indent=2) + "\n")
    (args.outputs / "manifest.json").write_text(
        json.dumps(
            {
                "weighings.json": {
                    "artifact_type": "Other",
                    "media_type": "application/json",
                }
            },
            indent=2,
        )
        + "\n"
    )
    print(f"read {len(weighings)} weighings for Lot {', '.join(lots)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
