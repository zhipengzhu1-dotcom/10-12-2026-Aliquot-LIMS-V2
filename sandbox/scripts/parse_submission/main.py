"""Stage 1 Script: read a Submission Form (docx) into a structured submission record.

Sandbox Script convention: run as
    python /script/main.py --inputs /inputs --outputs /outputs
with /inputs/manifest.json listing the inputs and each input at
/inputs/<artifact_id>/<filename>. The inputs must hold exactly one docx.

The Script reads two layouts into the same record and tells them apart by the
document's structure, never by the form's code. A form with a table labelled
`COMPANY NAME & ADDRESS` is in the lab's real layout; any other is in the
synthetic layout. In both, tables are found by their labels, not their
position, and columns by their header text.

The synthetic layout's labels are the ones `sandbox/synth/submission_form.py`
writes. The real layout's are the labels the lab prints, which the scrubbed
copy in `fixtures/real-layout/` kept, and the values sit where
`sandbox/synth/real_layout_form.py` writes them:

- The date submitted and the submitter's contact follow their labels in the
  first table's top row.
- `COMPANY NAME & ADDRESS` holds the submitter's company, street and city, in
  that order, one per line, and its phone and email on lines of their own.
  `BILL TO ADDRESS (if different)` holds the bill-to party the same way, its
  contact on a line starting `Attn:`. A blank bill-to block means the
  submitter's company, as its label says.
- `Attn`, `Phone #` and `Email ID` give the report-to contact, phone and email.
  The real form has no cell for the report-to company, street or city, so they
  are the submitter's.
- The Sample table's columns are `No.`, `SAMPLE DESCRIPTION`, `TYPE` (the
  Sample Type), `*TAT #`, `LOT NUMBER` and `PRODUCT CODE`, by the first line of
  their header.
- The lab-use table's date received, time received and receiver follow their
  labels, and its `Comments:` holds `<containers> x <Product Type>`, as the
  receiver writes the count and the package. The real form has no other cell
  for either.

Output `/outputs/submission.json` (keys sorted, two-space indent):

    {
      "schema": "submission/1",
      "date_submitted": "YYYY-MM-DD",
      "submitted_by": Party,
      "bill_to": Party,
      "report_to": Party,
      "sample": {
        "no", "sample_description", "sample_type", "tat_code",
        "tat_class", "lot", "product_code", "product_type",
        "no_of_containers"
      },
      "receipt": {"date_received": "YYYY-MM-DD", "time_received": "HH:MM",
                  "received_by"}
    }
    Party = {"company", "contact", "address", "city", "phone", "email"}

Every value is a string holding the cell's text with surrounding whitespace
removed; a blank cell, or a label the form lacks, gives "". Dates printed as
MM/DD/YYYY become ISO 8601 YYYY-MM-DD; any other date text is kept as
printed. `sample` is the first Sample row. `tat_class` is derived from
`tat_code` with the TAT code table below ("" for a code not in the table).
The Script reports what the form says and judges nothing: the stage 1
Evaluation decides whether the record is complete.

`/outputs/manifest.json` types submission.json as
`artifact_type: "SubmissionRecord"`.
"""

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.table import Table

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
SCHEMA = "submission/1"

# Copied from sandbox/synth/tat.py: the container mounts only this directory,
# so the Script cannot import `synth`.
TAT_CODES = {"1": "Regular", "2": "Rush1to2Days", "3": "Rush3to4Days"}

PARTIES = {
    "Submitted By": "submitted_by",
    "Bill To": "bill_to",
    "Report To": "report_to",
}
PARTY_FIELDS = {
    "Company": "company",
    "Contact": "contact",
    "Address": "address",
    "City": "city",
    "Phone": "phone",
    "Email": "email",
}
SAMPLE_COLUMNS = {
    "No.": "no",
    "Sample Description": "sample_description",
    "Sample Type": "sample_type",
    "TAT Code": "tat_code",
    "Lot": "lot",
    "Product Code": "product_code",
    "Product Type": "product_type",
    "No. of Containers": "no_of_containers",
}
RECEIPT_FIELDS = {
    "Date Received": "date_received",
    "Time Received": "time_received",
    "Received By (initials)": "received_by",
}
DATE_SUBMITTED = re.compile(r"^\s*Date Submitted:\s*(.*)$")

REAL_COMPANY = "COMPANY NAME & ADDRESS"
REAL_BILL_TO = "BILL TO ADDRESS"
REAL_SAMPLE_COLUMNS = {
    "No.": "no",
    "SAMPLE DESCRIPTION": "sample_description",
    "TYPE": "sample_type",
    "*TAT #": "tat_code",
    "LOT NUMBER": "lot",
    "PRODUCT CODE": "product_code",
}
REAL_SEND_TO = {"Attn": "contact", "Phone #": "phone", "Email ID": "email"}
REAL_RECEIPT = {
    "Date Sample(s) received:": "date_received",
    "Time received:": "time_received",
    "Received By:": "received_by",
}
PHONE = re.compile(r"^[\d\s()+.-]{7,}$")
COUNT_AND_PACKAGE = re.compile(r"^(\d+)\s*x\s*(.+)$", re.IGNORECASE)


class FormError(Exception):
    """The inputs are not a Submission Form this Script can read."""


def iso_date(text: str) -> str:
    try:
        return datetime.strptime(text, "%m/%d/%Y").date().isoformat()
    except ValueError:
        return text


def cells(row) -> list[str]:
    return [cell.text.strip() for cell in row.cells]


def find_table(tables: list[Table], label: str) -> Table | None:
    """The first table whose first row or first column holds `label`."""
    for table in tables:
        if label in cells(table.rows[0]) or label in [
            cells(row)[0] for row in table.rows
        ]:
            return table
    return None


def parties(table: Table | None) -> dict[str, dict[str, str]]:
    record = {key: dict.fromkeys(PARTY_FIELDS.values(), "") for key in PARTIES.values()}
    if table is None:
        return record
    header = cells(table.rows[0])
    for row in table.rows[1:]:
        values = cells(row)
        field = PARTY_FIELDS.get(values[0])
        if field is None:
            continue
        for col, title in enumerate(header):
            if title in PARTIES and col < len(values):
                record[PARTIES[title]][field] = values[col]
    return record


def first_row(table: Table | None, columns: dict[str, str]) -> dict[str, str]:
    record = dict.fromkeys(columns.values(), "")
    if table is None or len(table.rows) < 2:
        return record
    header = cells(table.rows[0])
    values = cells(table.rows[1])
    for col, title in enumerate(header):
        if title in columns and col < len(values):
            record[columns[title]] = values[col]
    return record


def parse_synthetic(doc, tables: list[Table]) -> dict:
    date_submitted = ""
    for paragraph in doc.paragraphs:
        match = DATE_SUBMITTED.match(paragraph.text)
        if match:
            date_submitted = iso_date(match.group(1).strip())
            break

    sample = first_row(find_table(tables, "Sample Description"), SAMPLE_COLUMNS)
    receipt = first_row(find_table(tables, "Date Received"), RECEIPT_FIELDS)
    return {
        "date_submitted": date_submitted,
        **parties(find_table(tables, "Submitted By")),
        "sample": sample,
        "receipt": receipt,
    }


def after(text: str, label: str) -> str | None:
    """The text after `label` when `text` starts with it, ignoring case."""
    text = text.strip()
    if not text.lower().startswith(label.lower()):
        return None
    return text[len(label) :].strip()


def labelled(rows, labels: dict[str, str]) -> dict[str, str]:
    """Each label's value, from the cells that start with the label."""
    values = dict.fromkeys(labels.values(), "")
    for row in rows:
        for text in cells(row):
            for label, key in labels.items():
                value = after(text, label)
                if value is not None:
                    values[key] = value
    return values


def party_block(text: str) -> dict[str, str]:
    """A party from a free-text address block, one item per line."""
    party = dict.fromkeys(PARTY_FIELDS.values(), "")
    lines = []
    for line in (line.strip() for line in text.splitlines()):
        contact = after(line, "Attn:")
        if not line:
            continue
        if contact is not None:
            party["contact"] = contact
        elif "@" in line:
            party["email"] = line
        elif PHONE.match(line):
            party["phone"] = line
        else:
            lines.append(line)
    for key, line in zip(("company", "address", "city"), lines, strict=False):
        party[key] = line
    return party


def real_table(tables: list[Table], label: str) -> Table | None:
    """The first table with a cell that starts with `label`."""
    for table in tables:
        for row in table.rows:
            if any(text.startswith(label) for text in cells(row)):
                return table
    return None


def real_parties(table: Table | None) -> dict:
    if table is None or len(table.rows) < 3:
        return {
            key: dict.fromkeys(PARTY_FIELDS.values(), "") for key in PARTIES.values()
        }
    top = labelled(table.rows[:1], {"Date submitted:": "date", "Submitted By:": "by"})
    headings = cells(table.rows[1])
    blocks = cells(table.rows[2])

    def block(heading: str) -> str:
        for col, title in enumerate(headings):
            if title.startswith(heading) and col < len(blocks):
                return blocks[col]
        return ""

    submitted_by = {**party_block(block(REAL_COMPANY)), "contact": top["by"]}
    bill_text = block(REAL_BILL_TO)
    bill_to = party_block(bill_text) if bill_text.strip() else dict(submitted_by)
    send_to = dict.fromkeys(REAL_SEND_TO.values(), "")
    for row in table.rows[2:]:
        values = cells(row)
        for col, text in enumerate(values[:-1]):
            if text in REAL_SEND_TO:
                send_to[REAL_SEND_TO[text]] = values[col + 1]
    report_to = {
        **{key: submitted_by[key] for key in ("company", "address", "city")},
        **send_to,
    }
    return {
        "date_submitted": iso_date(top["date"]),
        "submitted_by": submitted_by,
        "bill_to": bill_to,
        "report_to": report_to,
    }


def parse_real(tables: list[Table]) -> dict:
    samples = real_table(tables, "SAMPLE DESCRIPTION")
    sample = dict.fromkeys(SAMPLE_COLUMNS.values(), "")
    if samples is not None and len(samples.rows) >= 2:
        header = [
            text.splitlines()[0].strip() if text else ""
            for text in cells(samples.rows[0])
        ]
        values = cells(samples.rows[1])
        for col, title in enumerate(header):
            if title in REAL_SAMPLE_COLUMNS and col < len(values):
                sample[REAL_SAMPLE_COLUMNS[title]] = values[col]

    lab_use = real_table(tables, "Date Sample(s) received:")
    rows = lab_use.rows if lab_use is not None else []
    receipt = labelled(rows, REAL_RECEIPT)
    count = COUNT_AND_PACKAGE.match(labelled(rows, {"Comments:": "c"})["c"])
    if count:
        sample["no_of_containers"] = count.group(1)
        sample["product_type"] = count.group(2).strip()
    return {
        **real_parties(real_table(tables, REAL_COMPANY)),
        "sample": sample,
        "receipt": receipt,
    }


def parse(form: Path) -> dict:
    doc = Document(str(form))
    tables = list(doc.tables)
    if real_table(tables, REAL_COMPANY) is not None:
        record = parse_real(tables)
    else:
        record = parse_synthetic(doc, tables)
    record["sample"]["tat_class"] = TAT_CODES.get(record["sample"]["tat_code"], "")
    record["receipt"]["date_received"] = iso_date(record["receipt"]["date_received"])
    return {"schema": SCHEMA, **record}


def submission_form(inputs: Path) -> Path:
    manifest = json.loads((inputs / "manifest.json").read_text())
    forms = [
        entry
        for entry in manifest
        if entry["media_type"] == DOCX or entry["filename"].lower().endswith(".docx")
    ]
    if len(forms) != 1:
        raise FormError(
            f"expected exactly one Submission Form (docx) input, got {len(forms)}"
        )
    return inputs / forms[0]["artifact_id"] / forms[0]["filename"]


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--inputs", type=Path, required=True)
    parser.add_argument("--outputs", type=Path, required=True)
    args = parser.parse_args(argv)

    try:
        form = submission_form(args.inputs)
        record = parse(form)
    except FormError as error:
        print(error, file=sys.stderr)
        return 1

    (args.outputs / "submission.json").write_text(
        json.dumps(record, indent=2, sort_keys=True) + "\n"
    )
    (args.outputs / "manifest.json").write_text(
        json.dumps(
            {
                "submission.json": {
                    "artifact_type": "SubmissionRecord",
                    "media_type": "application/json",
                }
            },
            indent=2,
        )
        + "\n"
    )
    print(f"parsed {form.name}: {record['sample']['lot'] or 'no Lot'}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
