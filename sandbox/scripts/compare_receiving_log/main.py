"""Comparison Script: the Shadow's rows beside the rows the lab entered by hand.

Sandbox Script convention: run as
    python /script/main.py --inputs /inputs --outputs /outputs
with /inputs/manifest.json listing the inputs and each input at
/inputs/<artifact_id>/<filename>. The inputs must hold exactly one
`ShadowReceivingLog` (the Shadow stage 2 wrote), the `shadow.json` written
beside it (typed `Other`, naming the Shadow's new rows), and exactly one
`ReceivingLog`, a fresh read of the lab's workbook.

The Shadow's rows are the new rows of its year's sheet. Each is paired with a
row of the lab's sheet of the same name that has the same `Rec. DATE` and the
same `LOT#` or the same `Product Code`; among several, the one that agrees on
the most compared fields, then the first. The Lab Reference cannot pair them,
because the lab numbers its Samples its own way. Requiring only one of Lot and
Product Code lets a mistyped Lot or Product Code show as a differing field,
not as a missing row on both sides.

Compared fields, by position in the year's sheet (the lab's headers are not
kept in the Fixtures): C `Rec. Time`, D `Product Code`, E `PRODUCT NAME`,
F `TYPE`, G `LOT#`, H `CLIENT`, I `EXP DATE`, J `QTY RECD`, K `SENT BY`,
L `RECD BY`. Values are compared as text: dates as YYYY-MM-DD, times as HH:MM,
whole numbers without a decimal point, surrounding space removed.

Outputs, both typed `Comparison`:

- `comparison.json`: `{"schema": "comparison/1", "sheet", "rows": [...],
  "counts": {...}}`. Each row has a `status`: `matched`, `differs`,
  `missing_in_lab` (a Shadow row the lab has not entered) or
  `missing_in_shadow` (a lab row for the same Samples with no Shadow row);
  `shadow_row` and `lab_row` (null when missing), the Shadow's Lab Reference,
  and `differences`, one `{"column", "field", "shadow", "lab"}` per field.
- `comparison.txt`: the same as a table to read.

The Script refuses inputs it cannot read (exit 1, the reason on stderr) and
judges nothing: the Operator reads the Comparison.
"""

import argparse
import json
import sys
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

YEAR_FIRST_ROW = 4
YEAR_COLUMNS = 20
REC_DATE, PRODUCT_CODE, LOT, LAB_REFERENCE = "B", "D", "G", "M"
COMPARED = {
    "C": "Rec. Time",
    "D": "Product Code",
    "E": "PRODUCT NAME",
    "F": "TYPE",
    "G": "LOT#",
    "H": "CLIENT",
    "I": "EXP DATE",
    "J": "QTY RECD",
    "K": "SENT BY",
    "L": "RECD BY",
}
STATUS_TEXT = {
    "matched": "matched",
    "differs": "differs",
    "missing_in_lab": "missing in the lab's",
    "missing_in_shadow": "missing in the Shadow",
}

Row = dict[str, str]


class ComparisonError(Exception):
    """The inputs are not a Shadow and a Receiving Log this Script can read."""


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, ArrayFormula):
        return value.text or ""
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def row_of(ws, number: int) -> Row:
    return {
        get_column_letter(col): cell_text(ws.cell(number, col).value)
        for col in range(1, YEAR_COLUMNS + 1)
    }


def inputs_of(inputs: Path) -> tuple[Path, Path, Path]:
    manifest = json.loads((inputs / "manifest.json").read_text())

    def one(artifact_type: str, filename: str | None = None) -> Path:
        entries = [
            e
            for e in manifest
            if e["artifact_type"] == artifact_type
            and (filename is None or e["filename"] == filename)
        ]
        if len(entries) != 1:
            named = f" {filename}" if filename else ""
            raise ComparisonError(
                f"expected exactly one {artifact_type}{named} input, got {len(entries)}"
            )
        return inputs / entries[0]["artifact_id"] / entries[0]["filename"]

    return (
        one("ShadowReceivingLog"),
        one("Other", "shadow.json"),
        one("ReceivingLog"),
    )


def agreement(a: Row, b: Row) -> int:
    return sum(a.get(col) == b.get(col) for col in COMPARED)


def same_samples(shadow: Row, lab: Row) -> bool:
    return shadow[REC_DATE] == lab[REC_DATE] and (
        shadow[LOT] == lab[LOT] or shadow[PRODUCT_CODE] == lab[PRODUCT_CODE]
    )


def compare(shadow_rows: dict[int, Row], lab_rows: dict[int, Row]) -> list[dict]:
    unpaired = dict(lab_rows)
    results = []
    for number, shadow in shadow_rows.items():
        candidates = [n for n, lab in unpaired.items() if same_samples(shadow, lab)]
        if not candidates:
            results.append(
                {
                    "status": "missing_in_lab",
                    "shadow_row": number,
                    "lab_row": None,
                    "lab_reference": shadow[LAB_REFERENCE],
                    "differences": [],
                }
            )
            continue
        best = max(candidates, key=lambda n: (agreement(shadow, unpaired[n]), -n))
        lab = unpaired.pop(best)
        differences = [
            {"column": col, "field": field, "shadow": shadow[col], "lab": lab[col]}
            for col, field in COMPARED.items()
            if shadow[col] != lab[col]
        ]
        results.append(
            {
                "status": "differs" if differences else "matched",
                "shadow_row": number,
                "lab_row": best,
                "lab_reference": shadow[LAB_REFERENCE],
                "differences": differences,
            }
        )
    for number, lab in unpaired.items():
        if any(same_samples(shadow, lab) for shadow in shadow_rows.values()):
            results.append(
                {
                    "status": "missing_in_shadow",
                    "shadow_row": None,
                    "lab_row": number,
                    "lab_reference": "",
                    "differences": [],
                }
            )
    return results


def table(sheet: str, rows: list[dict], counts: dict[str, int]) -> str:
    header = (
        "Shadow row",
        "Lab row",
        "Lab Reference",
        "Status",
        "Field",
        "Shadow",
        "Lab",
    )
    lines = [header]
    for r in rows:
        first = (
            str(r["shadow_row"] or "-"),
            str(r["lab_row"] or "-"),
            r["lab_reference"] or "-",
            STATUS_TEXT[r["status"]],
        )
        if not r["differences"]:
            lines.append((*first, "", "", ""))
        for i, d in enumerate(r["differences"]):
            lead = first if i == 0 else ("", "", "", "")
            lines.append((*lead, d["field"], d["shadow"], d["lab"]))
    widths = [max(len(line[i]) for line in lines) for i in range(len(header))]
    text = [
        "  ".join(cell.ljust(w) for cell, w in zip(line, widths, strict=True)).rstrip()
        for line in lines
    ]
    summary = ", ".join(f"{counts[s]} {STATUS_TEXT[s]}" for s in STATUS_TEXT)
    return (
        "\n".join(
            [f"Comparison of the Shadow's sheet {sheet} with the lab's", *text, summary]
        )
        + "\n"
    )


def run(inputs: Path, outputs: Path) -> dict[str, int]:
    shadow_path, held_path, lab_path = inputs_of(inputs)
    held = json.loads(held_path.read_text())
    if held.get("schema") != "shadow/1":
        raise ComparisonError("shadow.json is not a shadow/1 read-back")
    sheet, numbers = held["year"]["sheet"], held["year"]["rows"]
    try:
        shadow_wb, lab_wb = load_workbook(shadow_path), load_workbook(lab_path)
    except Exception as error:
        raise ComparisonError(f"a workbook is unreadable: {error}") from error
    shadow_rows = {n: row_of(shadow_wb[sheet], n) for n in numbers}
    lab_rows: dict[int, Row] = {}
    if sheet in lab_wb.sheetnames:
        ws = lab_wb[sheet]
        for n in range(YEAR_FIRST_ROW, ws.max_row + 1):
            values = row_of(ws, n)
            if any(values.values()):
                lab_rows[n] = values

    rows = compare(shadow_rows, lab_rows)
    counts = {s: sum(r["status"] == s for r in rows) for s in STATUS_TEXT}
    report = {"schema": "comparison/1", "sheet": sheet, "rows": rows, "counts": counts}
    (outputs / "comparison.json").write_text(json.dumps(report, indent=2) + "\n")
    (outputs / "comparison.txt").write_text(table(sheet, rows, counts))
    manifest = {
        "comparison.json": {
            "artifact_type": "Comparison",
            "media_type": "application/json",
        },
        "comparison.txt": {"artifact_type": "Comparison", "media_type": "text/plain"},
    }
    (outputs / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    return counts


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--inputs", type=Path, required=True)
    parser.add_argument("--outputs", type=Path, required=True)
    args = parser.parse_args(argv)

    try:
        counts = run(args.inputs, args.outputs)
    except (ComparisonError, KeyError) as error:
        print(f"compare_receiving_log refused its inputs: {error!r}", file=sys.stderr)
        return 1
    print(", ".join(f"{n} {STATUS_TEXT[s]}" for s, n in counts.items()))
    return 0


if __name__ == "__main__":
    sys.exit(main())
