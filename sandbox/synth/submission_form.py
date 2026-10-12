"""The Client-facing Submission Form (docx).

Three tables, in document order:

1. Header table: a label column, then one column each for the submitter, the
   bill-to and the report-to party.
2. Sample table: 11 rows by 8 columns, a header row and ten Sample rows.
3. Lab-use-only receipt block: date, time and received-by initials.

The labels below are the contract `parse_submission` reads.
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt

from synth.lab import LAB_ADDRESS, LAB_CONTACT, LAB_NAME
from synth.record import Party, Record
from synth.reproducible import save_reproducibly

PARTY_COLUMNS = ("Submitted By", "Bill To", "Report To")
PARTY_FIELDS = ("Company", "Contact", "Address", "City", "Phone", "Email")

SAMPLE_COLUMNS = (
    "No.",
    "Sample Description",
    "Sample Type",
    "TAT Code",
    "Lot",
    "Product Code",
    "Product Type",
    "No. of Containers",
)
SAMPLE_ROWS = 10

RECEIPT_FIELDS = ("Date Received", "Time Received", "Received By (initials)")
LAB_USE_ONLY = "FOR LAB USE ONLY"

# Fields that must be filled for the form to be accepted, per table.
REQUIRED_PARTY_FIELDS: dict[str, tuple[str, ...]] = {
    "Submitted By": PARTY_FIELDS,
    "Bill To": ("Company", "Contact", "Address", "City"),
    "Report To": ("Company", "Contact", "Email"),
}
REQUIRED_SAMPLE_FIELDS = SAMPLE_COLUMNS[1:]
REQUIRED_RECEIPT_FIELDS = RECEIPT_FIELDS


def _party_values(party: Party) -> tuple[str, ...]:
    return (
        party.company,
        party.contact,
        party.street,
        party.city,
        party.phone,
        party.email,
    )


def write_submission_form(record: Record, path: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(9)

    doc.add_heading(LAB_NAME, level=1)
    doc.add_paragraph(f"{LAB_ADDRESS} · {LAB_CONTACT}")
    doc.add_heading("Sample Submission Form", level=2)
    doc.add_paragraph(f"Date Submitted: {record.submitted_on:%m/%d/%Y}")

    header = doc.add_table(rows=1 + len(PARTY_FIELDS), cols=1 + len(PARTY_COLUMNS))
    header.style = "Table Grid"
    for col, title in enumerate(PARTY_COLUMNS, start=1):
        header.cell(0, col).text = title
    parties = (record.submitter, record.bill_to, record.report_to)
    for row, field in enumerate(PARTY_FIELDS, start=1):
        header.cell(row, 0).text = field
        for col, party in enumerate(parties, start=1):
            header.cell(row, col).text = _party_values(party)[row - 1]

    doc.add_paragraph()
    samples = doc.add_table(rows=1 + SAMPLE_ROWS, cols=len(SAMPLE_COLUMNS))
    samples.style = "Table Grid"
    for col, title in enumerate(SAMPLE_COLUMNS):
        samples.cell(0, col).text = title
    for row in range(1, SAMPLE_ROWS + 1):
        samples.cell(row, 0).text = str(row)
    first_sample = (
        "1",
        record.product_name,
        record.sample_type,
        record.tat_code,
        record.lot,
        record.product_code,
        record.product_type,
        str(len(record.containers)),
    )
    for col, value in enumerate(first_sample):
        samples.cell(1, col).text = value

    doc.add_paragraph(
        "TAT codes: 1 = regular, 2 = rush in 1 to 2 days, 3 = rush in 3 to 4 days. "
        "Sample Types: FP, IP, RM, R&D."
    )

    doc.add_paragraph(LAB_USE_ONLY)
    receipt = doc.add_table(rows=2, cols=len(RECEIPT_FIELDS))
    receipt.style = "Table Grid"
    received = (
        f"{record.received_on:%m/%d/%Y}",
        f"{record.received_at:%H:%M}",
        record.received_by,
    )
    for col, (label, value) in enumerate(zip(RECEIPT_FIELDS, received, strict=True)):
        receipt.cell(0, col).text = label
        receipt.cell(1, col).text = value

    props = doc.core_properties
    props.title = "Sample Submission Form"
    props.author = record.submitter.contact
    props.last_modified_by = record.submitter.contact
    props.created = record.document_time
    props.modified = record.document_time
    props.last_printed = record.document_time
    props.revision = 1

    save_reproducibly(doc.save, path, record.document_time)
