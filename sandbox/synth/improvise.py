"""`python -m synth.improvise --log LOG --blank BLANK --out DIR`: forms from the log.

Fills the scrubbed blank Submission Form once for each submission the scrubbed
Receiving Log records, so a Comparison (ADR-0018) pairs the Shadow's rows with
rows the lab entered. Writes, per form, `<stem>.docx`, `<stem>.expected.json`
(the record `parse_submission` must read) and `<stem>.weighings.json`. The
same log and blank always give the same bytes.

From the log's year sheet, the rows the lab received on one date from one
Client make one submission, in the order the log holds them, and each row is
one Sample line on its form. The blank has ten Sample rows, so a longer
submission continues on another form, and it breaks there only between lines
of different Lots or Product Codes while it can: a Comparison reads one form
at a time, and pairs a Shadow row with the first lab row of its day, Lot and
Product Code that agrees best, so lines alike on two forms of one day would
pair with the same lab rows. Each line carries the row's PRODUCT
NAME, LOT# and Product Code, and its TYPE when that is a Sample Type code;
any other TYPE (a stability condition or a scrubbed label) is written as FP,
the only Sample Type stage 3 plans a Test for. The form's date and time
received, Client (the company line) and receiver's initials are the first
row's. Each line is one Sample in one container, as the lab logs one row per
line, so its `Comments:` entry is `1 x <unit>`, the unit of the row's QTY
RECD as the package. The weighings number the form's containers on from 1 in
line order, as stage 2 numbers them.

Everything the log does not hold is invented with one seeded generator: the
date submitted, contacts, streets, cities, phones and emails, the TAT code,
the Tests requested and limits, and the weighings. The TAT code, a report-to
contact other than the submitter, and a filled bill-to block rotate with the
form's place in the set, so the set has each kind.

Two forms are off on purpose. The earliest form's Lots are mistyped by one
character, and one more form lists an IP Sample the log does not hold.
"""

import argparse
import json
import random
import re
import string
from dataclasses import dataclass, replace
from datetime import date, datetime, time, timedelta
from itertools import groupby
from pathlib import Path

from openpyxl import load_workbook

from synth.real_layout_form import FilledForm, SampleLine, write_filled_form
from synth.record import (
    CLIENTS,
    FIRST_NAMES,
    LAST_NAMES,
    PRODUCTS,
    STREETS,
    TEST_REQUESTED,
    Party,
)
from synth.tat import TAT_CODES

SEED = 81
FIRST_ROW = 4  # the year sheet's headers are in row 3
FORM_LINES = 10  # the blank's Sample table
SAMPLE_TYPES = ("FP", "IP", "RM")
OTHER_TESTS = ("Assay", "Identification", "Water content")
QUANTITY = re.compile(r"^(\d+)\s+(.+)$")


@dataclass(frozen=True)
class LogRow:
    """One row of the year sheet, by the columns `log_receipt/shadow.py` names."""

    number: int
    received_on: date
    received_at: time
    product_code: str
    product_name: str
    type: str
    lot: str
    client: str
    quantity: str
    received_by: str


@dataclass(frozen=True)
class Improvised:
    stem: str
    form: FilledForm
    weighings: tuple[dict[str, float | int], ...]


def _text(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return "" if value is None else str(value).strip()


def read_log(path: Path) -> list[LogRow]:
    wb = load_workbook(path)
    year = next(ws for ws in wb.worksheets if ws.title.isdigit())
    rows = []
    for number in range(FIRST_ROW, year.max_row + 1):
        cells = [year.cell(number, col).value for col in range(1, 13)]
        received_on, received_at = cells[1], cells[2]
        if not isinstance(received_on, datetime) or not isinstance(received_at, time):
            continue
        rows.append(
            LogRow(
                number=number,
                received_on=received_on.date(),
                received_at=received_at,
                product_code=_text(cells[3]),
                product_name=_text(cells[4]),
                type=_text(cells[5]),
                lot=_text(cells[6]),
                client=_text(cells[7]),
                quantity=_text(cells[9]),
                received_by=_text(cells[11]),
            )
        )
    return rows


def _submission(row: LogRow) -> tuple[date, str]:
    return row.received_on, row.client


def _forms(rows: list[LogRow]) -> list[list[LogRow]]:
    """The rows in forms of at most `FORM_LINES`, each run of lines with one
    Lot and Product Code kept on one form while it fits."""
    forms: list[list[LogRow]] = [[]]
    for _, run in groupby(rows, key=lambda r: (r.lot, r.product_code)):
        lines = list(run)
        if forms[-1] and len(forms[-1]) + len(lines) > FORM_LINES:
            forms.append([])
        for line in lines:
            if len(forms[-1]) == FORM_LINES:
                forms.append([])
            forms[-1].append(line)
    return forms


def submissions(rows: list[LogRow]) -> list[tuple[str, list[LogRow]]]:
    """Each form's stem and rows: one submission, at most a form's worth of rows."""
    out = []
    ordered = sorted(rows, key=lambda r: (_submission(r), r.number))
    for (on, _), group in groupby(ordered, key=_submission):
        forms = _forms(list(group))
        stem = f"received-{on.isoformat()}"
        if len(forms) == 1:
            out.append((stem, forms[0]))
        else:
            out.extend((f"{stem}-form-{n}", f) for n, f in enumerate(forms, start=1))
    return out


class Inventor:
    """Every value the log does not hold, from one seeded generator."""

    def __init__(self, seed: int) -> None:
        self.rng = random.Random(seed)

    def party(self, company: str, city: str) -> Party:
        rng = self.rng
        first, last = rng.choice(FIRST_NAMES), rng.choice(LAST_NAMES)
        return Party(
            company=company,
            contact=f"{first} {last}",
            street=f"{rng.randint(10, 990)} {rng.choice(STREETS)}",
            city=city,
            phone=f"(201) 555-01{rng.randint(10, 99)}",
            email=f"{first.lower()}.{last.lower()}@{company.lower()}.example",
        )

    def city(self) -> str:
        return self.rng.choice(CLIENTS)[3]

    def tests(self, sample_type: str) -> tuple[str, str]:
        if sample_type == "FP":
            return TEST_REQUESTED, ""
        analyte = "".join(self.rng.choices(string.ascii_uppercase, k=7))
        limit = f"{analyte} < {self.rng.uniform(10, 99):.1f}ng/day"
        return self.rng.choice(OTHER_TESTS), limit

    def weighings(self, containers: int) -> tuple[dict[str, float | int], ...]:
        rng = self.rng
        capsule_g = rng.uniform(0.35, 0.85)
        return tuple(
            {
                "container": n,
                "weight_of_20_g": round(20 * capsule_g * rng.gauss(1, 0.01), 4),
                "weight_taken_g": round(capsule_g * rng.uniform(0.9, 1.1), 4),
            }
            for n in range(1, containers + 1)
        )


def _sample_type(row: LogRow) -> str:
    return row.type if row.type in SAMPLE_TYPES else "FP"


def _mistyped(lot: str) -> str:
    last = lot[-1].upper()
    pool = string.digits if last.isdigit() else string.ascii_uppercase
    return lot[:-1] + pool[(pool.index(last) + 1) % len(pool)]


def _form(
    rows: list[LogRow], place: int, inventor: Inventor, tat_code: str
) -> FilledForm:
    first = rows[0]
    city = inventor.city()
    submitter = inventor.party(first.client, city)
    report_to = inventor.party(first.client, city) if place % 2 == 0 else submitter
    bill_to = inventor.party(first.client, city) if place // 2 % 2 else None
    lines = []
    for row in rows:
        sample_type = _sample_type(row)
        tests, limits = inventor.tests(sample_type)
        quantity = QUANTITY.match(row.quantity)
        if quantity is None:
            raise ValueError(f"row {row.number}: QTY RECD {row.quantity!r} has no unit")
        lines.append(
            SampleLine(
                description=row.product_name,
                sample_type=sample_type,
                tat_code=tat_code,
                lot=row.lot,
                product_code=row.product_code,
                tests=tests,
                containers=1,
                product_type=quantity.group(2),
                limits=limits,
            )
        )
    return FilledForm(
        submitted_on=first.received_on - timedelta(days=inventor.rng.randint(1, 3)),
        submitter=submitter,
        bill_to=bill_to,
        report_to=report_to,
        lines=tuple(lines),
        received_on=first.received_on,
        received_at=first.received_at,
        received_by=first.received_by,
        document_time=datetime.combine(first.received_on, time(8, 0)),
    )


def _unlogged(rows: list[LogRow], rng: random.Random) -> LogRow:
    """An IP Sample on a weekday, with a Lot and Product Code the log does not hold."""
    dates = {r.received_on for r in rows}
    received_on = rows[0].received_on
    while received_on in dates or received_on.weekday() >= 5:
        received_on = date(2026, 1, 5) + timedelta(days=rng.randint(0, 300))
    lots, codes = {r.lot for r in rows}, {r.product_code for r in rows}
    lot = code = ""
    while lot in lots | {""}:
        lot = f"{rng.choice(string.ascii_uppercase)}{rng.randint(0, 999999):06d}"
    while code in codes | {""}:
        code = f"{rng.randint(0, 9999):04d}{rng.choice(string.ascii_uppercase)}"
    return replace(
        rows[0],
        number=0,
        received_on=received_on,
        received_at=time(rng.randint(8, 16), rng.randint(0, 59)),
        product_code=code,
        product_name=rng.choice(PRODUCTS),
        type="IP",
        lot=lot,
    )


def improvise(log: Path) -> list[Improvised]:
    rows = read_log(log)
    inventor = Inventor(SEED)
    planned = submissions(rows)
    planned.append(("unlogged-sample", [_unlogged(rows, inventor.rng)]))
    tat_codes = sorted(TAT_CODES)
    out = []
    for place, (stem, group) in enumerate(planned):
        form = _form(group, place, inventor, tat_codes[place % len(tat_codes)])
        if place == 0:
            stem += "-lot-mistyped"
            lines = tuple(replace(line, lot=_mistyped(line.lot)) for line in form.lines)
            form = replace(form, lines=lines)
        containers = sum(line.containers for line in form.lines)
        out.append(Improvised(stem, form, inventor.weighings(containers)))
    return out


def _party(party: Party) -> dict[str, str]:
    return {
        "company": party.company,
        "contact": party.contact,
        "address": party.street,
        "city": party.city,
        "phone": party.phone,
        "email": party.email,
    }


def expected_record(form: FilledForm) -> dict:
    """The `submission/2` record the form carries, as `parse_submission` reads it."""
    submitted_by = _party(form.submitter)
    if form.bill_to is None:
        bill_to = dict(submitted_by)
    else:
        bill_to = {**_party(form.bill_to), "phone": "", "email": ""}
    report_to = {
        **{k: submitted_by[k] for k in ("company", "address", "city")},
        "contact": form.report_to.contact,
        "phone": form.report_to.phone,
        "email": form.report_to.email,
    }
    return {
        "schema": "submission/2",
        "date_submitted": form.submitted_on.isoformat(),
        "submitted_by": submitted_by,
        "bill_to": bill_to,
        "report_to": report_to,
        "samples": [
            {
                "no": str(number),
                "sample_description": line.description,
                "sample_type": line.sample_type,
                "tat_code": line.tat_code,
                "tat_class": str(TAT_CODES[line.tat_code]),
                "lot": line.lot,
                "product_code": line.product_code,
                "product_type": line.product_type,
                "no_of_containers": str(line.containers),
            }
            for number, line in enumerate(form.lines, start=1)
        ],
        "receipt": {
            "date_received": form.received_on.isoformat(),
            "time_received": f"{form.received_at:%H:%M}",
            "received_by": form.received_by,
        },
    }


def write(log: Path, blank: Path, out: Path) -> list[str]:
    out.mkdir(parents=True, exist_ok=True)
    stems = []
    for item in improvise(log):
        write_filled_form(item.form, blank, out / f"{item.stem}.docx")
        (out / f"{item.stem}.expected.json").write_text(
            json.dumps(expected_record(item.form), indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )
        (out / f"{item.stem}.weighings.json").write_text(
            json.dumps(list(item.weighings), indent=2) + "\n", encoding="utf-8"
        )
        stems.append(item.stem)
    return stems


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(
        prog="improvise",
        description="Fill the real-layout Submission Form from the Receiving Log.",
    )
    parser.add_argument("--log", type=Path, required=True)
    parser.add_argument("--blank", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args(argv)
    for stem in write(args.log, args.blank, args.out):
        print(stem)


if __name__ == "__main__":
    main()
