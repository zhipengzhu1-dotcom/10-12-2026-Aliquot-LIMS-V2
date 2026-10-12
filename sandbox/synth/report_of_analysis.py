"""The Report of Analysis (xlsx) for one Sample, with static values.

Top to bottom: letterhead, Client block, metadata block, results table, signature
area. Labels sit in column A and values in column B for the Client and metadata
blocks. The metadata values copy the first Receiving Log row of the Sample.
"""

from pathlib import Path

from openpyxl.styles import Font

from synth.lab import AVERAGE_WEIGHT_METHOD, LAB_ADDRESS, LAB_CONTACT, LAB_NAME
from synth.record import TEST_REQUESTED, Record
from synth.workbook import new_workbook, save_workbook

SHEET = "Report of Analysis"
RESULT_COLUMNS = ("Test", "Method", "Specification", "Result", "Unit")
SIGNATURE_ROLES = ("Analyzed By", "Reviewed By", "Approved By")


def write_report_of_analysis(record: Record, path: Path) -> None:
    wb = new_workbook(record)
    ws = wb.active
    ws.title = SHEET

    ws["A1"] = LAB_NAME
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = LAB_ADDRESS
    ws["A3"] = LAB_CONTACT
    ws["A5"] = "REPORT OF ANALYSIS"
    ws["A5"].font = Font(size=12, bold=True)

    client_block = (
        ("Client", record.client),
        ("Attention", record.report_to.contact),
        ("Address", record.report_to.street),
        ("", record.report_to.city),
    )
    metadata_block = (
        ("Lab Reference", record.lab_reference),
        ("Product", record.product_name),
        ("Lot", record.lot),
        ("Product Code", record.product_code),
        ("Product Type", record.product_type),
        ("Document Revision", record.document_revision),
        ("Date Received", record.received_on),
        ("Time Received", record.received_at),
        ("Received By", record.received_by),
    )
    for row, (label, value) in enumerate(client_block, start=7):
        ws.cell(row, 1, label)
        ws.cell(row, 2, value)
    for row, (label, value) in enumerate(metadata_block, start=12):
        ws.cell(row, 1, label).font = Font(bold=True)
        ws.cell(row, 2, value)
    ws["B18"].number_format = "mm/dd/yyyy"
    ws["B19"].number_format = "hh:mm"

    results_header = 22
    for col, title in enumerate(RESULT_COLUMNS, start=1):
        ws.cell(results_header, col, title).font = Font(bold=True)
    result = (
        TEST_REQUESTED,
        AVERAGE_WEIGHT_METHOD,
        "Report value",
        round(record.mean_average_weight_g, 4),
        "g",
    )
    for col, value in enumerate(result, start=1):
        ws.cell(results_header + 1, col, value)

    signers = (record.analyst, record.reviewer, record.approver)
    for i, (role, name) in enumerate(zip(SIGNATURE_ROLES, signers, strict=True)):
        col = 1 + 2 * i
        ws.cell(26, col, role).font = Font(bold=True)
        ws.cell(27, col, name)
        ws.cell(28, col, "Signature: ______________")
        ws.cell(29, col, "Date: ______________")

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 44
    save_workbook(wb, record, path)
