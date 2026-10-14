"""Fill the scrubbed blank Submission Form, in the lab's real layout, with a record.

The blank is `fixtures/real-layout/submission-form.docx`. Its tables are found
by the labels the scrub kept and each value is written into the cell the lab
fills by hand, so the form keeps every table, merge and style of the blank.
How the record lands in the real form, which `parse_submission` reads back:

- First table, top row: the date submitted and the submitter's contact follow
  their labels. The PO# stays blank.
- First table, `COMPANY NAME & ADDRESS`: the submitter's company, street, city,
  phone and email, one per line.
- First table, `BILL TO ADDRESS (if different)`: the bill-to company,
  `Attn: <contact>`, street and city, one per line, or nothing when the
  bill-to party is the submitter's.
- First table, `Attn`, `Phone #`, `Email ID`: the report-to contact, phone and
  email. The real form has no cell for the report-to company or address.
- Sample table: one row per Sample line from the first row, each with its No.,
  description, Sample Type, TAT code, Lot, Product Code, Test requested and
  limits. `parse_submission` reads the first line only.
- Lab-use table: the date and time received and the receiver's initials follow
  their labels, and `Comments:` holds the count received and the package as
  `<containers> x <Product Type>`.
"""

from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path

from docx import Document
from docx.table import Table, _Cell

from synth.record import TEST_REQUESTED, Party, Record
from synth.reproducible import save_reproducibly

COMPANY = "COMPANY NAME & ADDRESS"
BILL_TO = "BILL TO ADDRESS"
SAMPLE_DESCRIPTION = "SAMPLE DESCRIPTION"
DATE_RECEIVED = "Date Sample(s) received:"


@dataclass(frozen=True)
class SampleLine:
    description: str
    sample_type: str
    tat_code: str
    lot: str
    product_code: str
    tests: str
    limits: str = ""


@dataclass(frozen=True)
class FilledForm:
    """What a Client and the receiver write on one real-layout form.

    `bill_to` is None when the bill-to block is left blank, which the form's
    label reads as the submitter's party.
    """

    submitted_on: date
    submitter: Party
    bill_to: Party | None
    report_to: Party
    lines: tuple[SampleLine, ...]
    received_on: date
    received_at: time
    received_by: str
    containers: int
    product_type: str
    document_time: datetime


def _texts(row) -> list[str]:
    return [cell.text.strip() for cell in row.cells]


def _table_with(tables: list[Table], label: str) -> Table:
    for table in tables:
        for row in table.rows:
            if any(text.startswith(label) for text in _texts(row)):
                return table
    raise ValueError(f"the blank form has no table labelled {label!r}")


def _write(cell: _Cell, *lines: str) -> None:
    """Append the lines to the cell's first paragraph, after any label in it."""
    paragraph = cell.paragraphs[0]
    for i, line in enumerate(lines):
        run = paragraph.add_run(line)
        if i < len(lines) - 1:
            run.add_break()


def _after_label(row, label: str, value: str) -> None:
    for cell in dict.fromkeys(row.cells):
        if cell.text.strip().startswith(label):
            _write(cell, f" {value}")
            return
    raise ValueError(f"the blank form has no cell labelled {label!r}")


def write_filled_form(form: FilledForm, blank: Path, path: Path) -> None:
    doc = Document(str(blank))
    tables = list(doc.tables)

    parties = _table_with(tables, COMPANY)
    top = parties.rows[0]
    _after_label(top, "Date submitted:", f"{form.submitted_on:%m/%d/%Y}")
    _after_label(top, "Submitted By:", form.submitter.contact)
    headings = _texts(parties.rows[1])
    blocks = parties.rows[2].cells
    submitter, bill_to, report_to = form.submitter, form.bill_to, form.report_to
    _write(
        blocks[headings.index(COMPANY)],
        submitter.company,
        submitter.street,
        submitter.city,
        submitter.phone,
        submitter.email,
    )
    if bill_to is not None:
        bill_column = next(i for i, h in enumerate(headings) if h.startswith(BILL_TO))
        _write(
            blocks[bill_column],
            bill_to.company,
            f"Attn: {bill_to.contact}",
            bill_to.street,
            bill_to.city,
        )
    send_to = {
        "Attn": report_to.contact,
        "Phone #": report_to.phone,
        "Email ID": report_to.email,
    }
    for row in parties.rows[2:]:
        cells = list(dict.fromkeys(row.cells))
        for i, cell in enumerate(cells[:-1]):
            value = send_to.get(cell.text.strip())
            if value is not None:
                _write(cells[i + 1], value)

    samples = _table_with(tables, SAMPLE_DESCRIPTION)
    rows = samples.rows[1:]
    if len(form.lines) > len(rows):
        raise ValueError(
            f"the blank form has {len(rows)} Sample rows, not {len(form.lines)}"
        )
    for number, line in enumerate(form.lines, start=1):
        row = rows[number - 1]
        values = (
            str(number),
            line.description,
            line.sample_type,
            line.tat_code,
            line.lot,
            line.product_code,
            line.tests,
            line.limits,
        )
        for cell, value in zip(row.cells, values, strict=True):
            if value:
                _write(cell, value)

    receipt = _table_with(tables, DATE_RECEIVED)
    _after_label(receipt.rows[0], DATE_RECEIVED, f"{form.received_on:%m/%d/%Y}")
    _after_label(receipt.rows[0], "Time received:", f"{form.received_at:%H:%M}")
    _after_label(receipt.rows[0], "Received By:", form.received_by)
    _after_label(
        receipt.rows[1], "Comments:", f"{form.containers} x {form.product_type}"
    )

    save_reproducibly(doc.save, path, form.document_time)


def write_real_layout_submission_form(record: Record, blank: Path, path: Path) -> None:
    line = SampleLine(
        description=record.product_name,
        sample_type=record.sample_type,
        tat_code=record.tat_code,
        lot=record.lot,
        product_code=record.product_code,
        tests=TEST_REQUESTED,
    )
    form = FilledForm(
        submitted_on=record.submitted_on,
        submitter=record.submitter,
        bill_to=record.bill_to,
        report_to=record.report_to,
        lines=(line,),
        received_on=record.received_on,
        received_at=record.received_at,
        received_by=record.received_by,
        containers=len(record.containers),
        product_type=record.product_type,
        document_time=record.document_time,
    )
    write_filled_form(form, blank, path)
